"""QA-03: exportar inventario con el atributo ORM unidad_medida."""

from datetime import date
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock, patch

from flask import Flask
from flask_jwt_extended import JWTManager, create_access_token
from openpyxl import load_workbook

from app.routes import reportes


def test_exportar_inventario_excel_usa_unidad_medida():
    app = Flask(__name__)
    app.config.update(TESTING=True, JWT_SECRET_KEY="inventario-qa03-synthetic")
    JWTManager(app)
    app.register_blueprint(reportes.reportes_bp)

    producto = SimpleNamespace(
        id=17,
        tipo="producto",
        nombre="Producto QA",
        descripcion="Sintético",
        marca="Marca",
        proveedor="Proveedor",
        categoria="Consumibles",
        unidad_medida="pieza",
        codigo_interno="QA-17",
        no_equipo=None,
        gasto_sem=0,
        gasto_mes=0,
        pedido_mes=0,
        semana_pedido=None,
        fecha_inventario=date(2026, 10, 21),
    )
    with app.app_context():
        token = create_access_token(identity="1")

    with patch.object(reportes, "_require_inventory_report_admin", return_value=None), \
         patch.object(
             reportes,
             "_analytical_inventory_general_query",
             return_value=Mock(all=Mock(return_value=[producto])),
         ):
        response = app.test_client().get(
            "/api/reportes/exportar-inventario",
            headers={"Authorization": f"Bearer {token}"},
        )

    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.data), read_only=True).active
    assert sheet["H1"].value == "Unidad"
    assert sheet["H2"].value == "pieza"
    assert "Inventario_GENERAL_" in response.headers["Content-Disposition"]



def test_exportar_inventario_sucursal_usa_scope_y_nombre():
    """Un archivo llamado por sucursal procede de una consulta de esa sucursal."""
    app = Flask(__name__)
    app.config.update(TESTING=True, JWT_SECRET_KEY="inventario-branch-qa-synthetic")
    JWTManager(app)
    app.register_blueprint(reportes.reportes_bp)

    producto = SimpleNamespace(
        id=33, tipo="producto", nombre="Artículo de Villas del Rey",
        descripcion="", marca="", proveedor="", categoria="Consumibles",
        unidad_medida="pieza", codigo_interno="QA-VR-33", no_equipo=None,
        gasto_sem=0, gasto_mes=0, pedido_mes=0, semana_pedido=None,
        fecha_inventario=date(2026, 10, 21),
    )
    scoped_query = Mock(all=Mock(return_value=[producto]))
    sucursal = SimpleNamespace(sucursal="Villas del Rey")
    with app.app_context():
        token = create_access_token(identity="1")

    with patch.object(reportes, "_require_inventory_report_admin", return_value=None), \
         patch.object(
             reportes,
             "_inventario_export_query_por_sucursal",
             return_value=(sucursal, scoped_query),
         ) as scoped, \
         patch.object(
             reportes,
             "_analytical_inventory_general_query",
             side_effect=AssertionError("No debe exportar el catálogo general"),
         ):
        response = app.test_client().get(
            "/api/reportes/exportar-inventario?sucursal_id=1",
            headers={"Authorization": f"Bearer {token}"},
        )

    assert response.status_code == 200
    scoped.assert_called_once_with(1)
    assert "Inventario_VILLAS_DEL_REY_" in response.headers["Content-Disposition"]
    sheet = load_workbook(BytesIO(response.data), read_only=True).active
    assert sheet["C2"].value == "Artículo de Villas del Rey"


def test_exportar_inventario_rechaza_sucursal_id_invalido():
    app = Flask(__name__)
    app.config.update(TESTING=True, JWT_SECRET_KEY="inventario-bad-branch-qa")
    JWTManager(app)
    app.register_blueprint(reportes.reportes_bp)
    with app.app_context():
        token = create_access_token(identity="1")

    with patch.object(reportes, "_require_inventory_report_admin", return_value=None):
        response = app.test_client().get(
            "/api/reportes/exportar-inventario?sucursal_id=abc",
            headers={"Authorization": f"Bearer {token}"},
        )

    assert response.status_code == 400
