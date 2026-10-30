"""E4 Bloque 1: compatibilidad v0.1/v0.2, evidencia histórica y no mutación V1."""
import csv
from hashlib import sha256
from io import BytesIO, StringIO
from uuid import uuid4
import importlib.util
import os
import secrets
from pathlib import Path

import pytest
import sqlalchemy as sa
from alembic.config import Config as AlembicConfig
from alembic.script import ScriptDirectory
from alembic.migration import MigrationContext
from alembic.operations import Operations
from flask import Flask
from flask_jwt_extended import JWTManager, create_access_token
from openpyxl import Workbook, load_workbook
from zipfile import ZipFile
import re

from app.extensions import db
from app.routes.inventario_levantamientos_routes import inventario_levantamientos_bp
from app.services.inventario_levantamiento_excel_parser import (
    parse_historical_csv, parse_xlsx_v01,
)

BASE = Path(__file__).resolve().parents[2]
MIGRATION = BASE / "migrations/versions/v2e4b1102826_inventario_levantamiento_recepcion.py"


def original_csv():
    """Encabezados del CSV histórico real de distribución (28/10/2026)."""
    return (
        "sucursal_id,sucursal,inventario_id,codigo_interno,categoria,aparato,marca,stock_sistema_qa\n"
        "7,Pabellon rosarito,501,07CA01,Cardio,Caminadora,Star Trac,3\n"
        "7,Pabellon rosarito,502,07EL01,Cardio,Eliptica,Precor,1\n"
        "7,Pabellon rosarito,503,07RE00,Cardio,Remo,Concept2,0\n"
        "8,Mision ensenada,601,08CA01,Cardio,Caminadora,Star Trac,1\n"
    ).encode()


def workbook(branch=7, second="", first="SI", code="07EL01", duplicate=False,
             initial=False, found=True, version="v0.1 · formato"):
    """Reconstruye las posiciones A5/C6, encabezado fila 7 de la plantilla oficial v0.1."""
    wb = Workbook()
    ctl = wb.active
    ctl.title = "CONTROL"
    for r, name, val in (
        (5, "Sucursal", "Pabellon rosarito"),
        (6, "ID sucursal", branch),
        (7, "Gerente responsable", "Responsable QA"),
        (8, "Fecha del recorrido", "2026-10-28"),
        (9, "Folio / corte E4", "PENDIENTE - GENERADOR E4"),
        (10, "Versión de entrega", version),
    ):
        ctl.cell(r, 1, name)
        ctl.cell(r, 3, val)
    a = wb.create_sheet("APARATOS")
    a.append(["APARATOS · VERIFICACIÓN FÍSICA POR CÓDIGO"])
    for r in range(2, 7):
        a.cell(r, 1, "Instrucciones" if r == 2 else None)
    a.append(["ID oficial", "CÓDIGO INTERNO", "CATEGORÍA", "APARATO",
              "MARCA", "¿ESTÁ? SI / NO", "OBSERVACIONES / UBICACIÓN"])
    if not initial:
        a.append([501, "07CA01", "Cardio", "Caminadora", "Star Trac", first, "Zona cardio"])
        a.append([502, code, "Cardio", "Eliptica", "Precor", second, ""])
        if duplicate:
            a.append([501, "07CA01", "Cardio", "Caminadora", "Star Trac", "NO", ""])
    else:
        a.append([None, "SIN CÓDIGO", "Cardio", "Caminadora", "Star Trac", "SI", "Alta inicial sin ID"])
    h = wb.create_sheet("HALLAZGOS")
    for r in range(1, 7):
        h.cell(r, 1, "HALLAZGOS · APARATOS NO INCLUIDOS" if r == 1 else None)
    h.append(["CÓDIGO EN ETIQUETA", "NOMBRE / APARATO", "CATEGORÍA",
              "MARCA", "UBICACIÓN / ZONA", "DESCRIPCIÓN DEL HALLAZGO"])
    if found:
        h.append(["SIN CÓDIGO", "Banco no registrado", "Peso Libre",
                  "Desconocida", "Zona pesas", "Encontrado sin referencia oficial"])
    target = BytesIO()
    wb.save(target)
    return target.getvalue()


def test_e4_cadena_alembic_no_abre_rama_paralela():
    cfg = AlembicConfig()
    cfg.set_main_option("script_location", str(BASE / "migrations"))
    script = ScriptDirectory.from_config(cfg)
    assert script.get_heads() == ["n26c10202626"]
    assert script.get_revision("n26c10202626").down_revision == "v2e4b1102826"


def test_plantilla_v01_parser_y_sin_revisar_por_id():
    rows, excluded = parse_historical_csv(original_csv(), 7)
    assert len(rows) == 2 and len(excluded) == 1
    report = parse_xlsx_v01(workbook(), 7, rows, "VERIFICACION")
    assert report["resultado"] == "INCOMPLETE"
    assert report["resumen"]["si"] == 1
    assert report["resumen"]["no"] == 0
    assert report["resumen"]["sin_revisar"] == 1
    assert report["resumen"]["pendientes"][0]["inventario_id"] == 502
    assert report["resumen"]["pendientes"][0]["fila_excel"] == 9
    assert report["resumen"]["hallazgos"] == 1
    assert report["control"]["sucursal_id"] == "7"


def test_v02_real_distribuida_vacia_se_mantiene_pendiente_y_versiones_desconocidas_invalidas():
    # v0.2 es la etiqueta real del XLSX distribuido el 28/10/2026;
    # el fixture no contiene datos reales ni archivos de gerentes.
    expected, _ = parse_historical_csv(original_csv(), 7)
    received = parse_xlsx_v01(
        workbook(first="", second="", found=False,
                 version="v0.2 · base CSV 2026-10-28"),
        7, expected, "VERIFICACION",
    )
    assert received["resultado"] == "INCOMPLETE"
    assert received["control"]["version_formato"] == "v0.2 · base CSV 2026-10-28"
    assert received["resumen"]["referencias_esperadas"] == 2
    assert received["resumen"]["si"] == 0
    assert received["resumen"]["no"] == 0
    assert received["resumen"]["sin_revisar"] == 2
    assert not received["diagnosticos"]

    unsupported = parse_xlsx_v01(
        workbook(first="", second="", found=False,
                 version="v0.3 · formato no certificado"),
        7, expected, "VERIFICACION",
    )
    assert unsupported["resultado"] == "INVALID"
    assert any(d["codigo"] == "VERSION_FORMATO_NO_RECONOCIDA"
               for d in unsupported["diagnosticos"])


def test_xlsx_sin_dimensiones_como_el_distribuido_real():
    """Los libros distribuidos de E4 carecen de XML <dimension>: max_row=None."""
    original = workbook(
        first="", second="", found=False, version="v0.2 · base CSV 2026-10-28"
    )
    output = BytesIO()
    hojas = 0
    with ZipFile(BytesIO(original)) as z_in, ZipFile(output, "w") as z_out:
        for info in z_in.infolist():
            content = z_in.read(info.filename)
            if re.fullmatch(r"xl/worksheets/sheet[0-9]+\.xml", info.filename):
                content, n = re.subn(rb"<dimension\b[^>]*/>", b"", content, count=1)
                assert n == 1
                hojas += 1
            z_out.writestr(info, content)
    assert hojas == 3
    raw = output.getvalue()
    book = load_workbook(BytesIO(raw), read_only=True, data_only=False)
    try:
        assert all(book[sheet].max_row is None for sheet in ("CONTROL", "APARATOS", "HALLAZGOS"))
    finally:
        book.close()
    expected, _ = parse_historical_csv(original_csv(), 7)
    report = parse_xlsx_v01(raw, 7, expected, "VERIFICACION")
    assert report["resultado"] == "INCOMPLETE"
    assert report["resumen"]["sin_revisar"] == 2
    assert report["resumen"]["si"] == 0
    assert report["resumen"]["no"] == 0
    assert not report["diagnosticos"]


def test_plantilla_rechaza_codigo_diferente_duplicados_y_sucursal_falsa():
    expected, _ = parse_historical_csv(original_csv(), 7)
    report = parse_xlsx_v01(workbook(branch=8, code="OTRO", duplicate=True), 7, expected, "VERIFICACION")
    assert report["resultado"] == "INVALID"
    codes = {x["codigo"] for x in report["diagnosticos"]}
    assert {"SUCURSAL_NO_COINCIDE", "CODIGO_DIFIERE_DEL_CORTE", "ID_DUPLICADO_EN_EXCEL"}.issubset(codes)
    assert report["resumen"]["sin_revisar"] == 2


def test_hallazgo_no_puede_reciclar_codigo_de_aparato_oficial():
    refs, _ = parse_historical_csv(original_csv(), 7)
    book = load_workbook(BytesIO(workbook(second="SI", found=True)))
    book["HALLAZGOS"]["A8"] = "07CA01"
    raw = BytesIO()
    book.save(raw)
    resultado = parse_xlsx_v01(raw.getvalue(), 7, refs, "VERIFICACION")
    assert resultado["resultado"] == "INVALID"
    assert resultado["resumen"]["errores"] == 1
    assert any(d["codigo"] == "CODIGO_HALLAZGO_EN_CORTE"
               and d["inventario_ids_existentes"] == [501]
               for d in resultado["diagnosticos"])


def test_prefijo_sucursal_ajeno_es_aviso_no_mutacion_de_codigo():
    historical = original_csv().replace(b"07EL01", b"05EL01")
    refs, _ = parse_historical_csv(historical, 7)
    book = load_workbook(BytesIO(workbook(second="SI", found=False)))
    book["APARATOS"]["B9"] = "05EL01"
    raw = BytesIO()
    book.save(raw)
    resultado = parse_xlsx_v01(raw.getvalue(), 7, refs, "VERIFICACION")
    assert resultado["resultado"] == "VALIDATED"
    assert resultado["resumen"]["avisos"] == 1
    assert any(d["codigo"] == "PREFIJO_SUCURSAL_INCONSISTENTE"
               and d["inventario_id"] == 502 and d["prefijo_esperado"] == "07"
               for d in resultado["diagnosticos"])
    assert resultado["resumen"]["si"] == 2


def test_inicial_no_equivale_a_conciliado_sin_referencias():
    report = parse_xlsx_v01(workbook(branch=8, initial=True), 8, [], "INICIAL")
    assert report["resultado"] == "INCOMPLETE"
    assert report["resumen"]["requiere_revision_inicial"] is True
    assert report["resumen"]["hallazgos"] == 1
    assert report["resumen"]["provisionales_aparatos"] == 1
    assert report["filas"][0]["inventario_id_declarado"] is None


@pytest.fixture()
def api(monkeypatch, tmp_path):
    url = os.getenv("INVENTARIO_E2_TEST_POSTGRES_URL")
    using_real_binary = bool(os.getenv("E4_REAL_XLSX_PATH") or os.getenv("E4_REAL_BASELINE_CSV_PATH"))
    if not url:
        if using_real_binary:
            pytest.fail("Gate real requiere INVENTARIO_E2_TEST_POSTGRES_URL aislado")
        pytest.skip("PostgreSQL efímero requerido para la API E4")
    if using_real_binary:
        parsed_url = sa.engine.make_url(url)
        if (parsed_url.host not in {"127.0.0.1", "localhost", "::1"}
                or not re.fullmatch(r"inventario_e4_(?:ci|qa)(?:_[a-z0-9_]+)?", parsed_url.database or "")):
            pytest.fail("Gate real SOLO permite PostgreSQL local y BD inventario_e4_qa/ci")
    schema = "inv_e4_" + secrets.token_hex(5)
    admin = sa.create_engine(url, future=True)
    with admin.begin() as con:
        con.exec_driver_sql(f"CREATE SCHEMA {schema}")
    monkeypatch.setenv("E4_UPLOAD_DIR", str(tmp_path / "e4-private"))
    app = Flask(__name__)
    app.config.update(
        TESTING=True, JWT_SECRET_KEY="e4-qa-isolated-only",
        SQLALCHEMY_DATABASE_URI=url,
        SQLALCHEMY_ENGINE_OPTIONS={"connect_args": {"options": f"-csearch_path={schema}"}},
    )
    db.init_app(app)
    JWTManager(app)
    app.register_blueprint(inventario_levantamientos_bp, url_prefix="/api/inventario/levantamientos")
    spec = importlib.util.spec_from_file_location("inv_e4_migration", MIGRATION)
    migration = importlib.util.module_from_spec(spec)
    assert spec and spec.loader
    spec.loader.exec_module(migration)
    try:
        with app.app_context():
            with db.engine.begin() as con:
                con.exec_driver_sql("CREATE TABLE sucursales (sucursal_id INTEGER PRIMARY KEY)")
                con.exec_driver_sql("CREATE TABLE users (id INTEGER PRIMARY KEY, rol VARCHAR(50))")
                con.exec_driver_sql("CREATE TABLE inventario_general (id INTEGER PRIMARY KEY, codigo_interno VARCHAR(100), tipo VARCHAR(20), nombre VARCHAR(255), marca VARCHAR(100), categoria VARCHAR(100))")
                con.exec_driver_sql("CREATE TABLE inventario_sucursal (inventario_id INTEGER, sucursal_id INTEGER, stock INTEGER, PRIMARY KEY(inventario_id,sucursal_id))")
                con.exec_driver_sql("CREATE TABLE movimientos_inventario (id INTEGER PRIMARY KEY)")
                con.exec_driver_sql("CREATE TABLE inventario_correccion_solicitudes (id INTEGER PRIMARY KEY)")
                monkeypatch.setattr(migration, "op", Operations(MigrationContext.configure(con)))
                migration.upgrade()
                con.exec_driver_sql("INSERT INTO sucursales VALUES (7),(8)")
                con.exec_driver_sql("INSERT INTO users VALUES (1,'ADMINISTRADOR'),(2,'GERENTE'),(3,'GERENTE_REGIONAL')")
                con.exec_driver_sql("INSERT INTO inventario_general (id,codigo_interno,tipo,nombre,marca) VALUES (501,'07CA01','APARATOS','Caminadora','Star Trac'),(502,'07EL01','APARATOS','Eliptica','Precor'),(503,'07RE00','APARATOS','Remo','Concept2')")
                con.exec_driver_sql("INSERT INTO inventario_sucursal VALUES (501,7,3),(502,7,1),(503,7,0)")
            token = {uid: {"Authorization": "Bearer " + create_access_token(identity=str(uid))}
                     for uid in (1,2,3)}
            yield app.test_client(), token, app
            db.session.remove()
            db.engine.dispose()
    finally:
        with admin.begin() as con:
            con.exec_driver_sql(f"DROP SCHEMA IF EXISTS {schema} CASCADE")
        admin.dispose()


def _snapshot(app):
    with app.app_context():
        return {
            "stock": db.session.execute(sa.text("SELECT inventario_id,stock FROM inventario_sucursal ORDER BY inventario_id")).all(),
            "master": db.session.execute(sa.text("SELECT id,codigo_interno FROM inventario_general ORDER BY id")).all(),
            "mov": db.session.execute(sa.text("SELECT count(*) FROM movimientos_inventario")).scalar(),
            "e2": db.session.execute(sa.text("SELECT count(*) FROM inventario_correccion_solicitudes")).scalar(),
        }


def test_api_historico_versionado_idempotente_sin_cambios_oficiales(api):
    client, tokens, app = api
    url = "/api/inventario/levantamientos"
    assert client.post(url, headers=tokens[2], data={"sucursal_id":"7"}).status_code == 403
    assert client.post(url, headers=tokens[3], data={"sucursal_id":"7"}).status_code == 403
    assert client.post(url, data={"sucursal_id":"7"}).status_code == 401
    # Bandeja corporativa: acceso, paginación estricta y filtros.
    assert client.get(url).status_code == 401
    assert client.get(url, headers=tokens[2]).status_code == 403
    assert client.get(url, headers=tokens[3]).status_code == 403
    assert client.get(url + "?per_page=101", headers=tokens[1]).status_code == 400
    assert client.get(url + "?page=0", headers=tokens[1]).status_code == 400
    assert client.get(url + "?sucursal_id=abc", headers=tokens[1]).status_code == 400
    empty = client.get(url, headers=tokens[1])
    assert empty.status_code == 200 and empty.get_json()["total"] == 0
    original = _snapshot(app)
    basekey = str(uuid4())
    created = client.post(url, headers=tokens[1], data={
        "sucursal_id": "7", "origen_corte": "CSV_HISTORICO", "cycle_key": basekey,
        "fecha_corte_local": "2026-10-28", "responsable": "Gerencia sucursal",
        "baseline_csv": (BytesIO(original_csv()), "Inventarios para revisar.csv"),
    })
    assert created.status_code == 201, created.get_json()
    lev = created.get_json()["levantamiento"]["id"]
    assert created.get_json()["levantamiento"]["baseline_resumen"]["excluidas_stock_cero"] == 1
    listing = client.get(url + "?sucursal_id=7", headers=tokens[1])
    assert listing.status_code == 200, listing.get_json()
    assert listing.get_json()["total"] == 1
    assert listing.get_json()["items"][0]["id"] == lev
    assert listing.get_json()["items"][0]["importacion_actual"] is None
    assert client.get(url + "?sucursal_id=8", headers=tokens[1]).get_json()["total"] == 0
    first_bytes = workbook()
    endpoint = f"{url}/{lev}/importaciones"
    key = str(uuid4())
    first = client.post(endpoint, headers=tokens[1], data={
        "idempotency_key": key, "file": (BytesIO(first_bytes), "pabellon.xlsx"),
    })
    assert first.status_code == 201, first.get_json()
    first_data = first.get_json()["importacion"]
    assert first_data["resultado"] == "INCOMPLETE"
    assert first_data["resumen"]["pendientes"][0]["inventario_id"] == 502
    replay = client.post(endpoint, headers=tokens[1], data={
        "idempotency_key": str(uuid4()), "file": (BytesIO(first_bytes), "retry.xlsx"),
    })
    assert replay.status_code == 200
    assert replay.get_json()["replayed"] is True
    assert replay.get_json()["importacion"]["id"] == first_data["id"]
    badkey = client.post(endpoint, headers=tokens[1], data={
        "idempotency_key": key, "file": (BytesIO(workbook(second="NO")), "modified.xlsx"),
    })
    assert badkey.status_code == 409
    second = client.post(endpoint, headers=tokens[1], data={
        "idempotency_key": str(uuid4()), "file": (BytesIO(workbook(second="NO")), "final.xlsx"),
    })
    assert second.status_code == 201, second.get_json()
    assert second.get_json()["importacion"]["version"] == 2
    assert second.get_json()["importacion"]["resultado"] == "VALIDATED"
    summary_list = client.get(url, headers=tokens[1]).get_json()["items"][0]
    assert summary_list["importacion_actual"]["resultado"] == "VALIDATED"
    assert (summary_list["importacion_actual"]["si"], summary_list["importacion_actual"]["no"],
            summary_list["importacion_actual"]["sin_revisar"]) == (1, 1, 0)
    # Archivo estructuralmente incompatible: archivar, diagnosticar y NO promover.
    malformed = load_workbook(BytesIO(workbook(second="NO")))
    del malformed["HALLAZGOS"]
    damaged = BytesIO()
    malformed.save(damaged)
    third = client.post(endpoint, headers=tokens[1], data={
        "idempotency_key": str(uuid4()), "file": (BytesIO(damaged.getvalue()), "faltan_hojas.xlsx"),
    })
    assert third.status_code == 201, third.get_json()
    assert third.get_json()["importacion"]["resultado"] == "INVALID"
    assert third.get_json()["importacion"]["diagnosticos"][0]["codigo"] == "ESTRUCTURA_EXCEL"
    listed = client.get(f"{url}/{lev}", headers=tokens[1])
    assert listed.status_code == 200
    assert len(listed.get_json()["versiones"]) == 3
    assert listed.get_json()["levantamiento"]["ultima_importacion_id"] == second.get_json()["importacion"]["id"]
    assert _snapshot(app) == original
    assert client.get(f"{url}/{lev}", headers=tokens[2]).status_code == 403


def test_inicial_restringido_y_evidencia_persistente(api):
    client, tokens, app = api
    root = "/api/inventario/levantamientos"
    created = client.post(root, headers=tokens[1], data={
        "sucursal_id":"8","origen_corte":"INICIAL","cycle_key":str(uuid4())})
    assert created.status_code == 201, created.get_json()
    lev = created.get_json()["levantamiento"]["id"]
    result = client.post(f"{root}/{lev}/importaciones", headers=tokens[1], data={
        "idempotency_key":str(uuid4()),"file":(BytesIO(workbook(branch=8,initial=True)),"sin_catalogo.xlsx")})
    assert result.status_code == 201, result.get_json()
    assert result.get_json()["importacion"]["resultado"] == "INCOMPLETE"
    assert result.get_json()["importacion"]["resumen"]["requiere_revision_inicial"]
    assert _snapshot(app)["mov"] == 0


def test_api_real_xlsx_http_aislado_con_historico_sin_mutaciones(api):
    """Gate opt-in: bytes EXACTOS del archivo distribuido, API real y PostgreSQL efímero.

    Requiere E4_REAL_XLSX_PATH y E4_REAL_BASELINE_CSV_PATH. Los originales
    permanecen fuera del repo, no se envían a CI, y solo se leen.
    """
    path_xlsx = os.getenv("E4_REAL_XLSX_PATH")
    path_csv = os.getenv("E4_REAL_BASELINE_CSV_PATH")
    if not path_xlsx and not path_csv:
        pytest.skip("Gate real no configurado: requiere rutas locales XLSX y CSV")
    assert path_xlsx and path_csv, "Se requieren ambos archivos históricos locales"
    xlsx = Path(path_xlsx).read_bytes()
    baseline = Path(path_csv).read_bytes()
    expected_sha = os.getenv(
        "E4_REAL_XLSX_SHA256",
        "c3f86c1f2a1aa6ede646451c81d5b35306757fec417dad7e45d89e17c4896dd3",
    )
    assert sha256(xlsx).hexdigest() == expected_sha, (
        "El XLSX no coincide con el original de Pabellón Rosarito 2026-10-28"
    )
    references, excluded = parse_historical_csv(baseline, 7)
    assert len(references) == 150, "El CSV no coincide con el corte distribuido de esta sucursal"
    assert len({r["inventario_id_original"] for r in references}) == len(references)

    client, tokens, app = api
    # La BD efímera conserva los 3 aparatos de ejemplo de CI, pero agregamos
    # el universo REAL únicamente como datos temporales de QA: no hay servidor productivo.
    with app.app_context():
        with db.engine.begin() as con:
            for ref in references:
                inv_id = ref["inventario_id_original"]
                codigo = ref["codigo_interno_original"]
                stock = ref["stock_original"]
                assert con.execute(
                    sa.text("SELECT 1 FROM inventario_general WHERE id=:id"), {"id": inv_id}
                ).scalar() is None, f"Colisión con aparato sintético: {inv_id}"
                con.execute(sa.text(
                    "INSERT INTO inventario_general (id,codigo_interno,tipo,nombre) "
                    "VALUES (:id,:codigo,'APARATOS','Referencia histórica E4 QA')"
                ), {"id": inv_id, "codigo": codigo})
                con.execute(sa.text(
                    "INSERT INTO inventario_sucursal (inventario_id,sucursal_id,stock) "
                    "VALUES (:id,7,:stock)"
                ), {"id": inv_id, "stock": stock})
        db.session.remove()
    before = _snapshot(app)
    root = "/api/inventario/levantamientos"
    created = client.post(root, headers=tokens[1], data={
        "sucursal_id": "7", "origen_corte": "CSV_HISTORICO",
        "cycle_key": str(uuid4()), "fecha_corte_local": "2026-10-28",
        "responsable": "QA E4 binario privado",
        "baseline_csv": (BytesIO(baseline), "inventarios_corte_distribucion.csv"),
    })
    assert created.status_code == 201, created.get_json()
    lev = created.get_json()["levantamiento"]
    assert lev["baseline_resumen"]["referencias_esperadas"] == 150
    assert lev["origen_corte"] == "CSV_HISTORICO"
    assert lev["estado"] == "IN_PROGRESS"
    endpoint = f"{root}/{lev['id']}/importaciones"
    first = client.post(endpoint, headers=tokens[1], data={
        "idempotency_key": str(uuid4()),
        "file": (BytesIO(xlsx), Path(path_xlsx).name),
    })
    assert first.status_code == 201, first.get_json()
    imp = first.get_json()["importacion"]
    summary = imp["resumen"]
    assert imp["resultado"] == "INCOMPLETE", imp
    assert imp["file_sha256"] == expected_sha
    assert summary["referencias_esperadas"] == 150
    assert (summary["si"], summary["no"], summary["sin_revisar"]) == (0, 0, 150)
    assert summary["hallazgos"] == 0
    assert summary["errores"] == 0, imp["diagnosticos"]
    assert summary["avisos"] == 0, imp["diagnosticos"]
    assert len(summary["pendientes"]) == 150
    assert {p["inventario_id"] for p in summary["pendientes"]} == {
        r["inventario_id_original"] for r in references
    }
    detail = client.get(
        f"{endpoint}/{imp['id']}", headers=tokens[1]
    )
    assert detail.status_code == 200, detail.get_json()
    filas = detail.get_json()["filas"]
    assert len(filas) == 150
    assert all(r["hoja"] == "APARATOS" and r["respuesta"] == "SIN_REVISAR" for r in filas)

    downloaded = client.get(
        f"{endpoint}/{imp['id']}/archivo", headers=tokens[1]
    )
    assert downloaded.status_code == 200
    assert downloaded.data == xlsx, "El original privado no coincide byte a byte"
    replay = client.post(endpoint, headers=tokens[1], data={
        "idempotency_key": str(uuid4()),
        "file": (BytesIO(xlsx), "segunda_entrega_igual.xlsx"),
    })
    assert replay.status_code == 200
    assert replay.get_json()["replayed"] is True
    assert replay.get_json()["importacion"]["id"] == imp["id"]
    listed = client.get(f"{root}/{lev['id']}", headers=tokens[1])
    assert listed.status_code == 200
    assert len(listed.get_json()["versiones"]) == 1
    assert _snapshot(app) == before, "La evidencia NO debe cambiar inventario V1/E2"
