from __future__ import annotations

from datetime import datetime, timedelta, timezone
from io import BytesIO
from unittest.mock import patch

import pytest
from openpyxl import load_workbook
from sqlalchemy import BigInteger, Column, Integer, MetaData, Table, create_engine, event
from sqlalchemy.orm import Session

from app.models.marketing import (
    MarketingCampaignV2ORM,
    MarketingCampaignV2ProviderRecipientObservationORM,
    MarketingCampaignV2ProviderStatsSnapshotORM,
    MarketingCampaignV2RecipientEvidenceORM,
    MarketingCampaignV2RecipientORM,
)
from app.services.marketing_campaign_v2_reporting_service import (
    MarketingCampaignV2ReportingValidationError,
    build_campaign_v2_consolidated_report,
    build_campaign_v2_individual_report,
    build_campaign_v2_reporting_export_dataset,
)
from app.services.marketing_campaign_v2_reporting_excel_service import (
    build_campaign_v2_reporting_excel,
)


BASE = datetime(2026, 10, 1, 18, 0, tzinfo=timezone.utc)


def _definition(scope):
    return {
        "filters": {"allowed_sucursal_keys": scope},
        "preview": {
            "fingerprint": "freeze",
            "fingerprint_version": "campaign-v2-freeze-v1",
        },
    }


@pytest.fixture
def session():
    engine = create_engine("sqlite://")
    metadata = MetaData()
    Table("users", metadata, Column("id", Integer, primary_key=True))
    Table(
        "socios_vencidos_cartera",
        metadata,
        Column("id", BigInteger, primary_key=True),
    )
    Table(
        "socios_activos_snapshot_rows",
        metadata,
        Column("id", BigInteger, primary_key=True),
    )
    Table(
        "socios_activos_snapshots",
        metadata,
        Column("id", BigInteger, primary_key=True),
    )
    MarketingCampaignV2ORM.__table__.to_metadata(metadata)
    Table(
        "marketing_campaign_v2_provider_campaigns",
        metadata,
        Column("id", BigInteger, primary_key=True),
    )
    MarketingCampaignV2RecipientORM.__table__.to_metadata(metadata)
    MarketingCampaignV2RecipientEvidenceORM.__table__.to_metadata(metadata)
    MarketingCampaignV2ProviderStatsSnapshotORM.__table__.to_metadata(metadata)
    MarketingCampaignV2ProviderRecipientObservationORM.__table__.to_metadata(
        metadata
    )
    metadata.create_all(engine)

    with Session(engine) as value:
        campaigns = [
            _campaign(
                1,
                "Small",
                purpose="REACTIVATION",
                source="EXPIRED_MEMBERS",
                provider="IVENTAS",
                scope=["BRANCH A"],
                frozen_at=BASE,
            ),
            _campaign(
                2,
                "Large",
                purpose="NEW_SALE",
                source="FUNNEL_PORTFOLIO",
                provider="IVENTAS",
                scope=["BRANCH A"],
                frozen_at=BASE + timedelta(minutes=1),
            ),
            _campaign(
                3,
                "Without snapshot",
                purpose="ACTIVE_MEMBERS",
                source="ACTIVE_MEMBERS",
                provider="OTHER",
                scope=["BRANCH A"],
                frozen_at=BASE + timedelta(minutes=2),
            ),
            _campaign(
                4,
                "Branch B",
                purpose="REACTIVATION",
                source="EXPIRED_MEMBERS",
                provider="IVENTAS",
                scope=["BRANCH B"],
                frozen_at=BASE + timedelta(minutes=3),
            ),
            _campaign(
                5,
                "Global",
                purpose="REACTIVATION",
                source="EXPIRED_MEMBERS",
                provider="IVENTAS",
                scope=None,
                frozen_at=BASE + timedelta(minutes=4),
            ),
            _campaign(
                6,
                "Multi",
                purpose="REACTIVATION",
                source="EXPIRED_MEMBERS",
                provider="IVENTAS",
                scope=["BRANCH A", "BRANCH B"],
                frozen_at=BASE + timedelta(minutes=5),
            ),
        ]
        value.add_all(campaigns)

        recipients = [
            _recipient(101, 1, "6860000001", "BRANCH A", "DOMICILIADO"),
            _recipient(102, 1, "6860000002", "BRANCH A", "DOMICILIADO"),
        ]
        large_phones = [
            "6860000001",
            "6860000003",
            "6860000004",
            "6860000005",
            "6860000006",
            "6860000007",
            "6860000008",
            "6860000009",
        ]
        for offset, phone in enumerate(large_phones):
            branch = "BRANCH A" if offset < 6 else None
            family = "CONVENIO" if offset < 4 else None
            recipients.append(
                _recipient(
                    200 + offset,
                    2,
                    phone,
                    branch,
                    family,
                )
            )
        recipients.extend(
            [
                _recipient(301, 3, "6860000010", "BRANCH A", "TRIMESTRAL"),
                _recipient(302, 3, "6860000011", None, None),
                _recipient(401, 4, "6860000012", "BRANCH B", "DOMICILIADO"),
                _recipient(501, 5, "6860000013", "BRANCH A", "DOMICILIADO"),
                _recipient(601, 6, "6860000014", "BRANCH A", "DOMICILIADO"),
                _recipient(602, 6, "6860000015", "BRANCH B", "DOMICILIADO"),
            ]
        )
        value.add_all(recipients)
        value.flush()

        value.add_all(
            [
                MarketingCampaignV2RecipientEvidenceORM(
                    id=1,
                    recipient_id=207,
                    evidence_order=0,
                    source="FUNNEL_PORTFOLIO",
                    phone_raw="6860000009",
                    phone_mx10="6860000009",
                    sucursal_key="BRANCH A",
                    audience_family="SEMESTRE",
                    evidence_json=[],
                    created_at=BASE,
                ),
                MarketingCampaignV2RecipientEvidenceORM(
                    id=2,
                    recipient_id=302,
                    evidence_order=0,
                    source="ACTIVE_MEMBERS",
                    phone_raw="6860000011",
                    phone_mx10="6860000011",
                    sucursal_key="BRANCH A",
                    audience_family="ESTUDIANTE",
                    evidence_json=[],
                    created_at=BASE,
                ),
            ]
        )

        snapshots = [
            _snapshot(
                10,
                campaign_id=1,
                fetched_at=BASE + timedelta(hours=1),
                fingerprint="1" * 64,
                raw=(2, 0, 0, 1, 1, 1, 1, 2),
                coverage=(2, 0, 0),
                analytics={
                    "responders": 2,
                    "interactions": {"freeText": 1},
                    "cost": {},
                },
            ),
            _snapshot(
                11,
                campaign_id=1,
                fetched_at=BASE + timedelta(hours=2),
                fingerprint="2" * 64,
                raw=(3, 1, 1, 1, 1, 2, 1, 3),
                coverage=(2, 1, 0),
                analytics={
                    "responders": 3,
                    "interactions": {"freeText": 1},
                    "cost": {},
                },
            ),
            _snapshot(
                20,
                campaign_id=2,
                fetched_at=BASE + timedelta(hours=3),
                fingerprint="3" * 64,
                raw=(5, 3, 2, 2, 1, 4, 1, 5),
                coverage=(4, 1, 4),
                analytics={
                    "responders": 5,
                    "interactions": {"freeText": 2},
                    "cost": {},
                },
            ),
            _snapshot(
                21,
                campaign_id=2,
                fetched_at=BASE + timedelta(hours=3),
                fingerprint="4" * 64,
                raw=(6, 4, 3, 2, 1, 5, 1, 6),
                coverage=(4, 1, 4),
                analytics={
                    "responders": 6,
                    "interactions": {"freeText": 2},
                    "cost": {},
                },
            ),
            _snapshot(
                40,
                campaign_id=4,
                fetched_at=BASE + timedelta(hours=4),
                fingerprint="5" * 64,
                raw=(1, 0, 1, 0, 0, 0, 0, 0),
                coverage=(1, 0, 0),
                analytics={"cost": {}},
            ),
            _snapshot(
                50,
                campaign_id=5,
                fetched_at=BASE + timedelta(hours=5),
                fingerprint="6" * 64,
                raw=(1, 0, 1, 0, 0, 0, 0, 0),
                coverage=(1, 0, 0),
                analytics={"cost": {}},
            ),
            _snapshot(
                60,
                campaign_id=6,
                fetched_at=BASE + timedelta(hours=6),
                fingerprint="7" * 64,
                raw=(2, 0, 2, 0, 0, 0, 0, 0),
                coverage=(2, 0, 0),
                analytics={"cost": {}},
            ),
        ]
        value.add_all(snapshots)
        value.flush()

        value.add_all(
            [
                _obs(11, "mx10:6860000001", 101, "SUCCESSFUL", "DELIVERED", ["CTA"]),
                _obs(11, "mx10:6860000002", 102, "SUCCESSFUL", "VIEWED"),
                _obs(11, "mx10:6860999999", None, "FAILED", None),
                _obs(21, "mx10:6860000001", 200, "SUCCESSFUL", "VIEWED", ["CTA"]),
                _obs(21, "mx10:6860000003", 201, "SUCCESSFUL", "SENT"),
                _obs(21, "mx10:6860000004", 202, "FAILED", None),
                _obs(21, "mx10:6860000005", 203, "FAILED", None),
                _obs(21, "mx10:6860999998", None, "FAILED", None),
                _obs(40, "mx10:6860000012", 401, "SUCCESSFUL", "SENT"),
                _obs(50, "mx10:6860000013", 501, "SUCCESSFUL", "SENT"),
                _obs(60, "mx10:6860000014", 601, "SUCCESSFUL", "SENT"),
                _obs(60, "mx10:6860000015", 602, "SUCCESSFUL", "SENT"),
            ]
        )
        value.commit()
        yield value

    engine.dispose()


def _campaign(
    campaign_id,
    name,
    *,
    purpose,
    source,
    provider,
    scope,
    frozen_at,
):
    return MarketingCampaignV2ORM(
        id=campaign_id,
        name=name,
        purpose=purpose,
        source=source,
        provider=provider,
        provider_campaign_id=f"external-{campaign_id}",
        audience_definition_json=_definition(scope),
        frozen_at=frozen_at,
        created_at=frozen_at,
        updated_at=frozen_at,
    )


def _recipient(recipient_id, campaign_id, phone, branch, family):
    return MarketingCampaignV2RecipientORM(
        id=recipient_id,
        campaign_id=campaign_id,
        phone_mx10=phone,
        source="EXPIRED_MEMBERS",
        sucursal=branch,
        audience_family=family,
        created_at=BASE,
    )


def _snapshot(
    snapshot_id,
    *,
    campaign_id,
    fetched_at,
    fingerprint,
    raw,
    coverage,
    analytics,
):
    (
        successful,
        failed,
        sent,
        delivered,
        viewed,
        answered,
        interaction_groups,
        interaction_items,
    ) = raw
    matched, unmatched, without_status = coverage
    return MarketingCampaignV2ProviderStatsSnapshotORM(
        id=snapshot_id,
        campaign_v2_id=campaign_id,
        provider="IVENTAS",
        provider_campaign_id=f"external-{campaign_id}",
        analytics_status="ok",
        fetched_at=fetched_at,
        raw_successful=successful,
        raw_failed=failed,
        raw_sent=sent,
        raw_delivered=delivered,
        raw_viewed=viewed,
        raw_answered=answered,
        raw_interaction_groups=interaction_groups,
        raw_interaction_items=interaction_items,
        analytics_json=analytics,
        button_interactions_json=[],
        provider_recipient_count=matched + unmatched,
        matched_recipient_count=matched,
        unmatched_provider_count=unmatched,
        frozen_recipient_without_provider_status_count=without_status,
        fingerprint=fingerprint,
        created_at=fetched_at,
    )


def _obs(
    snapshot_id,
    phone,
    recipient_id,
    outcome,
    delivery_bucket,
    labels=None,
):
    return MarketingCampaignV2ProviderRecipientObservationORM(
        snapshot_id=snapshot_id,
        normalized_phone=phone,
        campaign_recipient_id=recipient_id,
        outcome=outcome,
        delivery_bucket=delivery_bucket,
        button_labels_json=labels or [],
        created_at=BASE,
    )


def _report(session, *, scope=("BRANCH A",), **filters):
    return build_campaign_v2_consolidated_report(
        allowed_sucursal_keys=scope,
        filters=filters,
        session=session,
    )


def test_bulk_latest_and_ordering_with_tie_by_id(session):
    report = _report(session)

    assert [row["campaign"]["id"] for row in report["campaigns"]] == [2, 1, 3]
    assert report["campaigns"][0]["observation"]["snapshot_id"] == 21
    assert report["campaigns"][1]["observation"]["snapshot_id"] == 11
    assert report["campaigns"][2]["observation"]["snapshot_id"] is None
    assert all("evolution" not in row for row in report["campaigns"])
    assert all("dimensions" not in row for row in report["campaigns"])


def test_recipient_exposures_do_not_dedupe_cross_campaign(session):
    report = _report(session)

    assert report["summary"]["campaign_count"] == 3
    assert report["summary"]["total_recipients"] == 12
    assert report["campaigns"][0]["audience"]["total_recipients"] == 8
    assert report["campaigns"][1]["audience"]["total_recipients"] == 2
    # 6860000001 existe en Campaign 1 y 2 y aporta dos exposures.


def test_weighted_rates_are_recomputed_from_summed_numerators(session):
    report = _report(session)
    summary = report["summary"]

    assert summary["normalized"] == {
        "successful": 4,
        "failed": 2,
        "sent": 1,
        "delivered": 1,
        "viewed": 2,
        "reach_count": 3,
    }
    assert summary["rates"] == {
        "successful_rate": 4 / 12,
        "reach_rate": 3 / 12,
        "read_rate": 2 / 3,
        "failure_rate": 2 / 12,
    }

    campaign_rates = [
        row["rates"]["successful_rate"]
        for row in report["campaigns"]
        if row["rates"]["successful_rate"] is not None
    ]
    assert summary["rates"]["successful_rate"] != (
        sum(campaign_rates) / len(campaign_rates)
    )


def test_campaign_without_snapshot_contributes_denominator_and_coverage(session):
    report = _report(session)
    summary = report["summary"]

    assert summary["campaigns_with_snapshot"] == 2
    assert summary["campaigns_without_snapshot"] == 1
    assert summary["coverage"] == {
        "matched_recipient_count": 6,
        "unmatched_provider_count": 2,
        "frozen_recipient_without_provider_status_count": 6,
        "status_coverage_rate": 0.5,
    }
    without = next(
        row
        for row in report["campaigns"]
        if row["campaign"]["id"] == 3
    )
    assert without["provider_raw"] is None
    assert without["normalized"]["successful"] == 0
    assert without["coverage"]["frozen_recipient_without_provider_status_count"] == 2


def test_provider_raw_interactions_and_cost_summary_are_separate(session):
    report = _report(session)
    summary = report["summary"]

    assert summary["provider_raw"] == {
        "successful": 9,
        "failed": 5,
        "sent": 4,
        "delivered": 3,
        "viewed": 2,
        "answered": 7,
        "interaction_groups": 2,
        "interaction_items": 9,
    }
    assert summary["normalized"]["successful"] == 4
    assert summary["interactions"] == {
        "button_interaction_recipient_exposures": 2,
        "responders_aggregate": 9,
        "responders_campaigns_with_value": 2,
        "free_text_aggregate": 3,
        "free_text_campaigns_with_value": 2,
    }
    assert summary["cost"] == {
        "status": "unavailable",
        "currency": None,
        "total": None,
    }


def test_branch_and_family_breakdowns_partition_frozen_cohort(session):
    report = _report(session)
    branches = report["breakdowns"]["branches"]
    families = report["breakdowns"]["audience_families"]

    assert sum(row["recipient_exposures"] for row in branches) == 12
    assert sum(row["successful"] for row in branches) == 4
    assert sum(row["failed"] for row in branches) == 2
    assert sum(row["viewed"] for row in branches) == 2
    assert sum(row["reach_count"] for row in branches) == 3

    assert sum(row["recipient_exposures"] for row in families) == 12
    assert sum(row["successful"] for row in families) == 4
    assert sum(row["failed"] for row in families) == 2

    branch_unknown = next(row for row in branches if row["value"] == "UNKNOWN")
    family_unknown = next(row for row in families if row["value"] == "UNKNOWN")
    assert branch_unknown["recipient_exposures"] == 1
    assert family_unknown["recipient_exposures"] == 3


def test_breakdowns_do_not_contain_campaign_level_provider_aggregates(session):
    report = _report(session)
    for group in (
        report["breakdowns"]["branches"],
        report["breakdowns"]["audience_families"],
    ):
        for row in group:
            assert "provider_raw" not in row
            assert "responders_aggregate" not in row
            assert "free_text_aggregate" not in row
            assert "cost" not in row
            assert "unmatched_provider_count" not in row["coverage"]


def test_observed_window_filters_latest_global_snapshot_inclusively(session):
    exact_one = (BASE + timedelta(hours=2)).isoformat()
    exact_two = (BASE + timedelta(hours=3)).isoformat()

    at_from = _report(session, observed_from=exact_two)
    at_to = _report(session, observed_to=exact_one)
    exact = _report(
        session,
        observed_from=exact_one,
        observed_to=exact_two,
    )

    assert [row["campaign"]["id"] for row in at_from["campaigns"]] == [2]
    assert [row["campaign"]["id"] for row in at_to["campaigns"]] == [1]
    assert [row["campaign"]["id"] for row in exact["campaigns"]] == [2, 1]
    assert all(row["campaign"]["id"] != 3 for row in exact["campaigns"])


@pytest.mark.parametrize(
    ("filters", "expected"),
    [
        ({"purpose": "NEW_SALE"}, [2]),
        ({"source": "FUNNEL_PORTFOLIO"}, [2]),
        ({"provider": "OTHER"}, [3]),
        ({"snapshot_status": "WITH_SNAPSHOT"}, [2, 1]),
        ({"snapshot_status": "WITHOUT_SNAPSHOT"}, [3]),
    ],
)
def test_v1_filters(session, filters, expected):
    report = _report(session, **filters)
    assert [row["campaign"]["id"] for row in report["campaigns"]] == expected


@pytest.mark.parametrize(
    "filters",
    [
        {"observed_from": "2026-10-01T12:00:00"},
        {"observed_to": "not-a-date"},
        {
            "observed_from": "2026-10-02T00:00:00Z",
            "observed_to": "2026-10-01T00:00:00Z",
        },
        {"purpose": "BAD"},
        {"source": "BAD"},
        {"provider": "X" * 51},
        {"snapshot_status": "BAD"},
        {"branch": "BRANCH A"},
    ],
)
def test_invalid_filters_are_rejected(session, filters):
    with pytest.raises(MarketingCampaignV2ReportingValidationError):
        _report(session, **filters)


def test_scope_visibility_is_applied_before_aggregates(session):
    limited = _report(session, scope=("BRANCH A",))
    global_scope = _report(session, scope=None)
    multi_scope = _report(session, scope=("BRANCH A", "BRANCH B"))

    assert [row["campaign"]["id"] for row in limited["campaigns"]] == [2, 1, 3]
    assert {row["campaign"]["id"] for row in global_scope["campaigns"]} == {
        1,
        2,
        3,
        4,
        5,
        6,
    }
    assert {row["campaign"]["id"] for row in multi_scope["campaigns"]} == {
        1,
        2,
        3,
        4,
        6,
    }
    assert 5 not in {row["campaign"]["id"] for row in multi_scope["campaigns"]}


def test_same_persisted_state_and_filters_produce_same_response(session):
    first = _report(session, purpose="REACTIVATION")
    second = _report(session, purpose="REACTIVATION")
    assert first == second
    assert "generated_at" not in first


def test_provider_boundary_and_m13_are_not_used(session):
    with (
        patch(
            "app.services.marketing_campaign_v2_provider_stats_service."
            "get_campaign_v2_provider_stats",
            side_effect=AssertionError("no provider read-through"),
        ),
        patch(
            "app.services.marketing_campaign_v2_provider_stats_snapshot_service."
            "capture_campaign_v2_provider_stats_snapshot",
            side_effect=AssertionError("no provider capture"),
        ),
        patch(
            "app.services.marketing_campaign_v2_provider_history_service."
            "get_provider_history_for_phones",
            side_effect=AssertionError("no M13"),
        ),
    ):
        report = _report(session)

    assert report["summary"]["campaign_count"] == 3


def _statement_count(session, scope):
    statements = []

    def before_cursor_execute(
        _conn,
        _cursor,
        statement,
        _parameters,
        _context,
        _executemany,
    ):
        statements.append(statement)

    event.listen(
        session.get_bind(),
        "before_cursor_execute",
        before_cursor_execute,
    )
    try:
        report = _report(session, scope=scope)
    finally:
        event.remove(
            session.get_bind(),
            "before_cursor_execute",
            before_cursor_execute,
        )
    return len(statements), report


def test_query_count_is_approximately_constant_not_per_campaign(session):
    one_count, one = _statement_count(session, ("BRANCH B",))
    many_count, many = _statement_count(session, None)

    assert one["summary"]["campaign_count"] == 1
    assert many["summary"]["campaign_count"] == 6
    assert many_count <= one_count + 2
    assert many_count <= 6


def test_export_dataset_reuses_m28_selection_and_loads_historical_evolution_bulk(session):
    consolidated = _report(session, purpose="REACTIVATION")
    dataset = build_campaign_v2_reporting_export_dataset(
        allowed_sucursal_keys=("BRANCH A",),
        filters={"purpose": "REACTIVATION"},
        session=session,
    )

    assert dataset["report_type"] == "CONSOLIDATED"
    assert dataset["report"] == consolidated
    assert [
        row["campaign"]["id"]
        for row in dataset["report"]["campaigns"]
    ] == [1]
    assert [
        (
            row["campaign_id"],
            row["snapshot_id"],
            row["observed_at"],
        )
        for row in dataset["evolution"]
    ] == [
        (1, 10, (BASE + timedelta(hours=1)).isoformat()),
        (1, 11, (BASE + timedelta(hours=2)).isoformat()),
    ]
    assert dataset["evolution"][0]["normalized"] == {
        "successful": 0,
        "failed": 0,
        "sent": 0,
        "delivered": 0,
        "viewed": 0,
        "reach_count": 0,
    }
    assert dataset["evolution"][1]["normalized"] == {
        "successful": 2,
        "failed": 0,
        "sent": 0,
        "delivered": 1,
        "viewed": 1,
        "reach_count": 2,
    }
    assert dataset["scope"] == {
        "is_global": False,
        "allowed_sucursal_keys": ["BRANCH A"],
    }


def test_export_dataset_multi_campaign_evolution_order_is_deterministic(session):
    dataset = build_campaign_v2_reporting_export_dataset(
        allowed_sucursal_keys=None,
        filters=None,
        session=session,
    )

    evolution_keys = [
        (row["campaign_id"], row["observed_at"], row["snapshot_id"])
        for row in dataset["evolution"]
    ]
    assert evolution_keys == sorted(evolution_keys)
    assert {row["campaign_id"] for row in dataset["evolution"]} == {
        1,
        2,
        4,
        5,
        6,
    }


def test_export_dataset_empty_consolidated_is_valid(session):
    dataset = build_campaign_v2_reporting_export_dataset(
        allowed_sucursal_keys=("BRANCH A",),
        filters={"purpose": "UNCLASSIFIED"},
        session=session,
    )

    assert dataset["report"]["summary"]["campaign_count"] == 0
    assert dataset["report"]["summary"]["total_recipients"] == 0
    assert dataset["report"]["campaigns"] == []
    assert dataset["evolution"] == []


def _export_statement_count(session, scope):
    statements = []

    def before_cursor_execute(
        _conn,
        _cursor,
        statement,
        _parameters,
        _context,
        _executemany,
    ):
        statements.append(statement)

    event.listen(
        session.get_bind(),
        "before_cursor_execute",
        before_cursor_execute,
    )
    try:
        dataset = build_campaign_v2_reporting_export_dataset(
            allowed_sucursal_keys=scope,
            filters=None,
            session=session,
        )
    finally:
        event.remove(
            session.get_bind(),
            "before_cursor_execute",
            before_cursor_execute,
        )
    return len(statements), dataset


def test_export_dataset_query_count_is_bulk_not_per_campaign(session):
    one_count, one = _export_statement_count(session, ("BRANCH B",))
    many_count, many = _export_statement_count(session, None)

    assert one["report"]["summary"]["campaign_count"] == 1
    assert many["report"]["summary"]["campaign_count"] == 6
    assert one_count <= 7
    assert many_count <= 7
    assert many_count <= one_count + 1


def test_export_dataset_provider_boundary_and_m13_are_not_used(session):
    with (
        patch(
            "app.services.marketing_campaign_v2_provider_stats_service."
            "get_campaign_v2_provider_stats",
            side_effect=AssertionError("no provider read-through"),
        ),
        patch(
            "app.services.marketing_campaign_v2_provider_stats_snapshot_service."
            "capture_campaign_v2_provider_stats_snapshot",
            side_effect=AssertionError("no provider capture"),
        ),
        patch(
            "app.services.marketing_campaign_v2_provider_history_service."
            "get_provider_history_for_phones",
            side_effect=AssertionError("no M13"),
        ),
    ):
        dataset = build_campaign_v2_reporting_export_dataset(
            allowed_sucursal_keys=("BRANCH A",),
            filters=None,
            session=session,
        )

    assert dataset["report"]["summary"]["campaign_count"] == 3


def test_json_and_xlsx_use_the_same_selected_campaigns_and_summary(session):
    filters = {"purpose": "REACTIVATION"}
    json_report = build_campaign_v2_consolidated_report(
        allowed_sucursal_keys=("BRANCH A",),
        filters=filters,
        session=session,
    )
    dataset = build_campaign_v2_reporting_export_dataset(
        allowed_sucursal_keys=("BRANCH A",),
        filters=filters,
        session=session,
    )
    output, _ = build_campaign_v2_reporting_excel(
        report_type=dataset["report_type"],
        report=dataset["report"],
        evolution=dataset["evolution"],
        scope=dataset["scope"],
        generated_at=BASE,
    )
    workbook = load_workbook(BytesIO(output.getvalue()), data_only=False)

    campaign_sheet = workbook["Campañas"]
    campaign_id_col = next(
        cell.column
        for cell in campaign_sheet[1]
        if cell.value == "campaign_id"
    )
    xlsx_campaign_ids = [
        campaign_sheet.cell(row, campaign_id_col).value
        for row in range(2, campaign_sheet.max_row + 1)
    ]
    json_campaign_ids = [
        row["campaign"]["id"]
        for row in json_report["campaigns"]
    ]
    assert xlsx_campaign_ids == json_campaign_ids

    summary_sheet = workbook["Resumen"]
    summary_values = {
        summary_sheet.cell(row, 1).value: summary_sheet.cell(row, 2).value
        for row in range(2, summary_sheet.max_row + 1)
    }
    assert summary_values["campaign_count"] == json_report["summary"]["campaign_count"]
    assert (
        summary_values["Exposiciones de destinatarios"]
        == json_report["summary"]["total_recipients"]
    )
    assert (
        summary_values["successful"]
        == json_report["summary"]["normalized"]["successful"]
    )
    assert (
        summary_values["successful_rate"]
        == json_report["summary"]["rates"]["successful_rate"]
    )
    assert (
        summary_values["status_coverage_rate"]
        == json_report["summary"]["coverage"]["status_coverage_rate"]
    )


def test_phase2c_reporting_paths_are_read_only_and_offline(session):
    tracked_models = (
        MarketingCampaignV2ORM,
        MarketingCampaignV2RecipientORM,
        MarketingCampaignV2RecipientEvidenceORM,
        MarketingCampaignV2ProviderStatsSnapshotORM,
        MarketingCampaignV2ProviderRecipientObservationORM,
    )
    before = {
        model.__name__: session.query(model).count()
        for model in tracked_models
    }

    with (
        patch(
            "app.services.marketing_campaign_v2_provider_stats_service."
            "get_campaign_v2_provider_stats",
            side_effect=AssertionError("2C must stay offline"),
        ),
        patch(
            "app.services.marketing_campaign_v2_provider_stats_snapshot_service."
            "capture_campaign_v2_provider_stats_snapshot",
            side_effect=AssertionError("2C must not capture"),
        ),
        patch(
            "app.services.marketing_campaign_v2_provider_history_service."
            "get_provider_history_for_phones",
            side_effect=AssertionError("2C must not use M13"),
        ),
    ):
        individual = build_campaign_v2_individual_report(
            campaign_id=1,
            allowed_sucursal_keys=("BRANCH A",),
            session=session,
        )
        consolidated = build_campaign_v2_consolidated_report(
            allowed_sucursal_keys=("BRANCH A",),
            filters={"purpose": "REACTIVATION"},
            session=session,
        )
        dataset = build_campaign_v2_reporting_export_dataset(
            allowed_sucursal_keys=("BRANCH A",),
            filters={"purpose": "REACTIVATION"},
            session=session,
        )
        consolidated_output, _ = build_campaign_v2_reporting_excel(
            report_type=dataset["report_type"],
            report=dataset["report"],
            evolution=dataset["evolution"],
            scope=dataset["scope"],
            generated_at=BASE,
        )
        individual_evolution = [
            {
                "campaign_id": individual["campaign"]["id"],
                "campaign_name": individual["campaign"]["name"],
                **row,
            }
            for row in individual["evolution"]
        ]
        individual_output, _ = build_campaign_v2_reporting_excel(
            report_type="INDIVIDUAL",
            report=individual,
            evolution=individual_evolution,
            scope=individual["dimensions"]["scope"],
            generated_at=BASE,
        )

    assert consolidated["summary"]["campaign_count"] == 1
    assert load_workbook(
        BytesIO(consolidated_output.getvalue()),
        data_only=False,
    ).sheetnames == ["Resumen", "Campañas", "KPIs", "Evolución", "Metadata"]
    assert load_workbook(
        BytesIO(individual_output.getvalue()),
        data_only=False,
    ).sheetnames == ["Resumen", "Campañas", "KPIs", "Evolución", "Metadata"]

    after = {
        model.__name__: session.query(model).count()
        for model in tracked_models
    }
    assert after == before
    assert not session.new
    assert not session.dirty
    assert not session.deleted


def test_individual_json_and_xlsx_have_phase2c_parity(session):
    report = build_campaign_v2_individual_report(
        campaign_id=1,
        allowed_sucursal_keys=("BRANCH A",),
        session=session,
    )
    evolution = [
        {
            "campaign_id": report["campaign"]["id"],
            "campaign_name": report["campaign"]["name"],
            **row,
        }
        for row in report["evolution"]
    ]
    output, _ = build_campaign_v2_reporting_excel(
        report_type="INDIVIDUAL",
        report=report,
        evolution=evolution,
        scope=report["dimensions"]["scope"],
        generated_at=BASE,
    )
    workbook = load_workbook(BytesIO(output.getvalue()), data_only=False)

    campaigns = workbook["Campañas"]
    headers = {
        campaigns.cell(1, col).value: col
        for col in range(1, campaigns.max_column + 1)
    }
    row = {
        key: campaigns.cell(2, column).value
        for key, column in headers.items()
    }

    assert row["campaign_id"] == report["campaign"]["id"]
    assert row["name"] == report["campaign"]["name"]
    assert row["purpose"] == report["campaign"]["purpose"]
    assert row["source"] == report["campaign"]["source"]
    assert row["provider"] == report["campaign"]["provider"]
    assert row["provider_campaign_id"] == report["campaign"]["provider_campaign_id"]
    assert row["snapshot_id"] == report["observation"]["snapshot_id"]
    assert row["latest_observed_at"] == report["observation"]["latest_observed_at"]
    assert row["analytics_status"] == report["observation"]["analytics_status"]
    assert row["total_recipients"] == report["audience"]["total_recipients"]

    for metric in (
        "successful",
        "failed",
        "sent",
        "delivered",
        "viewed",
        "reach_count",
    ):
        assert row[metric] == report["normalized"][metric]
    for metric in (
        "successful_rate",
        "reach_rate",
        "read_rate",
        "failure_rate",
    ):
        assert row[metric] == report["rates"][metric]
    assert (
        row["matched_recipient_count"]
        == report["coverage"]["matched_recipient_count"]
    )
    assert (
        row["unmatched_provider_count"]
        == report["coverage"]["unmatched_provider_count"]
    )
    assert (
        row["frozen_without_status"]
        == report["coverage"]["frozen_recipient_without_provider_status_count"]
    )
    assert (
        row["status_coverage_rate"]
        == report["coverage"]["status_coverage_rate"]
    )
    assert row["raw_successful"] == report["provider_raw"]["successful"]
    assert row["raw_failed"] == report["provider_raw"]["failed"]
    assert (
        row["button_interaction_recipient_exposures"]
        == report["interactions"]["unique_button_recipients"]
    )
    assert (
        row["responders_aggregate"]
        == report["interactions"]["responders_aggregate"]
    )
    assert (
        row["free_text_aggregate"]
        == report["interactions"]["free_text_aggregate"]
    )
    assert row["cost_status"] == report["cost"]["status"]
    assert row["currency"] is report["cost"]["currency"]
    assert row["cost_total"] is report["cost"]["total"]

    evolution_sheet = workbook["Evolución"]
    evolution_headers = {
        evolution_sheet.cell(1, col).value: col
        for col in range(1, evolution_sheet.max_column + 1)
    }
    workbook_evolution = [
        (
            evolution_sheet.cell(r, evolution_headers["snapshot_id"]).value,
            evolution_sheet.cell(r, evolution_headers["observed_at"]).value,
            evolution_sheet.cell(
                r,
                evolution_headers["normalized_successful"],
            ).value,
        )
        for r in range(2, evolution_sheet.max_row + 1)
    ]
    assert workbook_evolution == [
        (
            point["snapshot_id"],
            point["observed_at"],
            point["normalized"]["successful"],
        )
        for point in report["evolution"]
    ]
