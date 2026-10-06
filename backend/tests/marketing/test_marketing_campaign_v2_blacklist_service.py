from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from types import SimpleNamespace as NS

from openpyxl import Workbook, load_workbook
import pytest

from app.models.marketing import MarketingCampaignV2BlacklistORM
from app.services import marketing_campaign_v2_blacklist_service as service


def _xlsx(values, *, header="telefono"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([header])
    for value in values:
        sheet.append([value])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class _Query:
    def __init__(self, rows):
        self.rows = list(rows)

    def filter(self, *_args):
        return self

    def order_by(self, *_args):
        return self

    def all(self):
        return list(self.rows)


class _Session:
    def __init__(self, rows=()):
        self.rows = list(rows)
        self.added = []
        self.commits = 0
        self.rollbacks = 0

    def query(self, *_args):
        return _Query(self.rows)

    def add(self, value):
        self.added.append(value)

    def commit(self):
        self.commits += 1

    def rollback(self):
        self.rollbacks += 1


def test_parse_blacklist_xlsx_normalizes_deduplicates_and_reports_invalid():
    parsed = service._parse_blacklist_workbook(
        _xlsx(
            [
                6861000001,
                "52 686 100 0002",
                "+52 1 686 100 0003",
                "6861000001",
                "no-es-telefono",
                None,
            ],
            header="Teléfono",
        )
    )

    assert parsed["phones"] == {
        "6861000001",
        "6861000002",
        "6861000003",
    }
    assert parsed["rows_read"] == 5
    assert parsed["duplicates_in_file"] == 1
    assert parsed["invalid"] == 1


def test_parse_blacklist_requires_supported_phone_header():
    with pytest.raises(
        service.MarketingCampaignV2BlacklistValidationError,
        match="columna telefono",
    ):
        service._parse_blacklist_workbook(_xlsx(["6861000001"], header="correo"))


def test_import_blacklist_is_incremental(monkeypatch):
    session = _Session()
    monkeypatch.setattr(
        service,
        "get_blacklisted_phones",
        lambda **_kwargs: {"6861000002"},
    )
    monkeypatch.setattr(
        service,
        "get_campaign_v2_blacklist_summary",
        lambda **_kwargs: {
            "total": 2,
            "latest_created_at": "2026-10-06T10:00:00+00:00",
        },
    )

    result = service.import_campaign_v2_blacklist_xlsx(
        file_bytes=_xlsx(
            [
                "6861000001",
                "6861000002",
                "6861000002",
                "invalido",
            ]
        ),
        filename="bloqueados.xlsx",
        created_by_user_id=7,
        session=session,
    )

    assert result == {
        "filename": "bloqueados.xlsx",
        "rows_read": 4,
        "valid_unique": 2,
        "added": 1,
        "already_existing": 1,
        "duplicates_in_file": 1,
        "invalid": 1,
        "blacklist_total": 2,
        "latest_created_at": "2026-10-06T10:00:00+00:00",
    }
    assert session.commits == 1
    assert session.rollbacks == 0
    assert len(session.added) == 1
    added = session.added[0]
    assert isinstance(added, MarketingCampaignV2BlacklistORM)
    assert added.phone_mx10 == "6861000001"
    assert added.source_filename == "bloqueados.xlsx"
    assert added.created_by_user_id == 7


def test_get_blacklisted_phones_returns_only_existing_matches():
    session = _Session(
        [
            ("6861000002",),
            ("6861000003",),
        ]
    )

    result = service.get_blacklisted_phones(
        phones=[
            "6861000001",
            "+52 686 100 0002",
            "6861000003",
        ],
        session=session,
    )

    assert result == {"6861000002", "6861000003"}


def test_export_blacklist_contains_full_incremental_list():
    rows = [
        NS(
            id=1,
            phone_mx10="6861000001",
            created_at=datetime(2026, 10, 5, 12, 0, tzinfo=timezone.utc),
            source_filename="primero.xlsx",
            created_by_user_id=7,
        ),
        NS(
            id=2,
            phone_mx10="6861000002",
            created_at=datetime(2026, 10, 6, 12, 0, tzinfo=timezone.utc),
            source_filename="segundo.xlsx",
            created_by_user_id=9,
        ),
    ]

    data, filename = service.export_campaign_v2_blacklist_xlsx(
        session=_Session(rows)
    )

    assert filename == "campaign_v2_lista_negra.xlsx"
    assert list(load_workbook(BytesIO(data)).active.values) == [
        ("telefono", "fecha_alta", "archivo_origen", "usuario_id"),
        (
            "6861000001",
            "2026-10-05T12:00:00+00:00",
            "primero.xlsx",
            7,
        ),
        (
            "6861000002",
            "2026-10-06T12:00:00+00:00",
            "segundo.xlsx",
            9,
        ),
    ]
