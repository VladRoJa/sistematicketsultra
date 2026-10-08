"""Ephemeral M3 HTTP + PostgreSQL integration smoke, never production.

Requires the dedicated CI database M3_TEST_POSTGRES_URL. Uses a real local
HTTP listener and real persisted reporting, with no outbound provider calls.
Authorization identities are simulated, but JWT is verified by Flask.
"""

from __future__ import annotations

from datetime import datetime, timedelta, timezone
from io import BytesIO
import os
import threading
from types import SimpleNamespace
from unittest.mock import patch

import pytest
import requests
import sqlalchemy as sa
from flask import Flask
from flask_jwt_extended import JWTManager, create_access_token
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from werkzeug.serving import make_server

from app.extensions import db
from app.models.marketing import (
    MarketingCampaignV2ORM as Campaign,
    MarketingCampaignV2RecipientORM as Recipient,
    MarketingCampaignV2RecipientEvidenceORM as Evidence,
    MarketingCampaignV2ProviderStatsSnapshotORM as Snapshot,
    MarketingCampaignV2ProviderRecipientObservationORM as Observation,
)
from app.routes import marketing_campaign_v2_routes as routes


URL = os.getenv("M3_TEST_POSTGRES_URL")
UTC = datetime(2026, 10, 8, 18, tzinfo=timezone.utc)
PREFIX = "/api/marketing/campaigns-v2"


def _guard_isolated_url(url: str) -> None:
    parsed = sa.engine.make_url(url)
    if (
        parsed.get_backend_name() != "postgresql"
        or parsed.database != "m3_child_snapshots"
        or parsed.host not in {"127.0.0.1", "localhost"}
        or parsed.username != "m3_test"
    ):
        pytest.fail("M3 HTTP QA refuses to touch a non-isolated database")


def _setup_schema(engine: sa.Engine) -> sa.MetaData:
    # Mirror only the required ORM schema, not the entire legacy Suite schema.
    metadata = sa.MetaData()
    sa.Table("users", metadata, sa.Column("id", sa.Integer, primary_key=True))
    for name in (
        "socios_vencidos_cartera",
        "socios_activos_snapshot_rows",
        "socios_activos_snapshots",
    ):
        sa.Table(name, metadata, sa.Column("id", sa.BigInteger, primary_key=True))
    Campaign.__table__.to_metadata(metadata)
    child_table = sa.Table(
        "marketing_campaign_v2_provider_campaigns",
        metadata,
        sa.Column("id", sa.BigInteger, primary_key=True),
        sa.Column("campaign_v2_id", sa.BigInteger, nullable=False),
        sa.Column("provider", sa.String, nullable=False),
        sa.Column("provider_campaign_id", sa.String),
        sa.Column("status", sa.String, nullable=False),
        sa.Column("recipient_count", sa.Integer, nullable=False),
        sa.Column("sucursal_canon", sa.String, nullable=False),
    )
    Recipient.__table__.to_metadata(metadata)
    Evidence.__table__.to_metadata(metadata)
    Snapshot.__table__.to_metadata(metadata)
    Observation.__table__.to_metadata(metadata)
    metadata.create_all(engine)
    with Session(engine) as session:
        session.add(Campaign(
            id=91, name="M3 QA CI - artificial two-branch cohort",
            source="EXPIRED_MEMBERS", purpose="REACTIVATION",
            provider=None, provider_campaign_id=None,
            audience_definition_json={
                "filters": {"allowed_sucursal_keys": ["BRANCH A", "BRANCH B"]},
                "preview": {"fingerprint": "test-only", "fingerprint_version": "campaign-v2-freeze-v1"},
            },
            created_at=UTC, updated_at=UTC, frozen_at=UTC,
        ))
        session.flush()
        session.add_all([
            Recipient(id=701, campaign_id=91, phone_mx10="6860000701",
                      source="EXPIRED_MEMBERS", sucursal="BRANCH A",
                      audience_family="DOMICILIADO", created_at=UTC),
            Recipient(id=702, campaign_id=91, phone_mx10="6860000702",
                      source="EXPIRED_MEMBERS", sucursal="BRANCH B",
                      audience_family="DOMICILIADO", created_at=UTC),
        ])
        session.flush()
        session.execute(child_table.insert(), [
            dict(id=901, campaign_v2_id=91, provider="IVENTAS",
                 provider_campaign_id="ci-fake-batch-a",
                 status="SUBMITTED", recipient_count=1, sucursal_canon="BRANCH A"),
            dict(id=902, campaign_v2_id=91, provider="IVENTAS",
                 provider_campaign_id="ci-fake-batch-b",
                 status="SUBMITTED", recipient_count=1, sucursal_canon="BRANCH B"),
        ])
        session.add(Snapshot(
            id=801, campaign_v2_id=91, provider_campaign_child_id=901,
            provider="IVENTAS", provider_campaign_id="ci-fake-batch-a",
            analytics_status="ok", fetched_at=UTC,
            raw_successful=1, raw_failed=0, raw_sent=0, raw_delivered=1,
            raw_viewed=0, raw_answered=0, raw_interaction_groups=0,
            raw_interaction_items=0, analytics_json=None,
            button_interactions_json=[], provider_recipient_count=1,
            matched_recipient_count=1, unmatched_provider_count=0,
            frozen_recipient_without_provider_status_count=0,
            fingerprint="a" * 64, created_at=UTC,
        ))
        session.flush()
        session.add(Observation(
            snapshot_id=801, normalized_phone="mx10:6860000701",
            campaign_recipient_id=701, outcome="SUCCESSFUL",
            delivery_bucket="DELIVERED", button_labels_json=[],
            created_at=UTC,
        ))
        session.commit()
    return metadata


@pytest.mark.skipif(not URL, reason="Dedicated PostgreSQL test DB is required")
def test_real_http_routes_use_postgres_and_enforce_kill_switch():
    _guard_isolated_url(URL)
    engine = sa.create_engine(URL, pool_pre_ping=True)
    metadata = _setup_schema(engine)
    app = Flask(__name__)
    app.config.update(
        TESTING=False,
        JWT_SECRET_KEY="isolated-http-ci-secret",
        SQLALCHEMY_DATABASE_URI=URL,
        SQLALCHEMY_TRACK_MODIFICATIONS=False,
        CAMPAIGN_V2_PROVIDER_SEND_ENABLED=False,
    )
    db.init_app(app)
    JWTManager(app)
    app.register_blueprint(
        routes.marketing_campaign_v2_bp, url_prefix="/api/marketing"
    )
    with app.app_context():
        token = create_access_token(identity="7")
    authorization = {"Authorization": f"Bearer {token}"}
    access = SimpleNamespace(
        is_global=True, branch_ids=(), can_manage_campaigns=True,
        can_edit_inputs=True, can_preflight_campaigns=True,
        can_manage_dispatch_config=True, can_send_campaigns=True,
    )
    current = {"access": access, "scope": None}
    server = make_server("127.0.0.1", 0, app, threaded=True)
    port = server.server_port
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    base = f"http://127.0.0.1:{port}"

    def api(method, path, **kwargs):
        return requests.request(
            method, base + path, timeout=8, **kwargs
        )

    try:
        with (
            patch.object(
                routes, "_resolve_request_access",
                side_effect=lambda: (SimpleNamespace(id=7), current["access"]),
            ),
            patch.object(
                routes, "_campaign_v2_allowed_sucursal_keys",
                side_effect=lambda _access: current["scope"],
            ),
            patch.object(
                routes, "IVentasBroadcastProvider",
                side_effect=AssertionError("No provider construction allowed"),
            ) as forbidden_provider,
        ):
            assert api("GET", PREFIX + "/reporting").status_code == 401
            assert api(
                "POST", PREFIX + "/91/submit",
                headers=authorization,
                json={"template_id": 1,
                      "expected_dispatch_fingerprint": "a" * 64},
            ).status_code == 423
            forbidden_provider.assert_not_called()

            report_response = api(
                "GET", PREFIX + "/reporting", headers=authorization,
            )
            assert report_response.status_code == 200, report_response.text
            report = report_response.json()
            qa = next(row for row in report["campaigns"] if row["campaign"]["id"] == 91)
            assert qa["provider_children"]["summary"]["child_count"] == 2
            assert qa["provider_children"]["summary"]["provider_raw_status"] == "partial"
            assert qa["provider_raw"]["successful"] == 1
            assert qa["rates"]["successful_rate"] is None
            assert qa["cost"]["total"] is None

            individual_response = api(
                "GET", PREFIX + "/91/report", headers=authorization,
            )
            assert individual_response.status_code == 200, individual_response.text
            assert individual_response.json()["provider_children"]["summary"]["child_count"] == 2

            xlsx_response = api(
                "GET", PREFIX + "/reporting/export", headers=authorization,
            )
            assert xlsx_response.status_code == 200, xlsx_response.text
            workbook = load_workbook(BytesIO(xlsx_response.content), read_only=True)
            assert "Envíos proveedor" in workbook.sheetnames
            workbook.close()

            # Wrong-child lookup must be rejected before any provider request.
            invalid_child = api(
                "GET", PREFIX + "/91/provider-stats?provider_campaign_child_id=9999",
                headers=authorization,
            )
            assert invalid_child.status_code == 409
            forbidden_provider.assert_not_called()

            # A user restricted to branch A cannot see an A+B frozen cohort.
            current["scope"] = ("BRANCH A",)
            current["access"] = SimpleNamespace(
                is_global=False, branch_ids=(1,), can_manage_campaigns=True,
                can_edit_inputs=True, can_preflight_campaigns=True,
                can_manage_dispatch_config=False, can_send_campaigns=False,
            )
            hidden = api(
                "GET", PREFIX + "/91/report", headers=authorization,
            )
            assert hidden.status_code == 404
            limited = api("GET", PREFIX + "/reporting", headers=authorization)
            assert limited.status_code == 200
            assert not limited.json()["campaigns"]

            # Explicit schedule remains an input to server-side preflight.
            current["access"] = access
            current["scope"] = None
            schedule = {
                "local_datetime": "2026-10-21T08:20",
                "timezone": "America/Tijuana",
            }
            with (
                patch.object(
                    routes, "build_campaign_v2_preflight",
                    return_value=SimpleNamespace(),
                ) as preflight,
                patch.object(
                    routes, "serialize_campaign_v2_preflight",
                    return_value={"campaign_id": 91, "ready": True},
                ),
                patch.object(
                    routes, "get_campaign_v2_submit_state",
                    return_value={
                        "status": "NOT_STARTED",
                        "has_provider_campaigns": False, "batches": [],
                    },
                ),
            ):
                planned = api(
                    "POST", PREFIX + "/91/preflight", headers=authorization,
                    json={"template_id": 10, "schedule": schedule},
                )
                assert planned.status_code == 200, planned.text
                assert planned.json()["submission"]["enabled"] is False
                assert preflight.call_args.kwargs["schedule"] == schedule
            forbidden_provider.assert_not_called()
    finally:
        server.shutdown()
        thread.join(timeout=5)
        with app.app_context():
            db.session.remove()
        metadata.drop_all(engine)
        engine.dispose()


@pytest.mark.skipif(not URL, reason="Dedicated PostgreSQL test DB is required")
def test_scheduler_multiple_postgresql_cycles_isolate_failures_and_dedupe():
    """Three real DB cycles, with only a simulated stats reader; never send."""
    from app.services.marketing_campaign_provider import (
        CampaignProviderRawCounts,
        CampaignProviderStats,
    )
    from app.services.marketing_campaign_v2_provider_stats_scheduler_service import (
        run_campaign_v2_provider_stats_capture_cycle,
    )
    from app.services.marketing_campaign_v2_provider_stats_service import (
        MarketingCampaignV2ProviderStatsUpstreamError,
    )
    from app.services.marketing_campaign_v2_provider_stats_snapshot_service import (
        capture_campaign_v2_provider_stats_snapshot,
    )
    from app.services.marketing_campaign_v2_reporting_service import (
        build_campaign_v2_individual_report,
    )

    _guard_isolated_url(URL)
    engine = sa.create_engine(URL, pool_pre_ping=True)
    metadata = _setup_schema(engine)
    now = UTC + timedelta(hours=1)
    calls = []
    state = {"fail_child_b": True}

    def provider_stats(phone: str) -> CampaignProviderStats:
        return CampaignProviderStats(
            analytics_status="ok",
            analytics={"responders": 0},
            raw_counts=CampaignProviderRawCounts(
                successful=1, failed=0, sent=0, delivered=1, viewed=0,
                answered=0, interaction_groups=0, interaction_items=0,
            ),
            successful_phones=frozenset({phone}),
            failed_phones=frozenset(),
            sent_phones=frozenset(),
            delivered_phones=frozenset({phone}),
            viewed_phones=frozenset(),
            button_interactions=(),
        )

    stats_by_child = {
        "ci-fake-batch-a": provider_stats("mx10:6860000701"),
        "ci-fake-batch-b": provider_stats("mx10:6860000702"),
    }

    class FakeStatsProvider:
        def get_campaign_stats(self, provider_campaign_id):
            calls.append(provider_campaign_id)
            if provider_campaign_id == "ci-fake-batch-b" and state["fail_child_b"]:
                raise MarketingCampaignV2ProviderStatsUpstreamError(
                    "test-only upstream failure", retryable=True,
                )
            return stats_by_child[provider_campaign_id]

    provider = FakeStatsProvider()
    try:
        with Session(engine) as session:
            def capture(**kwargs):
                return capture_campaign_v2_provider_stats_snapshot(
                    **kwargs, provider_resolver=lambda _key: provider,
                )

            first = run_campaign_v2_provider_stats_capture_cycle(
                now=now, horizon_hours=48, max_campaigns=10,
                session=session, capture_func=capture,
            )
            assert (first.selected, first.attempted, first.created,
                    first.unchanged, first.failed) == (2, 2, 1, 0, 1)
            # The successful child persists even when its sibling errors.
            assert session.query(Snapshot).filter_by(
                provider_campaign_child_id=901,
            ).count() == 2  # one seeded snapshot plus the newly captured one
            assert session.query(Snapshot).filter_by(
                provider_campaign_child_id=902,
            ).count() == 0

            state["fail_child_b"] = False
            second = run_campaign_v2_provider_stats_capture_cycle(
                now=now + timedelta(hours=1), horizon_hours=48, max_campaigns=10,
                session=session, capture_func=capture,
            )
            assert (second.selected, second.created, second.unchanged,
                    second.failed) == (2, 1, 1, 0)
            assert session.query(Snapshot).filter_by(
                provider_campaign_child_id=902,
            ).count() == 1

            third = run_campaign_v2_provider_stats_capture_cycle(
                now=now + timedelta(hours=2), horizon_hours=48, max_campaigns=10,
                session=session, capture_func=capture,
            )
            assert (third.selected, third.created, third.unchanged,
                    third.failed) == (2, 0, 2, 0)

            snapshots_before = session.query(Snapshot).count()
            observations_before = session.query(Observation).count()
            assert snapshots_before == 3
            assert observations_before == 3  # seeded + one per child
            assert calls == [
                "ci-fake-batch-a", "ci-fake-batch-b",
                "ci-fake-batch-a", "ci-fake-batch-b",
                "ci-fake-batch-a", "ci-fake-batch-b",
            ]

            report = build_campaign_v2_individual_report(
                campaign_id=91, allowed_sucursal_keys=None, session=session,
            )
            assert report["provider_children"]["summary"]["child_count"] == 2
            assert report["provider_children"]["summary"]["provider_raw_status"] == "complete"
            assert report["normalized"]["successful"] == 2
            assert report["provider_raw"]["successful"] == 2
            assert report["rates"]["successful_rate"] == 1
            assert report["cost"]["total"] is None

            # Verify persisted data remains unchanged by read-only reporting.
            assert session.query(Snapshot).count() == snapshots_before
            assert session.query(Observation).count() == observations_before
    finally:
        metadata.drop_all(engine)
        engine.dispose()
