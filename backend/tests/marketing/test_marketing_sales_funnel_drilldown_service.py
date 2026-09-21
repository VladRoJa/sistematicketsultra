from openpyxl import load_workbook

from app.services.marketing_sales_funnel_detail_service import (
    MarketingSalesFunnelDetailValidationError,
)
from app.services.marketing_sales_funnel_drilldown_service import (
    _build_excel_workbook,
    _normalize_sort,
    _sort_rows,
)


def test_normalize_sort_accepts_supported_sales_columns():
    assert _normalize_sort("revenue", "desc", "sales") == (
        "revenue",
        "desc",
    )
    assert _normalize_sort(None, None, "sales") == (None, "asc")


def test_normalize_sort_rejects_unknown_column():
    try:
        _normalize_sort("not_a_column", "asc", "sales")
    except MarketingSalesFunnelDetailValidationError:
        pass
    else:
        raise AssertionError("sort_by inválido debió rechazarse")


def test_normalize_sort_rejects_invalid_direction():
    try:
        _normalize_sort("branch", "sideways", "sales")
    except MarketingSalesFunnelDetailValidationError:
        pass
    else:
        raise AssertionError("sort_dir inválido debió rechazarse")


def test_sort_rows_keeps_empty_values_at_end_for_both_directions():
    rows = [
        {"branch": "Zulu", "revenue": 100},
        {"branch": None, "revenue": 300},
        {"branch": "Alpha", "revenue": 200},
    ]

    asc = _sort_rows(rows, "branch", "asc")
    desc = _sort_rows(rows, "branch", "desc")

    assert [row["branch"] for row in asc] == ["Alpha", "Zulu", None]
    assert [row["branch"] for row in desc] == ["Zulu", "Alpha", None]


def test_sort_rows_orders_numeric_values_numerically():
    rows = [
        {"revenue": 9},
        {"revenue": 100},
        {"revenue": 20},
    ]

    ordered = _sort_rows(rows, "revenue", "desc")
    assert [row["revenue"] for row in ordered] == [100, 20, 9]


def test_build_excel_workbook_contains_all_sales_rows():
    output = _build_excel_workbook(
        title="Venta nueva oficial",
        kind="sales",
        rows=[
            {
                "branch": "TEC MXL",
                "date": "2026-09-05",
                "name": "Persona Uno",
                "member_id": "123",
                "pin": "456",
                "phone": "6861234567",
                "folio": "F-1",
                "membership_type": "Mensual",
                "tariff": "1 MES $899",
                "revenue": 899,
                "origin": "Familiares o amigos",
                "survey": "Familiares o Amigos",
                "transaction_branch": "VILLA VERDE",
                "payment_method": "EFECTIVO",
                "id_order": "ORD-1",
                "payment_place": "CAJA",
            },
            {
                "branch": "PASEO LA PAZ",
                "date": "2026-09-08",
                "name": "Persona Dos",
                "revenue": 1499,
            },
        ],
    )

    workbook = load_workbook(output)
    worksheet = workbook["Detalle"]

    assert worksheet.max_row == 3
    assert worksheet["A1"].value == "Sucursal KPI"
    assert worksheet["A2"].value == "TEC MXL"
    assert worksheet["J2"].value == 899
    assert worksheet["A3"].value == "PASEO LA PAZ"


def test_meta_leads_sort_accepts_followup_status_only_for_meta_metric():
    assert _normalize_sort(
        "followup_status",
        "asc",
        "leads",
        "leads_meta",
    ) == ("followup_status", "asc")

    try:
        _normalize_sort("followup_status", "asc", "leads")
    except MarketingSalesFunnelDetailValidationError:
        pass
    else:
        raise AssertionError(
            "followup_status solo debe habilitarse para leads_meta"
        )


def test_meta_leads_export_includes_call_center_followup_columns():
    output = _build_excel_workbook(
        title="Leads Meta",
        kind="leads",
        metric="leads_meta",
        rows=[
            {
                "branch": "TEC MXL",
                "date": "2026-09-01",
                "name": "Lead Uno",
                "phone": "6861234567",
                "followup_status": "Visita sin compra",
                "visit_status": "Sí",
                "visit_date": "2026-09-04",
                "purchase_status": "No",
                "sale_date": None,
                "channel": "Ultragym Tecnologico",
                "contact_id": "contact-1",
            },
        ],
    )

    workbook = load_workbook(output)
    worksheet = workbook["Detalle"]
    headers = [cell.value for cell in worksheet[1]]

    assert headers == [
        "Sucursal KPI",
        "Fecha primer mensaje",
        "Nombre",
        "Teléfono",
        "Seguimiento",
        "Visitó",
        "Fecha visita",
        "Compró",
        "Fecha compra",
        "Canal",
        "ID contacto",
    ]
    assert worksheet["E2"].value == "Visita sin compra"
    assert worksheet["F2"].value == "Sí"
    assert worksheet["H2"].value == "No"

