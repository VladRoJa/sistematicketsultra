from __future__ import annotations

from datetime import datetime, timezone
import inspect
from io import BytesIO

from openpyxl import load_workbook

from app.services import marketing_campaign_v2_reporting_excel_service as excel_service
from app.services.marketing_campaign_v2_reporting_excel_service import (
    build_campaign_v2_reporting_excel,
)


GENERATED_AT = datetime(2026, 10, 3, 23, 45, tzinfo=timezone.utc)


def _campaign_row(
    campaign_id=1,
    *,
    name="Campaign Safe",
    snapshot_id=10,
    observed_at="2026-10-03T20:00:00+00:00",
    total=10,
    successful=6,
    failed=2,
    sent=1,
    delivered=2,
    viewed=3,
    reach=5,
    matched=8,
    unmatched=1,
    without_status=2,
    raw_successful=8,
    raw_failed=3,
):
    provider_raw = None
    if snapshot_id is not None:
        provider_raw = {
            "successful": raw_successful,
            "failed": raw_failed,
            "sent": 2,
            "delivered": 3,
            "viewed": 3,
            "answered": 4,
            "interaction_groups": 1,
            "interaction_items": 5,
        }
    return {
        "campaign": {
            "id": campaign_id,
            "name": name,
            "purpose": "REACTIVATION",
            "source": "EXPIRED_MEMBERS",
            "provider": "IVENTAS",
            "provider_campaign_id": f"external-{campaign_id}",
            "frozen_at": "2026-10-01T18:00:00+00:00",
        },
        "observation": {
            "snapshot_id": snapshot_id,
            "latest_observed_at": observed_at if snapshot_id is not None else None,
            "analytics_status": "ok" if snapshot_id is not None else None,
        },
        "audience": {"total_recipients": total},
        "normalized": {
            "successful": successful,
            "failed": failed,
            "sent": sent,
            "delivered": delivered,
            "viewed": viewed,
            "reach_count": reach,
        },
        "rates": {
            "successful_rate": successful / total if total else None,
            "reach_rate": reach / total if total else None,
            "read_rate": viewed / reach if reach else None,
            "failure_rate": failed / total if total else None,
        },
        "coverage": {
            "matched_recipient_count": matched,
            "unmatched_provider_count": unmatched,
            "frozen_recipient_without_provider_status_count": without_status,
            "status_coverage_rate": matched / total if total else None,
        },
        "provider_raw": provider_raw,
        "interactions": {
            "unique_button_recipients": 3,
            "responders_aggregate": 4,
            "free_text_aggregate": 1,
        },
        "cost": {
            "status": "unavailable",
            "currency": None,
            "total": None,
        },
    }


def _consolidated_report(*, empty=False, malicious_name=None):
    if empty:
        campaigns = []
        total = 0
        normalized = {
            "successful": 0,
            "failed": 0,
            "sent": 0,
            "delivered": 0,
            "viewed": 0,
            "reach_count": 0,
        }
        matched = 0
        unmatched = 0
        without_status = 0
        provider_raw = {
            "successful": 0,
            "failed": 0,
            "sent": 0,
            "delivered": 0,
            "viewed": 0,
            "answered": 0,
            "interaction_groups": 0,
            "interaction_items": 0,
        }
    else:
        campaigns = [
            _campaign_row(
                2,
                name=malicious_name or "Large",
                snapshot_id=21,
                observed_at="2026-10-03T22:00:00+00:00",
                total=8,
                successful=3,
                failed=2,
                sent=1,
                delivered=0,
                viewed=2,
                reach=2,
                matched=4,
                unmatched=1,
                without_status=4,
                raw_successful=6,
                raw_failed=4,
            ),
            _campaign_row(
                1,
                name="Small",
                snapshot_id=11,
                observed_at="2026-10-03T20:00:00+00:00",
                total=2,
                successful=2,
                failed=0,
                sent=0,
                delivered=1,
                viewed=1,
                reach=2,
                matched=2,
                unmatched=1,
                without_status=0,
                raw_successful=3,
                raw_failed=1,
            ),
            _campaign_row(
                3,
                name="Without snapshot",
                snapshot_id=None,
                observed_at=None,
                total=2,
                successful=0,
                failed=0,
                sent=0,
                delivered=0,
                viewed=0,
                reach=0,
                matched=0,
                unmatched=0,
                without_status=2,
                raw_successful=0,
                raw_failed=0,
            ),
        ]
        total = 12
        normalized = {
            "successful": 5,
            "failed": 2,
            "sent": 1,
            "delivered": 1,
            "viewed": 3,
            "reach_count": 4,
        }
        matched = 6
        unmatched = 2
        without_status = 6
        provider_raw = {
            "successful": 9,
            "failed": 5,
            "sent": 4,
            "delivered": 3,
            "viewed": 2,
            "answered": 7,
            "interaction_groups": 2,
            "interaction_items": 9,
        }

    return {
        "filters": {
            "observed_from": None,
            "observed_to": None,
            "purpose": "REACTIVATION" if not empty else None,
            "source": None,
            "provider": "IVENTAS" if not empty else None,
            "snapshot_status": None,
        },
        "summary": {
            "campaign_count": len(campaigns),
            "campaigns_with_snapshot": sum(
                1
                for row in campaigns
                if row["observation"]["snapshot_id"] is not None
            ),
            "campaigns_without_snapshot": sum(
                1
                for row in campaigns
                if row["observation"]["snapshot_id"] is None
            ),
            "total_recipients": total,
            "normalized": normalized,
            "rates": {
                "successful_rate": normalized["successful"] / total if total else None,
                "reach_rate": normalized["reach_count"] / total if total else None,
                "read_rate": (
                    normalized["viewed"] / normalized["reach_count"]
                    if normalized["reach_count"]
                    else None
                ),
                "failure_rate": normalized["failed"] / total if total else None,
            },
            "coverage": {
                "matched_recipient_count": matched,
                "unmatched_provider_count": unmatched,
                "frozen_recipient_without_provider_status_count": without_status,
                "status_coverage_rate": matched / total if total else None,
            },
            "provider_raw": provider_raw,
            "interactions": {
                "button_interaction_recipient_exposures": 4 if not empty else 0,
                "responders_aggregate": 9 if not empty else None,
                "responders_campaigns_with_value": 2 if not empty else 0,
                "free_text_aggregate": 3 if not empty else None,
                "free_text_campaigns_with_value": 2 if not empty else 0,
            },
            "cost": {
                "status": "unavailable",
                "currency": None,
                "total": None,
            },
        },
        "campaigns": campaigns,
        "breakdowns": {
            "branches": [],
            "audience_families": [],
        },
    }


def _evolution():
    return [
        {
            "campaign_id": 1,
            "campaign_name": "Small",
            "snapshot_id": 10,
            "observed_at": "2026-10-03T19:00:00+00:00",
            "normalized": {
                "successful": 1,
                "failed": 0,
                "sent": 1,
                "delivered": 0,
                "viewed": 0,
                "reach_count": 0,
            },
            "provider_raw": {
                "successful": 2,
                "failed": 0,
                "sent": 2,
                "delivered": 0,
                "viewed": 0,
            },
            "coverage": {
                "matched_recipient_count": 1,
                "unmatched_provider_count": 1,
            },
        },
        {
            "campaign_id": 1,
            "campaign_name": "Small",
            "snapshot_id": 11,
            "observed_at": "2026-10-03T20:00:00+00:00",
            "normalized": {
                "successful": 2,
                "failed": 0,
                "sent": 0,
                "delivered": 1,
                "viewed": 1,
                "reach_count": 2,
            },
            "provider_raw": {
                "successful": 3,
                "failed": 1,
                "sent": 1,
                "delivered": 1,
                "viewed": 1,
            },
            "coverage": {
                "matched_recipient_count": 2,
                "unmatched_provider_count": 1,
            },
        },
        {
            "campaign_id": 2,
            "campaign_name": "Large",
            "snapshot_id": 21,
            "observed_at": "2026-10-03T22:00:00+00:00",
            "normalized": {
                "successful": 3,
                "failed": 2,
                "sent": 1,
                "delivered": 0,
                "viewed": 2,
                "reach_count": 2,
            },
            "provider_raw": {
                "successful": 6,
                "failed": 4,
                "sent": 3,
                "delivered": 2,
                "viewed": 1,
            },
            "coverage": {
                "matched_recipient_count": 4,
                "unmatched_provider_count": 1,
            },
        },
    ]


def _load(output):
    return load_workbook(BytesIO(output.getvalue()), data_only=False)


def _sheet_rows_by_key(sheet):
    return {
        sheet.cell(row, 1).value: sheet.cell(row, 2).value
        for row in range(2, sheet.max_row + 1)
    }


def _header_map(sheet):
    return {
        sheet.cell(1, col).value: col
        for col in range(1, sheet.max_column + 1)
    }


def test_consolidated_workbook_has_exact_five_sheets_and_filename():
    output, filename = build_campaign_v2_reporting_excel(
        report_type="CONSOLIDATED",
        report=_consolidated_report(),
        evolution=_evolution(),
        scope={"is_global": False, "allowed_sucursal_keys": ["BRANCH A"]},
        generated_at=GENERATED_AT,
    )
    workbook = _load(output)

    assert workbook.sheetnames == [
        "Resumen",
        "Campañas",
        "KPIs",
        "Evolución",
        "Metadata",
    ]
    assert filename == "campaign_v2_reporting_20261003_234500.xlsx"


def test_summary_matches_consolidated_dto_and_labels_recipient_exposures():
    output, _ = build_campaign_v2_reporting_excel(
        report_type="CONSOLIDATED",
        report=_consolidated_report(),
        evolution=_evolution(),
        scope={"is_global": False, "allowed_sucursal_keys": ["BRANCH A"]},
        generated_at=GENERATED_AT,
    )
    rows = _sheet_rows_by_key(_load(output)["Resumen"])

    assert rows["generated_at"] == "2026-10-03T23:45:00+00:00"
    assert rows["report_type"] == "CONSOLIDATED"
    assert rows["latest_observed_at"] == "2026-10-03T22:00:00+00:00"
    assert rows["campaign_count"] == 3
    assert rows["Exposiciones de destinatarios"] == 12
    assert rows["successful"] == 5
    assert rows["reach_count"] == 4
    assert rows["successful_rate"] == 5 / 12
    assert rows["cost_status"] == "unavailable"
    assert rows["filter.purpose"] == "REACTIVATION"


def test_campaign_sheet_preserves_raw_normalized_and_snapshotless_semantics():
    output, _ = build_campaign_v2_reporting_excel(
        report_type="CONSOLIDATED",
        report=_consolidated_report(),
        evolution=_evolution(),
        scope={"is_global": False, "allowed_sucursal_keys": ["BRANCH A"]},
        generated_at=GENERATED_AT,
    )
    sheet = _load(output)["Campañas"]
    headers = _header_map(sheet)

    assert sheet.max_row == 4
    first = {
        key: sheet.cell(2, column).value
        for key, column in headers.items()
    }
    assert first["campaign_id"] == 2
    assert first["successful"] == 3
    assert first["raw_successful"] == 6
    assert first["successful"] != first["raw_successful"]
    assert first["successful_rate"] == 3 / 8
    assert sheet.cell(2, headers["successful_rate"]).number_format == "0.00%"

    snapshotless = {
        key: sheet.cell(4, column).value
        for key, column in headers.items()
    }
    assert snapshotless["snapshot_id"] is None
    assert snapshotless["latest_observed_at"] is None
    assert snapshotless["successful"] == 0
    assert snapshotless["raw_successful"] is None


def test_kpi_sheet_uses_dto_rates_and_real_numerator_denominator():
    report = _consolidated_report()
    output, _ = build_campaign_v2_reporting_excel(
        report_type="CONSOLIDATED",
        report=report,
        evolution=_evolution(),
        scope={"is_global": False, "allowed_sucursal_keys": ["BRANCH A"]},
        generated_at=GENERATED_AT,
    )
    sheet = _load(output)["KPIs"]
    headers = _header_map(sheet)
    rows = {
        sheet.cell(row, headers["metric"]).value: {
            **{
                key: sheet.cell(row, column).value
                for key, column in headers.items()
            },
            "_row": row,
        }
        for row in range(2, sheet.max_row + 1)
    }

    assert {row["type"] for row in rows.values()} == {
        "NORMALIZED",
        "PROVIDER_RAW",
        "PROVIDER_AGGREGATE",
        "SUITE_DERIVED",
    }
    assert rows["successful"]["value"] == 5
    assert rows["raw_successful"]["value"] == 9
    assert rows["responders_aggregate"]["value"] == 9

    successful = rows["successful_rate"]
    assert successful["value"] == report["summary"]["rates"]["successful_rate"]
    assert successful["numerator"] == 5
    assert successful["denominator"] == 12
    assert successful["type"] == "SUITE_DERIVED"
    assert (
        sheet.cell(successful["_row"], headers["value"]).number_format
        == "0.00%"
    )

    read_rate = rows["read_rate"]
    assert read_rate["numerator"] == 3
    assert read_rate["denominator"] == 4


def test_empty_consolidated_export_is_valid_with_blank_rates_and_headers_only():
    output, _ = build_campaign_v2_reporting_excel(
        report_type="CONSOLIDATED",
        report=_consolidated_report(empty=True),
        evolution=[],
        scope={"is_global": True, "allowed_sucursal_keys": None},
        generated_at=GENERATED_AT,
    )
    workbook = _load(output)
    summary = _sheet_rows_by_key(workbook["Resumen"])
    kpis = workbook["KPIs"]

    assert summary["campaign_count"] == 0
    assert summary["Exposiciones de destinatarios"] == 0
    assert summary["successful_rate"] is None
    assert workbook["Campañas"].max_row == 1
    assert workbook["Evolución"].max_row == 1
    headers = _header_map(kpis)
    rate_row = next(
        row
        for row in range(2, kpis.max_row + 1)
        if kpis.cell(row, headers["metric"]).value == "successful_rate"
    )
    assert kpis.cell(rate_row, headers["value"]).value is None


def test_evolution_multi_campaign_preserves_historical_values_and_observed_label():
    output, _ = build_campaign_v2_reporting_excel(
        report_type="CONSOLIDATED",
        report=_consolidated_report(),
        evolution=_evolution(),
        scope={"is_global": False, "allowed_sucursal_keys": ["BRANCH A"]},
        generated_at=GENERATED_AT,
    )
    sheet = _load(output)["Evolución"]
    headers = _header_map(sheet)

    assert sheet.max_row == 4
    assert [sheet.cell(row, headers["campaign_id"]).value for row in range(2, 5)] == [
        1,
        1,
        2,
    ]
    assert sheet.cell(2, headers["snapshot_id"]).value == 10
    assert sheet.cell(3, headers["snapshot_id"]).value == 11
    assert sheet.cell(2, headers["normalized_successful"]).value == 1
    assert sheet.cell(3, headers["normalized_successful"]).value == 2
    assert sheet.cell(4, headers["observed_at_meaning"]).value == "Observed by Suite"


def test_metadata_contains_scope_filters_definitions_and_cost_unavailable():
    output, _ = build_campaign_v2_reporting_excel(
        report_type="CONSOLIDATED",
        report=_consolidated_report(),
        evolution=_evolution(),
        scope={"is_global": False, "allowed_sucursal_keys": ["BRANCH A"]},
        generated_at=GENERATED_AT,
    )
    rows = _sheet_rows_by_key(_load(output)["Metadata"])

    assert rows["scope.is_global"] is False
    assert rows["scope.allowed_sucursal_keys"] == "BRANCH A"
    assert rows["definition.total_recipients"].startswith("Consolidado = recipient exposures")
    assert rows["definition.reach_count"] == "Unique frozen recipients in DELIVERED union VIEWED"
    assert rows["definition.provider_raw"].startswith("Provider raw counts")
    assert rows["definition.observed_at"] == "Observed by Suite (snapshot.fetched_at)"
    assert rows["definition.cost"] == "Costos actualmente no disponibles"
    assert rows["filter.provider"] == "IVENTAS"


def test_individual_export_has_one_campaign_row_and_uses_m27_rates():
    report = _campaign_row(9, name="One")
    report["dimensions"] = {
        "scope": {
            "is_global": False,
            "allowed_sucursal_keys": ["BRANCH A"],
        }
    }
    report["evolution"] = [
        {
            "snapshot_id": 10,
            "observed_at": "2026-10-03T20:00:00+00:00",
            "normalized": report["normalized"],
            "provider_raw": {
                "successful": 8,
                "failed": 3,
                "sent": 2,
                "delivered": 3,
                "viewed": 3,
            },
            "coverage": {
                "matched_recipient_count": 8,
                "unmatched_provider_count": 1,
            },
        }
    ]
    evolution = [
        {
            "campaign_id": 9,
            "campaign_name": "One",
            **report["evolution"][0],
        }
    ]

    output, filename = build_campaign_v2_reporting_excel(
        report_type="INDIVIDUAL",
        report=report,
        evolution=evolution,
        scope=report["dimensions"]["scope"],
        generated_at=GENERATED_AT,
    )
    workbook = _load(output)
    campaigns = workbook["Campañas"]
    summary = _sheet_rows_by_key(workbook["Resumen"])

    assert campaigns.max_row == 2
    assert summary["campaign_id"] == 9
    assert summary["campaign_name"] == "One"
    assert summary["successful_rate"] == report["rates"]["successful_rate"]
    assert filename == "campaign_v2_9_report_20261003_234500.xlsx"


def test_formula_injection_is_sanitized_as_text():
    malicious = '=HYPERLINK("https://example.invalid","click")'
    output, _ = build_campaign_v2_reporting_excel(
        report_type="CONSOLIDATED",
        report=_consolidated_report(malicious_name=malicious),
        evolution=[
            {
                **_evolution()[0],
                "campaign_name": "@SUM(1,1)",
            }
        ],
        scope={"is_global": False, "allowed_sucursal_keys": ["BRANCH A"]},
        generated_at=GENERATED_AT,
    )
    workbook = _load(output)
    campaigns = workbook["Campañas"]
    headers = _header_map(campaigns)
    campaign_name_cell = campaigns.cell(2, headers["name"])
    evolution_name_cell = workbook["Evolución"].cell(
        2,
        _header_map(workbook["Evolución"])["campaign_name"],
    )

    assert campaign_name_cell.data_type != "f"
    assert campaign_name_cell.value.startswith("'=")
    assert evolution_name_cell.data_type != "f"
    assert evolution_name_cell.value.startswith("'@")


def test_workbook_contains_no_recipient_pii_or_provider_payload():
    fixture_phone = "6860000001"
    report = _consolidated_report()
    report["campaigns"][0]["secret_test_value"] = fixture_phone

    output, _ = build_campaign_v2_reporting_excel(
        report_type="CONSOLIDATED",
        report=report,
        evolution=_evolution(),
        scope={"is_global": False, "allowed_sucursal_keys": ["BRANCH A"]},
        generated_at=GENERATED_AT,
    )
    workbook = _load(output)
    values = [
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]
    joined = "\n".join(values)

    assert fixture_phone not in joined
    assert "phone_mx10" not in joined
    assert "member_pin" not in joined
    assert "member_name" not in joined
    assert "analytics_json" not in joined
    assert "audience_definition_json" not in joined
    assert "Authorization" not in joined


def test_same_dto_and_generated_at_have_same_logical_workbook_values():
    kwargs = {
        "report_type": "CONSOLIDATED",
        "report": _consolidated_report(),
        "evolution": _evolution(),
        "scope": {"is_global": False, "allowed_sucursal_keys": ["BRANCH A"]},
        "generated_at": GENERATED_AT,
    }
    first, _ = build_campaign_v2_reporting_excel(**kwargs)
    second, _ = build_campaign_v2_reporting_excel(**kwargs)
    first_book = _load(first)
    second_book = _load(second)

    for sheet_name in first_book.sheetnames:
        first_values = [
            [cell.value for cell in row]
            for row in first_book[sheet_name].iter_rows()
        ]
        second_values = [
            [cell.value for cell in row]
            for row in second_book[sheet_name].iter_rows()
        ]
        assert first_values == second_values


def test_excel_service_is_pure_and_has_no_db_reporting_or_provider_dependencies():
    source = inspect.getsource(excel_service)
    for forbidden in (
        "app.extensions",
        "app.models",
        "marketing_campaign_v2_reporting_service",
        "marketing_campaign_v2_provider",
        "get_campaign_v2_provider_stats",
        "capture_campaign_v2_provider_stats_snapshot",
        "get_provider_history_for_phones",
        "requests.",
        "httpx",
    ):
        assert forbidden not in source
