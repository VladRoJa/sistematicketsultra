from __future__ import annotations

from datetime import date
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook

from app.services.system_daily_check_export_service import (
    build_system_daily_check_excel_export,
)


def _summary():
    return {
        "filters": {
            "date_from": "2026-10-01",
            "date_to": "2026-10-09",
            "branch_id": 100,
        },
        "universe": {
            "expected_checklists": 9,
            "expected_branches": 1,
            "out_of_rollout_submissions": 0,
        },
        "summary": {
            "completed": 8,
            "pending": 1,
            "compliance_pct": 88.9,
            "normal": 6,
            "minor_failure": 2,
            "operational_impact": 0,
        },
        "failures": {
            "checks_with_failure": 2,
            "no_answers": 3,
            "issues_total": 3,
            "reported": 2,
            "unreported": 1,
            "reported_pct": 66.7,
        },
        "postponements": {
            "completed_without_postpone": 6,
            "completed_after_1": 1,
            "completed_after_2": 1,
            "reached_mandatory": 1,
        },
    }


def _history():
    return {
        "page": 1,
        "page_size": 200,
        "total": 1,
        "items": [
            {
                "id": 71,
                "sucursal_id": 100,
                "sucursal": "Corporativo",
                "performed_by_user_id": 47,
                "performed_by_username": "admcorp",
                "business_date": "2026-10-09",
                "general_status": "MINOR_FAILURE",
                "submitted_at": "2026-10-09T13:08:43-07:00",
                "postpone_count": 0,
                "reached_mandatory": False,
            }
        ],
    }


def _issues():
    return {
        "page": 1,
        "page_size": 200,
        "total": 1,
        "items": [
            {
                "issue_id": 91,
                "check_id": 71,
                "sucursal_id": 100,
                "sucursal": "Corporativo",
                "business_date": "2026-10-09",
                "question_key": "BI_REPORTS_WORKING",
                "question_label": "¿Los reportes BI funcionan?",
                "reported_to_support": True,
                "affected_scope": None,
                "description": "Un tablero no cargó.",
                "attachment_count": 1,
            }
        ],
    }


def _detail():
    return {
        "id": 71,
        "sucursal_id": 100,
        "sucursal": "Corporativo",
        "performed_by_user_id": 47,
        "performed_by_username": "admcorp",
        "business_date": "2026-10-09",
        "general_status": "MINOR_FAILURE",
        "answers": [
            {
                "id": 501,
                "question_key": "BI_REPORTS_WORKING",
                "question_label": "¿Los reportes BI funcionan?",
                "category_key": "CONNECTIVITY_SYSTEMS",
                "answer": "NO",
                "issue": {
                    "id": 91,
                    "affected_scope": None,
                    "reported_to_support": True,
                    "description": "Un tablero no cargó.",
                    "attachments": [
                        {
                            "id": 12,
                            "original_filename": "evidencia.png",
                            "mime_type": "image/png",
                            "file_size_bytes": 1500,
                            "sha256": "abc",
                            "url": "/api/system-daily-checks/bi/attachments/12",
                        }
                    ],
                },
            }
        ],
    }


def test_excel_export_contains_reconciled_operational_sheets():
    actor = SimpleNamespace(
        id=1,
        rol="SISTEMAS",
        username="sistemas",
    )

    with (
        patch(
            "app.services.system_daily_check_export_service."
            "build_system_daily_check_bi_summary",
            return_value=_summary(),
        ),
        patch(
            "app.services.system_daily_check_export_service."
            "list_system_daily_check_bi_history",
            return_value=_history(),
        ),
        patch(
            "app.services.system_daily_check_export_service."
            "list_system_daily_check_bi_issues",
            return_value=_issues(),
        ),
        patch(
            "app.services.system_daily_check_export_service."
            "get_system_daily_check_bi_detail",
            return_value=_detail(),
        ),
    ):
        output, filename = build_system_daily_check_excel_export(
            actor,
            date_from=date(2026, 10, 1),
            date_to=date(2026, 10, 9),
            branch_id=100,
        )

    assert isinstance(output, BytesIO)
    assert filename == (
        "salud_sistemas_2026-10-01_2026-10-09.xlsx"
    )

    workbook = load_workbook(output, data_only=True)
    assert workbook.sheetnames == [
        "Resumen",
        "Checklists",
        "Incidencias",
        "Respuestas",
    ]

    summary = workbook["Resumen"]
    values = {
        summary.cell(row=row, column=1).value:
        summary.cell(row=row, column=2).value
        for row in range(1, summary.max_row + 1)
    }
    assert values["Checklists esperados"] == 9
    assert values["Completados"] == 8
    assert values["Pendientes"] == 1
    assert values["Respuestas NO"] == 3
    assert values["Incidencias reportadas"] == 2

    checklists = workbook["Checklists"]
    assert checklists.max_row == 2
    assert checklists["A2"].value == 71
    assert checklists["D2"].value == "Corporativo"
    assert checklists["E2"].value == "MINOR_FAILURE"

    issues = workbook["Incidencias"]
    assert issues.max_row == 2
    assert issues["A2"].value == 91
    assert issues["F2"].value == "BI_REPORTS_WORKING"
    assert issues["H2"].value == "Sí"
    assert issues["K2"].value == 1

    answers = workbook["Respuestas"]
    assert answers.max_row == 2
    assert answers["A2"].value == 71
    assert answers["F2"].value == "BI_REPORTS_WORKING"
    assert answers["I2"].value == "NO"
    assert answers["J2"].value == "Sí"
    assert answers["N2"].value == 1
