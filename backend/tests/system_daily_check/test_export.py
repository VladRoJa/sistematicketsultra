"""Regresión de exportación del checklist BI sin dependencia de base de datos."""

from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import load_workbook
import pytest

from app.services import system_daily_check_export_service as exporter
from app.services.system_daily_check_service import SystemDailyCheckValidationError


def _summary():
    return {
        "universe": {
            "expected_branches": 1,
            "expected_checklists": 101,
            "out_of_rollout_submissions": 0,
        },
        "summary": {
            "completed": 101,
            "pending": 0,
            "compliance_pct": 100.0,
            "normal": 100,
            "minor_failure": 1,
            "operational_impact": 0,
        },
        "failures": {
            "checks_with_failure": 1,
            "no_answers": 1,
            "issues_total": 201,
            "reported": 200,
            "unreported": 1,
            "reported_pct": 99.5,
        },
        "postponements": {
            "completed_without_postpone": 101,
            "completed_after_1": 0,
            "completed_after_2": 0,
            "reached_mandatory": 0,
        },
    }


def _history(index):
    return {
        "id": index,
        "sucursal": "PRUEBA",
        "business_date": "2026-10-09",
        "performed_by_username": "operador",
        "general_status": "MINOR_FAILURE" if index == 1 else "NORMAL",
        "submitted_at": "2026-10-09T20:08:43+00:00",
        "postpone_count": 0,
        "reached_mandatory": False,
    }


def _issue(index):
    return {
        "issue_id": index,
        "check_id": 1,
        "business_date": "2026-10-09",
        "sucursal": "PRUEBA",
        "question_label": "=HYPERLINK(\"https://evil.test\")" if index == 1 else "Pregunta",
        "affected_scope": "Terminales",
        "description": "Falla observada",
        "reported_to_support": index != 1,
        "attachment_count": 1,
    }


def test_export_all_pages_and_four_sheets(monkeypatch):
    seen_history_pages = []
    seen_issue_pages = []
    seen_details = []

    monkeypatch.setattr(
        exporter, "build_system_daily_check_bi_summary",
        lambda actor, **kwargs: _summary(),
    )

    def history(actor, **kwargs):
        page = kwargs["page"]
        seen_history_pages.append(page)
        first = (page - 1) * 100
        return {
            "total": 101,
            "items": [_history(i) for i in range(first + 1, min(101, first + 100) + 1)],
        }

    def issues(actor, **kwargs):
        page = kwargs["page"]
        seen_issue_pages.append(page)
        first = (page - 1) * 200
        return {
            "total": 201,
            "items": [_issue(i) for i in range(first + 1, min(201, first + 200) + 1)],
        }

    def detail(actor, *, check_id):
        seen_details.append(check_id)
        return {
            **_history(check_id),
            "answers": [{
                "category_key": "INTERNET",
                "question_label": "¿Funciona internet?",
                "answer": "NO" if check_id == 1 else "YES",
                "issue": {
                    "description": "-Cuidado: texto libre",
                    "affected_scope": "Sucursal",
                    "reported_to_support": False,
                    "attachments": [{"id": 1}],
                } if check_id == 1 else None,
            }],
        }

    monkeypatch.setattr(exporter, "list_system_daily_check_bi_history", history)
    monkeypatch.setattr(exporter, "list_system_daily_check_bi_issues", issues)
    monkeypatch.setattr(exporter, "get_system_daily_check_bi_detail", detail)

    result = exporter.export_system_daily_check_bi_workbook(
        object(), date_from=date(2026, 10, 1), date_to=date(2026, 10, 9),
        branch_id=1000,
    )
    assert isinstance(result, BytesIO)
    book = load_workbook(result, data_only=False)
    assert book.sheetnames == [
        "Resumen", "Checklists", "Incidencias", "Respuestas",
    ]
    assert book["Checklists"].max_row == 4 + 101
    assert book["Incidencias"].max_row == 4 + 201
    assert book["Respuestas"].max_row == 4 + 101
    assert seen_history_pages == [1, 2]
    assert seen_issue_pages == [1, 2]
    assert len(seen_details) == 101
    assert book["Checklists"]["F5"].value.hour == 13  # Hora Tijuana
    assert book["Incidencias"]["E5"].data_type != "f"
    assert book["Incidencias"]["E5"].value.startswith("'=")
    assert book["Respuestas"]["H5"].value.startswith("'-")
    assert book["Resumen"]["B9"].value == 100.0  # Cumplimiento
    assert all(
        book[sheet].freeze_panes == "A5"
        for sheet in book.sheetnames
    )


@pytest.mark.parametrize(
    "start,end",
    [
        (date(2026, 10, 10), date(2026, 10, 9)),
        (date(2026, 1, 1), date(2026, 10, 9)),
    ],
)
def test_export_rejects_invalid_or_excessive_range(start, end):
    with pytest.raises(SystemDailyCheckValidationError):
        exporter.export_system_daily_check_bi_workbook(
            object(), date_from=start, date_to=end,
        )


def test_export_rejects_too_many_checks_before_loading_details(monkeypatch):
    monkeypatch.setattr(
        exporter, "build_system_daily_check_bi_summary",
        lambda actor, **kwargs: _summary(),
    )
    monkeypatch.setattr(
        exporter, "list_system_daily_check_bi_history",
        lambda actor, **kwargs: {"total": 2501, "items": []},
    )
    with pytest.raises(SystemDailyCheckValidationError, match="2500"):
        exporter.export_system_daily_check_bi_workbook(
            object(), date_from=date(2026, 10, 1),
            date_to=date(2026, 10, 9),
        )
