"""Scheduled Google Ads email XLSX: guarded flat-day parsing and pair preflight.

All workbooks are synthetic. The historical titled exports must continue to
require official subtotals. No DB writes or mutable production evidence.
"""
from __future__ import annotations

from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.warehouse.services.google_ads_daily_xlsx_parser import (
    GoogleAdsXlsxValidationError, parse_google_ads_daily_xlsx,
)
from app.warehouse.services.google_ads_ad_group_daily_xlsx_parser import (
    parse_google_ads_ad_group_daily_xlsx,
)
from test_google_ads_m6_manual_preview import daily_xlsx, post

CAMPAIGN = "campaign"
GROUP = "group"


def flat_email(kind: str, *, mutate=None) -> bytes:
    """Make the actual email layout: first row headings, no official totals.

    Groups have the duplicated Código de moneda heading observed in the email.
    Other columns and their numeric values come from strict synthetic source.
    """
    originals = daily_xlsx()
    source = originals[0] if kind == CAMPAIGN else originals[1]
    old = load_workbook(BytesIO(source), read_only=True, data_only=True)
    try:
        sheet = old.active
        headings = list(next(sheet.iter_rows(min_row=3, max_row=3, values_only=True)))
        headings[1] = ("Estado de la campaña" if kind == CAMPAIGN
                       else "Estado del grupo de anuncios")
        if kind == GROUP:
            headings.append("Código de moneda")
        rows = []
        for row in sheet.iter_rows(min_row=4, values_only=True):
            if str(row[1]).startswith("Total:"):
                continue
            values = list(row)
            if kind == GROUP:
                values.append("MXN")
            rows.append(values)
    finally:
        old.close()

    if mutate is not None:
        mutate(headings, rows)

    result = Workbook()
    ws = result.active
    ws.append(headings)
    for values in rows:
        ws.append(values)
    buf = BytesIO()
    result.save(buf)
    result.close()
    return buf.getvalue()


def test_legacy_titled_files_keep_official_total_reconciliation():
    campaign, group = daily_xlsx()
    a = parse_google_ads_daily_xlsx(campaign)
    b = parse_google_ads_ad_group_daily_xlsx(group)
    assert a.totals_provenance == b.totals_provenance == "GOOGLE_ADS_DAILY_TOTALS"
    assert a.date_from == b.date_from == date(2026, 10, 13)
    assert len(a.rows) == 27
    assert len(b.rows) == 43


def test_email_flat_parse_both_types_single_closed_day():
    a = parse_google_ads_daily_xlsx(flat_email(CAMPAIGN))
    b = parse_google_ads_ad_group_daily_xlsx(flat_email(GROUP))
    assert (a.date_from, a.date_to) == (date(2026, 10, 13),) * 2
    assert (b.date_from, b.date_to) == (date(2026, 10, 13),) * 2
    assert len(a.rows) == 27 and len(b.rows) == 43
    assert a.totals_provenance == b.totals_provenance == "DETAIL_SUM_ONLY"
    assert a.total_cost == b.total_cost


def test_m6_preview_accepts_emailed_pair_and_exposes_provenance(setup_client):
    client, token, _ = setup_client
    response = post(client, token, files=(flat_email(CAMPAIGN), flat_email(GROUP)))
    assert response.status_code == 200, response.get_json()
    data = response.get_json()
    assert data["status"] == "PREFLIGHT_OK"
    assert data["campaign_rows"] == 27 and data["group_rows"] == 43
    assert data["campaign_totals_provenance"] == "DETAIL_SUM_ONLY"
    assert data["group_totals_provenance"] == "DETAIL_SUM_ONLY"
    assert data["ready_to_stage_raw"] is True
    assert data["eligible_for_publication"] is False


@pytest.mark.parametrize("kind", [CAMPAIGN, GROUP])
def test_email_rejects_two_dates_without_unverified_period(kind):
    def mutate(_heads, rows):
        rows[0][0] = "2026-10-14"
    parser = (parse_google_ads_daily_xlsx if kind == CAMPAIGN
              else parse_google_ads_ad_group_daily_xlsx)
    with pytest.raises(GoogleAdsXlsxValidationError, match="exactamente un día"):
        parser(flat_email(kind, mutate=mutate))


@pytest.mark.parametrize("kind", [CAMPAIGN, GROUP])
def test_email_rejects_wrong_or_missing_google_report_headers(kind):
    def mutate(heads, _rows):
        heads[1] = "Estado"
    parser = (parse_google_ads_daily_xlsx if kind == CAMPAIGN
              else parse_google_ads_ad_group_daily_xlsx)
    with pytest.raises(GoogleAdsXlsxValidationError, match="No es un"):
        parser(flat_email(kind, mutate=mutate))


def test_email_requires_both_group_currency_columns():
    def mutate(headings, rows):
        headings.pop()
        for row in rows:
            row.pop()
    with pytest.raises(GoogleAdsXlsxValidationError, match="No es un"):
        parse_google_ads_ad_group_daily_xlsx(flat_email(GROUP, mutate=mutate))


@pytest.mark.parametrize("kind", [CAMPAIGN, GROUP])
def test_email_rejects_duplicate_identity(kind):
    def mutate(_headings, rows):
        rows.append(list(rows[0]))
    parser = (parse_google_ads_daily_xlsx if kind == CAMPAIGN
              else parse_google_ads_ad_group_daily_xlsx)
    with pytest.raises(GoogleAdsXlsxValidationError, match="duplicad"):
        parser(flat_email(kind, mutate=mutate))


@pytest.mark.parametrize("kind", [CAMPAIGN, GROUP])
def test_email_rejects_currency_other_than_mxn(kind):
    def mutate(_headings, rows):
        currency_column = 3 if kind == CAMPAIGN else 4
        rows[0][currency_column] = "USD"
    parser = (parse_google_ads_daily_xlsx if kind == CAMPAIGN
              else parse_google_ads_ad_group_daily_xlsx)
    with pytest.raises(GoogleAdsXlsxValidationError, match="MXN|moneda"):
        parser(flat_email(kind, mutate=mutate))


def test_flat_pair_spend_drift_fails_real_m6_preflight(setup_client):
    client, token, _ = setup_client
    def mutate(_headings, rows):
        rows[0][4] = float(rows[0][4]) + 1.0
    response = post(client, token, files=(
        flat_email(CAMPAIGN, mutate=mutate), flat_email(GROUP)
    ))
    assert response.status_code == 422, response.get_json()
    assert response.get_json()["code"] == "CAMPAIGN_GROUP_RECONCILIATION_FAILED"


def test_titled_file_without_official_daily_total_still_rejected():
    from openpyxl import load_workbook
    old, _ = daily_xlsx()
    workbook = load_workbook(BytesIO(old))
    sheet = workbook.active
    sheet.delete_rows(sheet.max_row)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    with pytest.raises(GoogleAdsXlsxValidationError, match="totales diarios"):
        parse_google_ads_daily_xlsx(buffer.getvalue())
