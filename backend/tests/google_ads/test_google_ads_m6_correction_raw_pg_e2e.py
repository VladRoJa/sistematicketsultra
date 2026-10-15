"""Real PG16/FS test: signed correction-only RAW paired intake.

All uploads are synthetic and stored under tmp_path. No production API, data or
files are used, and no staged/canonical facts may be written by this feature.
"""
from __future__ import annotations

from datetime import date
from io import BytesIO
from hashlib import sha256

import pytest
import sqlalchemy as sa
from openpyxl import load_workbook

from app.warehouse.services import google_ads_m6_correction_raw_service as svc
from app.warehouse.services import google_ads_m6_manual_confirm_service as ordinary
from app.warehouse.services.warehouse_document_upload_service import create_warehouse_document_upload
from test_google_ads_m6_manual_raw_postgres import real_warehouse, _confirm
from test_google_ads_m6_manual_preview import daily_xlsx, post as preview_post

PREVIEW = "/api/integrations/google-ads/corrections/raw/preview"
CONFIRM = "/api/integrations/google-ads/corrections/raw/confirm"


def _changed_files():
    # Change one campaign + one group + BOTH account footer rows, keeping the
    # original 27/43 identities and reconciliation exactly balanced.
    original = daily_xlsx()
    outputs = []
    for idx, raw in enumerate(original):
        book = load_workbook(BytesIO(raw))
        sheet = book.active
        cost_col = 5 if idx == 0 else 6
        for row in range(4, sheet.max_row + 1):
            if row == 4 or (
                str(sheet.cell(row, 2).value or "").startswith("Total:")
            ):
                cell = sheet.cell(row, cost_col)
                cell.value = round(float(cell.value) + 10.0, 2)
        result = BytesIO()
        book.save(result)
        book.close()
        outputs.append(result.getvalue())
    return tuple(outputs)


def _request(client, token, url, files, *, reason=None, ticket=None):
    data = {
        "campaign": (BytesIO(files[0]), "campaign.xlsx"),
        "ad_group": (BytesIO(files[1]), "groups.xlsx"),
    }
    if reason is not None:
        data["reason"] = reason
    if ticket is not None:
        data["correction_ticket"] = ticket
    return client.post(
        url, data=data, content_type="multipart/form-data",
        headers={"Authorization": "Bearer " + token} if token else {},
    )


@pytest.fixture
def lab(real_warehouse, monkeypatch):
    client, actor, other, engine, root = real_warehouse
    original = daily_xlsx()
    before = preview_post(client, actor, files=original)
    assert before.status_code == 200, before.get_json()
    prior = _confirm(client, actor, before.json["raw_confirm_ticket"], original)
    assert prior.status_code == 201, prior.get_json()
    # This test only requires read access to publication metadata. It creates
    # a shallow table to avoid touching full production schema.
    metadata = sa.MetaData()
    pub = sa.Table("google_ads_m6_daily_publications", metadata,
                   sa.Column("id", sa.BigInteger, primary_key=True),
                   sa.Column("customer_id", sa.String(10), nullable=False),
                   sa.Column("business_date", sa.Date, nullable=False),
                   sa.Column("revision", sa.Integer, nullable=False),
                   sa.Column("is_canonical", sa.Boolean, nullable=False))
    with engine.begin() as conn:
        assert not sa.inspect(conn).has_table(pub.name)
        pub.create(conn)
    monkeypatch.setenv(svc.FLAG, "true")
    try:
        yield client, actor, other, engine, pub, original, _changed_files()
    finally:
        with engine.begin() as conn:
            pub.drop(conn)


def _snap(engine):
    with engine.connect() as conn:
        rows = conn.execute(sa.text("""
            SELECT u.id, u.status, u.file_hash_sha256, a.details
            FROM warehouse_uploads u LEFT JOIN warehouse_audit_logs a
             ON a.upload_id = u.id AND a.action = 'UPLOAD' ORDER BY u.id
        """)).mappings().all()
        return [dict(r) for r in rows]


def test_exclusive_preview_and_pair_staging_without_funnel_writes(lab):
    client, actor, other, engine, pub, original, files = lab
    from test_google_ads_m6_manual_preview import DAY
    first = _snap(engine)
    assert len(first) == 2
    reason = "Correccion extraordinaria validada con ambos informes"
    preview = _request(client, actor, PREVIEW, files, reason=reason)
    assert preview.status_code == 200, preview.get_json()
    assert preview.json["status"] == "CORRECTION_PREFLIGHT_OK"
    assert preview.json["revision"] == 1
    assert preview.json["eligible_for_publication"] is False
    assert _snap(engine) == first  # preview is strictly read only

    out = _request(client, actor, CONFIRM, files,
                   ticket=preview.json["correction_ticket"])
    assert out.status_code == 201, out.get_json()
    assert out.json["status"] == "CORRECTION_RAW_PAIR_STAGED"
    assert out.json["eligible_for_structured_ingestion"] is False
    assert out.json["eligible_for_publication"] is False
    assert out.json["funnel_publication"] == "NOT_AUTHORIZED"
    rows = _snap(engine)
    assert len(rows) == 4
    corrections = [r for r in rows if isinstance(r["details"], dict)
                   and r["details"].get("upload_origin") == svc.ORIGIN]
    assert len(corrections) == 2
    assert all(r["details"]["publication"] == "NOT_AUTHORIZED" for r in corrections)
    assert all(r["details"]["correction_case_intent"] == "CANDIDATE_ONLY"
               for r in corrections)
    assert {r["file_hash_sha256"] for r in corrections} == {
        sha256(f).hexdigest() for f in files
    }
    again = _request(client, actor, CONFIRM, files,
                     ticket=preview.json["correction_ticket"])
    assert again.status_code == 200, again.get_json()
    assert all(v["reused_existing"] for v in again.json["uploads"].values())
    assert _snap(engine) == rows


def test_disabled_permissions_duplicate_and_signed_ticket(lab, monkeypatch):
    client, actor, other, engine, pub, original, files = lab
    baseline = _snap(engine)
    monkeypatch.delenv(svc.FLAG)
    assert _request(client, actor, PREVIEW, files, reason="Motivo de correccion").status_code == 503
    monkeypatch.setenv(svc.FLAG, "true")
    assert _request(client, None, PREVIEW, files, reason="Motivo de correccion").status_code == 401
    assert _request(client, other, PREVIEW, files, reason="Motivo de correccion").status_code == 403
    # Reusing a normal original RAW would bypass the correction audit.
    same = _request(client, actor, PREVIEW, original, reason="Motivo de correccion")
    assert same.status_code == 409
    assert same.json["code"] == "CORRECTION_DUPLICATE_RAW_CONFLICT"
    p = _request(client, actor, PREVIEW, files, reason="Motivo de correccion")
    assert p.status_code == 200, p.get_json()
    signed = p.json["correction_ticket"]
    bad = _request(client, actor, CONFIRM, files, ticket=signed + "tamper")
    assert bad.status_code == 409
    changed = _request(client, actor, CONFIRM, original, ticket=signed)
    assert changed.status_code == 409
    assert _snap(engine) == baseline
    with engine.begin() as conn:
        conn.execute(sa.text("""
            UPDATE warehouse_operators SET can_upload = false WHERE user_id = 47
        """))
    assert _request(client, actor, PREVIEW, files, reason="Motivo de correccion").status_code == 403


def test_partial_second_raw_can_be_resumed_without_new_first_audit(lab, monkeypatch):
    client, actor, other, engine, pub, original, files = lab
    p = _request(client, actor, PREVIEW, files, reason="Motivo de correccion")
    assert p.status_code == 200
    def fail_group(*, report_type_key, **kwargs):
        if report_type_key == svc.REPORT_TYPES[1]:
            raise RuntimeError("synthetic second file IO failed")
        return create_warehouse_document_upload(report_type_key=report_type_key, **kwargs)

    monkeypatch.setattr(svc, "create_warehouse_document_upload", fail_group)
    partial = _request(client, actor, CONFIRM, files,
                       ticket=p.json["correction_ticket"])
    assert partial.status_code == 409
    assert partial.json["status"] == "PARTIAL_RAW"
    assert len(_snap(engine)) == 3
    monkeypatch.setattr(svc, "create_warehouse_document_upload",
                        create_warehouse_document_upload)
    resumed = _request(client, actor, CONFIRM, files,
                       ticket=p.json["correction_ticket"])
    assert resumed.status_code == 201, resumed.get_json()
    assert resumed.json["uploads"][svc.REPORT_TYPES[0]]["reused_existing"] is True
    assert resumed.json["uploads"][svc.REPORT_TYPES[1]]["reused_existing"] is False
    assert len(_snap(engine)) == 4


def test_archived_old_evidence_and_canonical_version_binding(lab):
    client, actor, other, engine, pub, original, files = lab
    with engine.begin() as conn:
        conn.execute(sa.text("UPDATE warehouse_uploads SET status = 'ARCHIVED'"))
        conn.execute(pub.insert(), {
            "id": 901, "customer_id": svc.CUSTOMER_ID,
            "business_date": date(2026, 10, 13),
            "revision": 1, "is_canonical": True,
        })
    p = _request(client, actor, PREVIEW, files,
                 reason="Motivo historico con fuentes archivadas")
    assert p.status_code == 200, p.get_json()
    assert p.json["revision"] == 2
    assert p.json["expected_publication_id"] == 901
    with engine.begin() as conn:
        conn.execute(sa.text("""
            UPDATE google_ads_m6_daily_publications
            SET revision = 2 WHERE id = 901
        """))
    stale = _request(client, actor, CONFIRM, files,
                     ticket=p.json["correction_ticket"])
    assert stale.status_code == 409
    assert stale.json["code"] == "CORRECTION_STATE_CHANGED"
    assert len(_snap(engine)) == 2
