from datetime import date

import pytest
from openpyxl import load_workbook

from app.sports_analysis import (
    attendance_base_health_service as health,
)
from app.sports_analysis.attendance_access import (
    SportsAnalysisScope,
)
from app.sports_analysis.attendance_query_service import (
    SportsAnalysisValidationError,
)


def _scope() -> SportsAnalysisScope:
    return SportsAnalysisScope(
        role="ADMINISTRADOR",
        is_global=True,
        allowed_branch_ids=(1, 2, 3),
        fixed_branch_id=None,
    )


def _member(
    id_socio: str,
    *,
    branch_id: int = 1,
    branch_name: str = "Sucursal Uno",
    name: str | None = None,
    member_since: date | None = date(2026, 1, 1),
    expiration_date: date = date(2026, 12, 31),
    applies_kpi: bool = True,
    tariff: str | None = None,
    phone: str | None = None,
) -> health._EligibleMember:
    return health._EligibleMember(
        id_socio=id_socio,
        pin=f"PIN-{id_socio}",
        name=name or f"Socio {id_socio}",
        branch_id=branch_id,
        branch_name=branch_name,
        member_since=member_since,
        expiration_date=expiration_date,
        snapshot_date=date(2026, 9, 25),
        applies_kpi=applies_kpi,
        tariff=tariff,
        phone=phone,
    )


def _state() -> health._BaseHealthState:
    return health._BaseHealthState(
        eligible_members={
            "100": _member("100"),
            "200": _member(
                "200",
                branch_id=2,
                branch_name="Sucursal Dos",
            ),
            "300": _member("300"),
            "400": _member("400"),
        },
        visited_member_ids=frozenset(
            {"100", "300", "999"}
        ),
        snapshot_dates=(
            date(2026, 9, 1),
            date(2026, 9, 25),
        ),
        total_visits_checked=100,
        resolved_visits=96,
    )


def test_base_health_utilization_uses_eligible_members(
    monkeypatch,
):
    monkeypatch.setattr(
        health,
        "_build_state",
        lambda *_args, **_kwargs: (
            _state(),
            (1, 2),
        ),
    )

    result = health.attendance_base_health(
        _scope(),
        date_from=date(2026, 9, 1),
        date_to=date(2026, 9, 25),
        branch_id=None,
        region_key=None,
    )

    assert result["summary"] == {
        "eligible_members": 4,
        "members_with_visit": 2,
        "members_without_visit": 2,
        "utilization_pct": 50.0,
        "without_visit_pct": 50.0,
        "frequency_avg_per_week": 0.0,
        "frequency_median_per_week": 0.0,
        "members_14_plus_days_without_visit": 0,
        "follow_up_members": 4,
        "members_less_than_one_visit_per_week": 0,
    }
    assert (
        result["source"]["identity_coverage_pct"]
        == 96.0
    )
    assert result["source"]["snapshot_count"] == 2
    assert (
        result["filters"]["attendance_type"]
        == "SOCIO"
    )


def test_base_health_ignores_resolved_visits_outside_base(
    monkeypatch,
):
    state = _state()
    monkeypatch.setattr(
        health,
        "_build_state",
        lambda *_args, **_kwargs: (
            state,
            (1, 2),
        ),
    )

    result = health.attendance_base_health(
        _scope(),
        date_from=date(2026, 9, 1),
        date_to=date(2026, 9, 25),
        branch_id=1,
        region_key=None,
    )

    assert result["summary"]["members_with_visit"] == 2
    assert result["summary"]["eligible_members"] == 4


def test_base_health_empty_source_is_safe(
    monkeypatch,
):
    empty = health._BaseHealthState(
        eligible_members={},
        visited_member_ids=frozenset(),
        snapshot_dates=(),
        total_visits_checked=0,
        resolved_visits=0,
    )
    monkeypatch.setattr(
        health,
        "_build_state",
        lambda *_args, **_kwargs: (
            empty,
            (1,),
        ),
    )

    result = health.attendance_base_health(
        _scope(),
        date_from=date(2026, 9, 1),
        date_to=date(2026, 9, 25),
        branch_id=1,
        region_key=None,
    )

    assert result["summary"]["eligible_members"] == 0
    assert result["summary"]["utilization_pct"] == 0.0
    assert result["source"]["available"] is False


def test_member_detail_splits_with_and_without_visit(
    monkeypatch,
):
    monkeypatch.setattr(
        health,
        "_build_state",
        lambda *_args, **_kwargs: (
            _state(),
            (1, 2),
        ),
    )

    with_visit = (
        health.attendance_base_health_members(
            _scope(),
            date_from=date(2026, 9, 1),
            date_to=date(2026, 9, 25),
            branch_id=None,
            region_key=None,
            status="WITH_VISIT",
            page=1,
            page_size=50,
        )
    )
    without_visit = (
        health.attendance_base_health_members(
            _scope(),
            date_from=date(2026, 9, 1),
            date_to=date(2026, 9, 25),
            branch_id=None,
            region_key=None,
            status="WITHOUT_VISIT",
            page=1,
            page_size=50,
        )
    )

    assert with_visit["count"] == 2
    assert {
        row["id_socio"]
        for row in with_visit["rows"]
    } == {"100", "300"}
    assert all(
        row["has_visit"] is True
        for row in with_visit["rows"]
    )

    assert without_visit["count"] == 2
    assert {
        row["id_socio"]
        for row in without_visit["rows"]
    } == {"200", "400"}
    assert all(
        row["has_visit"] is False
        for row in without_visit["rows"]
    )


def test_member_detail_rejects_unknown_status(
    monkeypatch,
):
    monkeypatch.setattr(
        health,
        "_build_state",
        lambda *_args, **_kwargs: (
            _state(),
            (1,),
        ),
    )

    with pytest.raises(
        SportsAnalysisValidationError
    ):
        health.attendance_base_health_members(
            _scope(),
            date_from=date(2026, 9, 1),
            date_to=date(2026, 9, 25),
            branch_id=None,
            region_key=None,
            status="ALL",
        )


def test_record_member_visits_accumulates_count_and_dates():
    stats: dict[str, health._MemberVisitStats] = {}

    health._record_member_visits(
        stats,
        id_socio="100",
        visit_count=2,
        business_date=date(2026, 9, 10),
    )
    health._record_member_visits(
        stats,
        id_socio="100",
        visit_count=3,
        business_date=date(2026, 9, 5),
    )
    health._record_member_visits(
        stats,
        id_socio="100",
        visit_count=1,
        business_date=date(2026, 9, 20),
    )

    assert stats["100"] == health._MemberVisitStats(
        visit_count=6,
        first_visit_date=date(2026, 9, 5),
        last_visit_date=date(2026, 9, 20),
    )


def test_frequency_uses_real_eligible_days():
    member = _member(
        "100",
        member_since=date(2026, 9, 15),
        expiration_date=date(2026, 9, 21),
    )

    days = health._member_eligible_days(
        member,
        date_from=date(2026, 9, 1),
        date_to=date(2026, 9, 30),
    )

    frequency = health._member_weekly_frequency(
        member,
        health._MemberVisitStats(
            visit_count=2,
            first_visit_date=date(2026, 9, 15),
            last_visit_date=date(2026, 9, 20),
        ),
        date_from=date(2026, 9, 1),
        date_to=date(2026, 9, 30),
    )

    assert days == 7
    assert frequency == 2.0


def test_frequency_metrics_include_zero_and_partial_members():
    state = health._BaseHealthState(
        eligible_members={
            "100": _member(
                "100",
                member_since=date(2026, 9, 1),
                expiration_date=date(2026, 9, 7),
            ),
            "200": _member(
                "200",
                member_since=date(2026, 9, 1),
                expiration_date=date(2026, 9, 14),
            ),
            "300": _member(
                "300",
                member_since=date(2026, 9, 1),
                expiration_date=date(2026, 9, 7),
            ),
        },
        visited_member_ids=frozenset({"100", "200"}),
        snapshot_dates=(date(2026, 9, 1),),
        total_visits_checked=4,
        resolved_visits=4,
        member_visit_stats={
            "100": health._MemberVisitStats(
                visit_count=2,
                first_visit_date=date(2026, 9, 1),
                last_visit_date=date(2026, 9, 5),
            ),
            "200": health._MemberVisitStats(
                visit_count=1,
                first_visit_date=date(2026, 9, 2),
                last_visit_date=date(2026, 9, 2),
            ),
        },
    )

    metrics = health._frequency_metrics(
        state,
        date_from=date(2026, 9, 1),
        date_to=date(2026, 9, 30),
    )

    assert metrics["average_per_week"] == 0.83
    assert metrics["median_per_week"] == 0.5
    assert (
        metrics["less_than_one_per_week_count"]
        == 1
    )
    assert [
        row["count"]
        for row in metrics["distribution"]
    ] == [1, 1, 0, 1, 0]


def test_recency_uses_date_to_and_known_last_visit():
    state = health._BaseHealthState(
        eligible_members={
            "100": _member("100"),
            "200": _member("200"),
            "300": _member("300"),
            "400": _member("400"),
            "500": _member("500"),
        },
        visited_member_ids=frozenset(
            {"100", "200", "300", "400"}
        ),
        snapshot_dates=(date(2026, 9, 1),),
        total_visits_checked=4,
        resolved_visits=4,
        member_visit_stats={},
        last_known_visit_by_member={
            "100": date(2026, 9, 25),
            "200": date(2026, 9, 17),
            "300": date(2026, 9, 11),
            "400": date(2026, 9, 1),
        },
    )

    result = health._recency_metrics(
        state,
        date_to=date(2026, 9, 25),
    )

    assert (
        result[
            "members_14_plus_days_without_visit"
        ]
        == 2
    )
    assert [
        row["count"]
        for row in result["distribution"]
    ] == [1, 2, 0, 1, 1]


def test_recency_exactly_14_days_counts_in_14_plus():
    state = health._BaseHealthState(
        eligible_members={
            "100": _member("100"),
        },
        visited_member_ids=frozenset(),
        snapshot_dates=(date(2026, 9, 1),),
        total_visits_checked=0,
        resolved_visits=0,
        last_known_visit_by_member={
            "100": date(2026, 9, 11),
        },
    )

    result = health._recency_metrics(
        state,
        date_to=date(2026, 9, 25),
    )

    assert (
        result[
            "members_14_plus_days_without_visit"
        ]
        == 1
    )
    assert result["distribution"][1]["count"] == 1


def test_historical_recency_skips_when_history_is_unavailable(
    monkeypatch,
):
    called = False

    def _unexpected_query(*_args, **_kwargs):
        nonlocal called
        called = True
        raise AssertionError(
            "No debe consultar histórico sin datos previos."
        )

    monkeypatch.setattr(
        health,
        "_has_historical_attendance_before",
        lambda *_args, **_kwargs: False,
    )
    monkeypatch.setattr(
        health.db.session,
        "query",
        _unexpected_query,
    )

    result = health._load_historical_last_visits(
        {
            "100": _member("100"),
        },
        {},
        date_before=date(2026, 9, 1),
    )

    assert result == {}
    assert called is False


def test_activation_classifies_new_members():
    state = health._BaseHealthState(
        eligible_members={
            "100": _member(
                "100",
                member_since=date(2026, 9, 1),
            ),
            "200": _member(
                "200",
                member_since=date(2026, 9, 1),
            ),
            "300": _member(
                "300",
                member_since=date(2026, 9, 1),
            ),
            "400": _member(
                "400",
                member_since=date(2026, 9, 1),
            ),
            "500": _member(
                "500",
                member_since=date(2026, 9, 1),
            ),
            "600": _member(
                "600",
                member_since=date(2026, 9, 20),
            ),
            "700": _member(
                "700",
                member_since=date(2026, 8, 20),
            ),
        },
        visited_member_ids=frozenset(
            {"100", "200", "300", "400", "700"}
        ),
        snapshot_dates=(date(2026, 9, 1),),
        total_visits_checked=5,
        resolved_visits=5,
        member_visit_stats={
            "100": health._MemberVisitStats(
                visit_count=1,
                first_visit_date=date(2026, 9, 1),
                last_visit_date=date(2026, 9, 1),
            ),
            "200": health._MemberVisitStats(
                visit_count=1,
                first_visit_date=date(2026, 9, 3),
                last_visit_date=date(2026, 9, 3),
            ),
            "300": health._MemberVisitStats(
                visit_count=1,
                first_visit_date=date(2026, 9, 6),
                last_visit_date=date(2026, 9, 6),
            ),
            "400": health._MemberVisitStats(
                visit_count=1,
                first_visit_date=date(2026, 9, 10),
                last_visit_date=date(2026, 9, 10),
            ),
            "700": health._MemberVisitStats(
                visit_count=1,
                first_visit_date=date(2026, 9, 2),
                last_visit_date=date(2026, 9, 2),
            ),
        },
    )

    result = health._activation_metrics(
        state,
        date_from=date(2026, 9, 1),
        date_to=date(2026, 9, 25),
    )

    assert result["new_members"] == 6
    assert [
        row["count"]
        for row in result["distribution"]
    ] == [1, 1, 1, 1, 1, 1]


def test_activation_exactly_seven_days_is_not_pending():
    state = health._BaseHealthState(
        eligible_members={
            "100": _member(
                "100",
                member_since=date(2026, 9, 18),
            ),
        },
        visited_member_ids=frozenset(),
        snapshot_dates=(date(2026, 9, 18),),
        total_visits_checked=0,
        resolved_visits=0,
    )

    result = health._activation_metrics(
        state,
        date_from=date(2026, 9, 1),
        date_to=date(2026, 9, 25),
    )

    assert result["distribution"][4]["count"] == 1
    assert (
        result["distribution"][4]["label"]
        == "7+ días sin primera visita"
    )
    assert result["distribution"][5]["count"] == 0
    assert (
        result["distribution"][5]["label"]
        == "Alta reciente, aún sin visita"
    )


def test_follow_up_uses_recency_or_low_frequency_once():
    state = health._BaseHealthState(
        eligible_members={
            "100": _member(
                "100",
                expiration_date=date(2026, 12, 31),
            ),
            "200": _member(
                "200",
                expiration_date=date(2026, 12, 31),
            ),
            "300": _member(
                "300",
                expiration_date=date(2026, 12, 31),
            ),
            "400": _member(
                "400",
                expiration_date=date(2026, 9, 20),
            ),
        },
        visited_member_ids=frozenset(
            {"100", "200", "300"}
        ),
        snapshot_dates=(date(2026, 9, 1),),
        total_visits_checked=6,
        resolved_visits=6,
        member_visit_stats={
            "100": health._MemberVisitStats(
                visit_count=4,
                first_visit_date=date(2026, 9, 1),
                last_visit_date=date(2026, 9, 5),
            ),
            "200": health._MemberVisitStats(
                visit_count=1,
                first_visit_date=date(2026, 9, 20),
                last_visit_date=date(2026, 9, 20),
            ),
            "300": health._MemberVisitStats(
                visit_count=1,
                first_visit_date=date(2026, 9, 1),
                last_visit_date=date(2026, 9, 1),
            ),
        },
        last_known_visit_by_member={
            "100": date(2026, 9, 5),
            "200": date(2026, 9, 20),
            "300": date(2026, 9, 1),
        },
    )

    result = health._follow_up_metrics(
        state,
        date_from=date(2026, 9, 1),
        date_to=date(2026, 9, 25),
    )

    assert result["count"] == 3
    assert result["member_ids"] == frozenset(
        {"100", "200", "300"}
    )
    assert result["by_recency"] == 2
    assert result["by_low_frequency"] == 2


def test_follow_up_excludes_member_not_active_at_date_to():
    state = health._BaseHealthState(
        eligible_members={
            "100": _member(
                "100",
                expiration_date=date(2026, 9, 20),
            ),
        },
        visited_member_ids=frozenset(),
        snapshot_dates=(date(2026, 9, 1),),
        total_visits_checked=0,
        resolved_visits=0,
    )

    result = health._follow_up_metrics(
        state,
        date_from=date(2026, 9, 1),
        date_to=date(2026, 9, 25),
    )

    assert result["count"] == 0
    assert result["member_ids"] == frozenset()


def test_member_detail_supports_follow_up_with_behavior_fields(
    monkeypatch,
):
    state = health._BaseHealthState(
        eligible_members={
            "100": _member("100"),
            "200": _member("200"),
        },
        visited_member_ids=frozenset({"100", "200"}),
        snapshot_dates=(date(2026, 9, 1),),
        total_visits_checked=3,
        resolved_visits=3,
        member_visit_stats={
            "100": health._MemberVisitStats(
                visit_count=1,
                first_visit_date=date(2026, 9, 1),
                last_visit_date=date(2026, 9, 1),
            ),
            "200": health._MemberVisitStats(
                visit_count=2,
                first_visit_date=date(2026, 9, 20),
                last_visit_date=date(2026, 9, 25),
            ),
        },
        last_known_visit_by_member={
            "100": date(2026, 9, 1),
            "200": date(2026, 9, 25),
        },
    )

    monkeypatch.setattr(
        health,
        "_build_state",
        lambda *_args, **_kwargs: (
            state,
            (1,),
        ),
    )

    result = health.attendance_base_health_members(
        _scope(),
        date_from=date(2026, 9, 1),
        date_to=date(2026, 9, 25),
        branch_id=None,
        region_key=None,
        status="FOLLOW_UP",
    )

    assert result["count"] == 2
    row = next(
        item
        for item in result["rows"]
        if item["id_socio"] == "100"
    )
    assert row["visit_count"] == 1
    assert row["last_visit_date"] == "2026-09-01"
    assert row["days_since_last_visit"] == 24
    assert row["frequency_per_week"] == 0.28


def test_member_detail_supports_recency_and_frequency_buckets(
    monkeypatch,
):
    state = health._BaseHealthState(
        eligible_members={
            "100": _member("100"),
            "200": _member("200"),
            "300": _member("300"),
        },
        visited_member_ids=frozenset({"100", "200"}),
        snapshot_dates=(date(2026, 9, 1),),
        total_visits_checked=4,
        resolved_visits=4,
        member_visit_stats={
            "100": health._MemberVisitStats(
                visit_count=1,
                first_visit_date=date(2026, 9, 11),
                last_visit_date=date(2026, 9, 11),
            ),
            "200": health._MemberVisitStats(
                visit_count=11,
                first_visit_date=date(2026, 9, 1),
                last_visit_date=date(2026, 9, 25),
            ),
        },
        last_known_visit_by_member={
            "100": date(2026, 9, 11),
            "200": date(2026, 9, 25),
        },
    )

    monkeypatch.setattr(
        health,
        "_build_state",
        lambda *_args, **_kwargs: (
            state,
            (1,),
        ),
    )

    recency = health.attendance_base_health_members(
        _scope(),
        date_from=date(2026, 9, 1),
        date_to=date(2026, 9, 25),
        branch_id=None,
        region_key=None,
        status="RECENCY_8_14",
    )
    frequency = health.attendance_base_health_members(
        _scope(),
        date_from=date(2026, 9, 1),
        date_to=date(2026, 9, 25),
        branch_id=None,
        region_key=None,
        status="FREQUENCY_GTE_3",
    )
    no_recorded = health.attendance_base_health_members(
        _scope(),
        date_from=date(2026, 9, 1),
        date_to=date(2026, 9, 25),
        branch_id=None,
        region_key=None,
        status="RECENCY_NO_RECORDED",
    )

    assert [
        row["id_socio"]
        for row in recency["rows"]
    ] == ["100"]
    assert [
        row["id_socio"]
        for row in frequency["rows"]
    ] == ["200"]
    assert [
        row["id_socio"]
        for row in no_recorded["rows"]
    ] == ["300"]



def test_base_health_keeps_non_kpi_members_separate(
    monkeypatch,
):
    state = health._BaseHealthState(
        eligible_members={
            "100": _member("100"),
            "200": _member("200"),
        },
        visited_member_ids=frozenset(
            {"100", "900"}
        ),
        snapshot_dates=(date(2026, 9, 25),),
        total_visits_checked=2,
        resolved_visits=2,
        other_active_members={
            "900": _member(
                "900",
                applies_kpi=False,
                tariff="SEMANA $299",
            ),
            "901": _member(
                "901",
                applies_kpi=False,
                tariff="BECA 6 MESES",
            ),
            "902": _member(
                "902",
                applies_kpi=False,
                tariff="SEMANA $299",
            ),
        },
    )

    monkeypatch.setattr(
        health,
        "_build_state",
        lambda *_args, **_kwargs: (
            state,
            (1,),
        ),
    )

    result = health.attendance_base_health(
        _scope(),
        date_from=date(2026, 9, 1),
        date_to=date(2026, 9, 25),
        branch_id=None,
        region_key=None,
    )

    assert result["summary"]["eligible_members"] == 2
    assert result["summary"]["members_with_visit"] == 1
    assert (
        result["other_active_accesses"]["count"]
        == 3
    )
    assert (
        result["other_active_accesses"][
            "tariff_distribution"
        ]
        == [
            {
                "tariff": "SEMANA $299",
                "count": 2,
            },
            {
                "tariff": "BECA 6 MESES",
                "count": 1,
            },
        ]
    )



def test_other_active_access_detail_uses_non_kpi_universe(
    monkeypatch,
):
    state = health._BaseHealthState(
        eligible_members={
            "100": _member("100"),
        },
        visited_member_ids=frozenset(),
        snapshot_dates=(date(2026, 9, 25),),
        total_visits_checked=0,
        resolved_visits=0,
        other_active_members={
            "900": _member(
                "900",
                applies_kpi=False,
                tariff="PASE GYMPASS",
            ),
            "901": _member(
                "901",
                applies_kpi=False,
                tariff="SEMANA $299",
            ),
        },
    )

    monkeypatch.setattr(
        health,
        "_build_state",
        lambda *_args, **_kwargs: (
            state,
            (1,),
        ),
    )

    result = health.attendance_base_health_members(
        _scope(),
        date_from=date(2026, 9, 1),
        date_to=date(2026, 9, 25),
        branch_id=None,
        region_key=None,
        status="OTHER_ACTIVE_ACCESS",
        page=1,
        page_size=50,
    )

    assert result["count"] == 2
    assert result["title"] == (
        "Otros accesos activos en el periodo"
    )
    assert {
        row["id_socio"]
        for row in result["rows"]
    } == {"900", "901"}
    assert {
        row["tariff"]
        for row in result["rows"]
    } == {"PASE GYMPASS", "SEMANA $299"}



def test_base_health_export_includes_phone_without_exposing_it_in_json(
    monkeypatch,
):
    state = health._BaseHealthState(
        eligible_members={
            "100": _member(
                "100",
                phone="6861111111",
            ),
            "200": _member(
                "200",
                phone="6862222222",
            ),
        },
        visited_member_ids=frozenset(
            {"100", "200"}
        ),
        snapshot_dates=(date(2026, 9, 25),),
        total_visits_checked=2,
        resolved_visits=2,
    )

    monkeypatch.setattr(
        health,
        "_build_state",
        lambda *_args, **_kwargs: (
            state,
            (1,),
        ),
    )

    detail = health.attendance_base_health_members(
        _scope(),
        date_from=date(2026, 9, 1),
        date_to=date(2026, 9, 25),
        branch_id=None,
        region_key=None,
        status="WITH_VISIT",
        page=1,
        page_size=1,
    )

    assert detail["count"] == 2
    assert len(detail["rows"]) == 1
    assert "phone" not in detail["rows"][0]

    output, filename = (
        health.attendance_base_health_members_export(
            _scope(),
            date_from=date(2026, 9, 1),
            date_to=date(2026, 9, 25),
            branch_id=None,
            region_key=None,
            status="WITH_VISIT",
        )
    )

    workbook = load_workbook(output)
    worksheet = workbook.active
    rows = list(
        worksheet.iter_rows(values_only=True)
    )

    assert filename == (
        "salud_base_with_visit_"
        "2026-09-01_2026-09-25.xlsx"
    )
    assert rows[0] == (
        "Socio",
        "ID socio",
        "PIN",
        "Teléfono",
        "Sucursal",
        "Alta",
        "Visitas",
        "Última visita",
        "Días sin venir",
        "Frecuencia/semana",
    )
    assert len(rows) == 3
    assert {
        row[3]
        for row in rows[1:]
    } == {
        "6861111111",
        "6862222222",
    }


def test_other_active_access_export_includes_tariff_and_phone(
    monkeypatch,
):
    state = health._BaseHealthState(
        eligible_members={
            "100": _member("100"),
        },
        visited_member_ids=frozenset(),
        snapshot_dates=(date(2026, 9, 25),),
        total_visits_checked=0,
        resolved_visits=0,
        other_active_members={
            "900": _member(
                "900",
                applies_kpi=False,
                tariff="PASE GYMPASS",
                phone="6869000000",
            ),
        },
    )

    monkeypatch.setattr(
        health,
        "_build_state",
        lambda *_args, **_kwargs: (
            state,
            (1,),
        ),
    )

    output, _ = (
        health.attendance_base_health_members_export(
            _scope(),
            date_from=date(2026, 9, 1),
            date_to=date(2026, 9, 25),
            branch_id=None,
            region_key=None,
            status="OTHER_ACTIVE_ACCESS",
        )
    )

    workbook = load_workbook(output)
    worksheet = workbook.active
    rows = list(
        worksheet.iter_rows(values_only=True)
    )

    assert rows[0] == (
        "Socio",
        "ID socio",
        "PIN",
        "Teléfono",
        "Sucursal",
        "Alta",
        "Tarifa",
    )
    assert rows[1][3] == "6869000000"
    assert rows[1][6] == "PASE GYMPASS"



def test_base_health_detail_records_server_timing_stages(
    monkeypatch,
):
    member = _member("100")
    state = health._BaseHealthState(
        eligible_members={"100": member},
        visited_member_ids=frozenset({"100"}),
        snapshot_dates=(date(2026, 9, 25),),
        total_visits_checked=1,
        resolved_visits=1,
        member_visit_stats={
            "100": health._MemberVisitStats(
                visit_count=1,
                first_visit_date=date(2026, 9, 10),
                last_visit_date=date(2026, 9, 10),
            ),
        },
        last_known_visit_by_member={
            "100": date(2026, 9, 10),
        },
    )

    monkeypatch.setattr(
        health,
        "scoped_branch_catalog",
        lambda _scope: (),
    )
    monkeypatch.setattr(
        health,
        "_effective_branch_ids",
        lambda *_args, **_kwargs: (1,),
    )
    monkeypatch.setattr(
        health,
        "_load_eligible_members",
        lambda *_args, **_kwargs: (
            {"100": member},
            {},
            (date(2026, 9, 25),),
        ),
    )
    monkeypatch.setattr(
        health,
        "_load_member_visit_stats",
        lambda *_args, **_kwargs: (
            state.member_visit_stats,
            1,
            1,
        ),
    )
    monkeypatch.setattr(
        health,
        "_load_historical_last_visits",
        lambda *_args, **_kwargs: {},
    )

    timings: dict[str, float] = {}
    result = health.attendance_base_health_members(
        _scope(),
        date_from=date(2026, 9, 1),
        date_to=date(2026, 9, 25),
        branch_id=None,
        region_key=None,
        status="WITH_VISIT",
        page=1,
        page_size=50,
        timings=timings,
    )

    assert result["count"] == 1
    assert set(timings) == {
        "scope",
        "socios_activos",
        "visitas_periodo",
        "visitas_consulta",
        "visitas_identidad",
        "visitas_fallback",
        "historico",
        "historico_identidad",
        "historico_visitas",
        "cohorte",
        "serializacion",
        "total",
    }
    assert all(
        value >= 0
        for value in timings.values()
    )



def test_base_health_dashboard_records_server_timing(
    monkeypatch,
):
    state = _state()

    def fake_build_state(*_args, **kwargs):
        timings = kwargs.get("timings")
        if timings is not None:
            timings.update(
                {
                    "scope": 1.0,
                    "socios_activos": 2.0,
                    "visitas_periodo": 3.0,
                    "visitas_consulta": 1.0,
                    "visitas_identidad": 1.0,
                    "visitas_fallback": 1.0,
                    "historico": 4.0,
                    "historico_identidad": 1.0,
                    "historico_visitas": 3.0,
                }
            )
        return state, (1, 2)

    monkeypatch.setattr(
        health,
        "_build_state",
        fake_build_state,
    )

    timings: dict[str, float] = {}
    result = health.attendance_base_health(
        _scope(),
        date_from=date(2026, 9, 1),
        date_to=date(2026, 9, 25),
        branch_id=None,
        region_key=None,
        timings=timings,
    )

    assert result["summary"]["eligible_members"] == 4
    assert "metricas" in timings
    assert "total" in timings
    assert timings["metricas"] >= 0
    assert timings["total"] >= timings["metricas"]



def test_historical_last_visit_query_uses_distinct_on_latest(
    monkeypatch,
):
    class FakeQuery:
        def __init__(self):
            self.join_target = None
            self.filter_args = ()
            self.distinct_args = ()
            self.order_args = ()

        def join(
            self,
            target,
            *_args,
            **_kwargs,
        ):
            self.join_target = target
            return self

        def filter(self, *args):
            self.filter_args = args
            return self

        def distinct(self, *args):
            self.distinct_args = args
            return self

        def order_by(self, *args):
            self.order_args = args
            return self

        def all(self):
            return [
                (
                    "00123",
                    date(2025, 1, 1),
                    date(2026, 8, 31),
                ),
            ]

    query = FakeQuery()
    monkeypatch.setattr(
        health.db.session,
        "query",
        lambda *_args: query,
    )

    rows = health._load_historical_last_visit_rows(
        [
            (
                "00123",
                date(2025, 1, 1),
            ),
        ],
        date_before=date(2026, 9, 1),
    )

    assert rows[0][2] == date(2026, 8, 31)
    assert (
        getattr(
            query.join_target,
            "name",
            None,
        )
        == "requested_keys"
    )
    assert len(query.filter_args) == 2
    assert len(query.distinct_args) == 2
    assert "member_pin" in str(
        query.distinct_args[0]
    )
    assert "member_since" in str(
        query.distinct_args[1]
    )
    assert len(query.order_args) == 3
    assert "business_date DESC" in str(
        query.order_args[2]
    )



def test_historical_visit_batches_record_individual_timings(
    monkeypatch,
):
    class FakeQuery:
        def join(self, *_args, **_kwargs):
            return self

        def filter(self, *_args, **_kwargs):
            return self

        def group_by(self, *_args, **_kwargs):
            return self

        def all(self):
            return [
                (
                    "PIN-100",
                    date(2026, 1, 1),
                    1,
                    "100",
                ),
                (
                    "PIN-200",
                    date(2026, 1, 1),
                    1,
                    "200",
                ),
            ]

    monkeypatch.setattr(
        health,
        "_has_historical_attendance_before",
        lambda *_args, **_kwargs: True,
    )
    monkeypatch.setattr(
        health,
        "IDENTITY_BATCH_SIZE",
        1,
    )
    monkeypatch.setattr(
        health.db.session,
        "query",
        lambda *_args, **_kwargs: FakeQuery(),
    )

    def fake_last_visit_rows(
        keys,
        *,
        date_before,
    ):
        del date_before
        pin, member_since = keys[0]
        visit_date = (
            date(2026, 8, 20)
            if pin == "PIN-100"
            else date(2026, 8, 22)
        )
        return [
            (
                pin,
                member_since,
                visit_date,
            ),
        ]

    monkeypatch.setattr(
        health,
        "_load_historical_last_visit_rows",
        fake_last_visit_rows,
    )

    timings: dict[str, float] = {}
    result = health._load_historical_last_visits(
        {
            "100": _member("100"),
            "200": _member("200"),
        },
        {},
        date_before=date(2026, 9, 1),
        timings=timings,
    )

    assert result == {
        "100": date(2026, 8, 20),
        "200": date(2026, 8, 22),
    }
    assert "hist_q1" in timings
    assert "hist_q2" in timings
    assert "hist_q3" not in timings
    assert timings["hist_q1"] >= 0
    assert timings["hist_q2"] >= 0
