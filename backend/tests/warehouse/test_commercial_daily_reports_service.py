from datetime import date, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.warehouse.services.commercial_daily_reports_service import (
    CommercialDailyReportError,
    _daily_deltas,
    _week_start,
    render_commercial_daily_xlsx,
)


def test_week_starts_on_sunday():
    assert _week_start(date(2026, 10, 7)) == date(2026, 10, 4)
    assert _week_start(date(2026, 10, 3)) == date(2026, 9, 27)
    assert _week_start(date(2026, 10, 4)) == date(2026, 10, 4)


def test_mtd_month_boundary():
    assert _daily_deltas({
        date(2026, 9, 29): 100,
        date(2026, 9, 30): 106,
        date(2026, 10, 1): 4,
        date(2026, 10, 2): 8,
    }) == {
        date(2026, 9, 29): None,
        date(2026, 9, 30): 6,
        date(2026, 10, 1): 4,
        date(2026, 10, 2): 4,
    }


def test_mtd_gap_not_interpreted_as_zero():
    result = _daily_deltas({
        date(2026, 10, 1): 5,
        date(2026, 10, 3): 9,
    })
    assert result[date(2026, 10, 3)] is None


def test_mtd_decrease_rejected():
    with pytest.raises(CommercialDailyReportError):
        _daily_deltas({
            date(2026, 10, 1): 8,
            date(2026, 10, 2): 5,
        })


def test_report_has_three_sheets_and_week_to_date_value():
    cutoff = date(2026, 10, 7)
    d = date(2026, 10, 1)
    series = {}
    for n in range(7):
        series[d.replace(day=n + 1)] = n + 1
    payload = render_commercial_daily_xlsx(
        metric="venta_nueva",
        cutoff=cutoff,
        mtd_series={"VILLA_VERDE": series},
    )
    wb = load_workbook(BytesIO(payload), data_only=True)
    assert wb.sheetnames == ["Diario", "Semanal", "Totales Mensuales"]
    assert wb["Diario"]["A3"].value == "VILLA_VERDE"
    assert wb["Diario"]["B3"].value == 1
    weekly = wb["Semanal"]
    assert weekly.cell(row=3, column=weekly.max_column).value == 4

def test_subtotals_and_region_summary_follow_catalog_metadata():
    cutoff = date(2026, 10, 7)
    days = [date(2026, 10, i) for i in range(1, 8)]
    data = {
        "VILLAS": {day: i for i, day in enumerate(days, 1)},
        "SERRANIA": {day: 2 * i for i, day in enumerate(days, 1)},
    }
    metadata = {
        "VILLAS": {
            "order": 1, "label": "VILLAS DEL REY",
            "group": "Subtotal 21 gyms", "region": "Mexicali"
        },
        "SERRANIA": {
            "order": 26, "label": "SERRANIA",
            "group": "Subtotal gyms nuevos", "region": "Costa"
        },
    }
    book = load_workbook(
        BytesIO(render_commercial_daily_xlsx(
            metric="reactivaciones",
            cutoff=cutoff,
            mtd_series=data,
            branch_metadata=metadata,
        )),
        data_only=True,
    )
    daily = book["Diario"]
    assert [daily[f"A{row}"].value for row in range(3, 8)] == [
        "VILLAS DEL REY", "Subtotal 21 gyms",
        "SERRANIA", "Subtotal gyms nuevos", "TOTAL",
    ]
    weekly = book["Semanal"]
    assert weekly["A9"].value == "REGIONES"
    assert weekly["A10"].value == "Costa"
    assert weekly["A11"].value == "Mexicali"
    last = weekly.max_column
    assert weekly.cell(row=7, column=last).value == 12
    assert weekly.cell(row=10, column=last).value == 8
    assert weekly.cell(row=11, column=last).value == 4


def test_reactivaciones_february_correction_does_not_fake_negative_daily():
    """Monthly historical snapshots stay readable despite old MTD corrections."""
    series = {
        date(2026, 2, 12): 33,
        date(2026, 2, 13): 31,  # A corrected cumulative total, not -2 visits.
        date(2026, 2, 28): 99,
        date(2026, 9, 30): 150,
        **{date(2026, 10, day): 2 * day for day in range(1, 8)},
    }
    payload = render_commercial_daily_xlsx(
        metric="reactivaciones",
        cutoff=date(2026, 10, 7),
        mtd_series={"VILLA_VERDE": series},
    )
    book = load_workbook(BytesIO(payload), data_only=True)
    weekly = book["Semanal"]
    daily = book["Diario"]
    monthly = book["Totales Mensuales"]

    assert book.sheetnames == ["Diario", "Semanal", "Totales Mensuales"]
    assert weekly.max_column == 12  # July 26–August 1 through the October current week.
    assert weekly["B2"].value == "26/07-01/08"
    assert weekly["L2"].value == "04/10-10/10 (corte 07/10)"
    assert weekly["L3"].value == 8
    assert daily["B3"].value == 2
    assert daily["E3"].value == 2
    assert daily["F3"].value is None  # Future day never counted as zero.
    assert monthly["C3"].value == 99  # February closing MTD, not daily net deltas.
    assert monthly["J3"].value == 150  # September historical monthly close.
    assert monthly["K3"].value == 14  # October MTD at cutoff.
    assert "04/10/2026" in daily["A1"].value
    assert "26/07/2026 AL 07/10/2026" in weekly["A1"].value
    assert "07/10/2026" in monthly["A1"].value


def test_reactivaciones_current_period_correction_fails_closed():
    days = {
        date(2026, 10, 1): 11,
        date(2026, 10, 2): 8,
        date(2026, 10, 3): 15,
        date(2026, 10, 4): 18,
    }
    with pytest.raises(CommercialDailyReportError, match="2026-10-02"):
        render_commercial_daily_xlsx(
            metric="reactivaciones",
            cutoff=date(2026, 10, 4),
            mtd_series={"VILLA_VERDE": days},
        )


def test_reactivaciones_monthly_history_without_old_daily_dates():
    series = {
        date(2026, 1, 31): 80,
        date(2026, 2, 28): 90,
        date(2026, 10, 1): 3,
        date(2026, 10, 2): 4,
        date(2026, 10, 3): 6,
        date(2026, 10, 4): 7,
    }
    report = load_workbook(
        BytesIO(render_commercial_daily_xlsx(
            metric="reactivaciones",
            cutoff=date(2026, 10, 4),
            mtd_series={"VILLA_VERDE": series},
        )),
        data_only=True,
    )
    assert report["Semanal"].cell(row=3, column=12).value == 1
    assert report["Totales Mensuales"]["B3"].value == 80
    assert report["Totales Mensuales"]["C3"].value == 90
    assert report["Totales Mensuales"]["K3"].value == 7


def test_reactivaciones_ignores_retired_historical_club_but_requires_current_ones(monkeypatch):
    from types import SimpleNamespace
    from app.warehouse.services import commercial_daily_reports_service as service

    jan31 = date(2026, 1, 31)
    cutoff = date(2026, 10, 7)
    version_by_day = {
        jan31: SimpleNamespace(id=11, status="success"),
        cutoff: SimpleNamespace(id=27, status="success"),
    }
    monkeypatch.setattr(service, "_resolve_canonical_version", version_by_day.get)

    class FakeColumn:
        def in_(self, keys):
            return True

    class FakeQuery:
        def __init__(self, rows):
            self.rows = rows
        def filter(self, *args):
            return self
        def all(self):
            return self.rows

    rows = [
        SimpleNamespace(
            sucursal_canon="LA_VIGA", track_date=jan31,
            track_daily_version_id=11, reactivaciones_real_mtd=40,
        ),
        SimpleNamespace(
            sucursal_canon="VILLA_VERDE", track_date=jan31,
            track_daily_version_id=11, reactivaciones_real_mtd=30,
        ),
        SimpleNamespace(
            sucursal_canon="VILLA_VERDE", track_date=cutoff,
            track_daily_version_id=27, reactivaciones_real_mtd=15,
        ),
        SimpleNamespace(
            sucursal_canon="SERRANIA", track_date=cutoff,
            track_daily_version_id=27, reactivaciones_real_mtd=8,
        ),
    ]
    monkeypatch.setattr(
        service, "TrackDailyMartORM",
        SimpleNamespace(
            track_daily_version_id=FakeColumn(),
            query=FakeQuery(rows),
        ),
    )
    output = service._load_series(
        "reactivaciones",
        cutoff,
        active_branches={"VILLA_VERDE", "SERRANIA"},
    )
    assert set(output) == {"VILLA_VERDE", "SERRANIA"}
    assert output["VILLA_VERDE"][jan31] == 30
    assert output["SERRANIA"][cutoff] == 8

    with pytest.raises(CommercialDailyReportError, match="MISSING_CLUB"):
        service._load_series(
            "reactivaciones",
            cutoff,
            active_branches={"VILLA_VERDE", "SERRANIA", "MISSING_CLUB"},
        )


def test_jan_apr_reactivaciones_historical_baseline_has_all_branches():
    from app.warehouse.services.reactivaciones_monthly_baseline_service import (
        load_early_reactivaciones_monthly,
    )
    snapshot = load_early_reactivaciones_monthly()
    assert len(snapshot) == 26
    assert snapshot["VILLA_VERDE"][date(2026, 1, 1)] == 169
    for month, expected in (
        (1, 4311), (2, 2684), (3, 3789), (4, 2684)
    ):
        assert sum(
            values[date(2026, month, 1)]
            for values in snapshot.values()
        ) == expected


def test_jan_apr_reactivaciones_fills_only_missing_monthly_track_data():
    from app.warehouse.services.reactivaciones_monthly_baseline_service import (
        load_early_reactivaciones_monthly,
    )
    snapshot = load_early_reactivaciones_monthly()
    series = {
        branch: {
            date(2026, 5, 31): 27,
            **{date(2026, 10, day): day * 2 for day in range(1, 8)},
        }
        for branch in snapshot
    }
    payload = render_commercial_daily_xlsx(
        metric="reactivaciones",
        cutoff=date(2026, 10, 7),
        mtd_series=series,
        historical_monthly_by_branch=snapshot,
    )
    book = load_workbook(BytesIO(payload), data_only=True)
    monthly = book["Totales Mensuales"]
    total_row = next(
        row for row in range(3, monthly.max_row + 1)
        if monthly.cell(row, 1).value == "TOTAL"
    )
    assert [
        monthly.cell(total_row, col).value for col in range(2, 6)
    ] == [4311, 2684, 3789, 2684]
    assert monthly.cell(total_row, 6).value == 26 * 27
    assert monthly.cell(total_row, 11).value == 26 * 14


def test_jan_apr_reactivaciones_rejects_conflict_with_existing_canonical_close():
    from app.warehouse.services.reactivaciones_monthly_baseline_service import (
        load_early_reactivaciones_monthly,
    )
    snapshot = load_early_reactivaciones_monthly()
    series = {
        branch: {
            date(2026, 10, 1): 2,
            date(2026, 10, 2): 4,
            date(2026, 10, 3): 6,
            date(2026, 10, 4): 8,
        }
        for branch in snapshot
    }
    series["VILLA_VERDE"][date(2026, 3, 31)] = 144
    with pytest.raises(CommercialDailyReportError, match="Historical Reactivaciones conflict"):
        render_commercial_daily_xlsx(
            metric="reactivaciones",
            cutoff=date(2026, 10, 4),
            mtd_series=series,
            historical_monthly_by_branch=snapshot,
        )


def test_january_february_explicit_historical_source_overrides_old_track_unchanged():
    from app.warehouse.services.reactivaciones_monthly_baseline_service import (
        load_early_reactivaciones_monthly,
    )

    historical = load_early_reactivaciones_monthly()
    series = {
        branch: {
            date(2026, 10, 1): 3,
            date(2026, 10, 2): 5,
            date(2026, 10, 3): 8,
            date(2026, 10, 4): 11,
            date(2026, 10, 5): 14,
            date(2026, 10, 6): 17,
            date(2026, 10, 7): 20,
        }
        for branch in historical
    }
    # Actual historical canon: snapshots 176/162 were older than the BI book.
    series["VILLAS_DEL_REY"][date(2026, 1, 31)] = 184
    series["VILLAS_DEL_REY"][date(2026, 2, 28)] = 123
    series["VILLA_VERDE"][date(2026, 1, 31)] = 161
    series["VILLA_VERDE"][date(2026, 2, 28)] = 121
    series["VILLA_VERDE"][date(2026, 3, 31)] = 143
    series["VILLA_VERDE"][date(2026, 4, 30)] = 116

    workbook = load_workbook(
        BytesIO(render_commercial_daily_xlsx(
            metric="reactivaciones",
            cutoff=date(2026, 10, 7),
            mtd_series=series,
            historical_monthly_by_branch=historical,
        )),
        data_only=False,
    )
    monthly = workbook["Totales Mensuales"]
    row_by_name = {
        str(monthly.cell(row, 1).value): row
        for row in range(3, monthly.max_row + 1)
    }
    villas = row_by_name["VILLAS_DEL_REY"]
    verde = row_by_name["VILLA_VERDE"]
    total = row_by_name["TOTAL"]
    assert monthly.cell(villas, 2).value == 211
    assert monthly.cell(verde, 2).value == 169
    assert monthly.cell(verde, 3).value == 123
    assert monthly.cell(total, 2).value == 4311
    assert monthly.cell(total, 3).value == 2684
    assert monthly.cell(total, 4).value == 3789
    assert monthly.cell(total, 5).value == 2684
    assert monthly["B2"].comment is not None
    assert monthly["C2"].comment is not None
    assert "Excel historico" in monthly["B2"].comment.text
    assert any(
        isinstance(monthly.cell(row, 1).value, str)
        and "FUENTES: Ene-Feb 2026" in monthly.cell(row, 1).value
        for row in range(total + 1, monthly.max_row + 1)
    )
    assert workbook.sheetnames == ["Diario", "Semanal", "Totales Mensuales"]


def test_weekly_reactivaciones_comparison_covers_august_september_october():
    start = date(2026, 7, 25)  # preceding MTD for the first Sunday
    cutoff = date(2026, 10, 7)
    days = [start + timedelta(days=i) for i in range((cutoff - start).days + 1)]
    data = {
        "VILLA_VERDE": {day: day.day for day in days},
        "VILLAS_DEL_REY": {day: day.day * 2 for day in days},
    }
    workbook = load_workbook(
        BytesIO(render_commercial_daily_xlsx(
            metric="reactivaciones",
            cutoff=cutoff,
            mtd_series=data,
        )),
        data_only=True,
    )
    weekly = workbook["Semanal"]
    assert weekly.max_column == 12
    assert weekly["B2"].value == "26/07-01/08"
    assert weekly["C2"].value == "02/08-08/08"
    assert weekly["G2"].value == "30/08-05/09"
    assert weekly["K2"].value == "27/09-03/10"
    assert weekly["L2"].value == "04/10-10/10 (corte 07/10)"
    # Branch order is deterministic but alphabetical: VILLAS_DEL_REY
    # precedes VILLA_VERDE in this synthetic fixture.
    assert weekly["A3"].value == "VILLAS_DEL_REY"
    assert weekly["A4"].value == "VILLA_VERDE"
    assert weekly["B3"].value == 14
    assert weekly["B4"].value == 7
    assert weekly["B5"].value == 21
    assert weekly["C3"].value == 14
    assert weekly["L3"].value == 8
    assert weekly["L4"].value == 4
    assert weekly["L5"].value == 12
    assert workbook["Totales Mensuales"]["I3"].value == 62
    assert workbook["Totales Mensuales"]["J3"].value == 60
    assert workbook["Totales Mensuales"]["K3"].value == 14
    assert "26/07/2026 AL 07/10/2026" in weekly["A1"].value


def test_older_negative_reactivaciones_is_blank_not_negative_or_zero():
    start = date(2026, 8, 1)
    cutoff = date(2026, 10, 7)
    days = [start + timedelta(days=i) for i in range((cutoff - start).days + 1)]
    values = {day: day.day for day in days}
    values[date(2026, 8, 10)] = 4  # corrected down from 9
    report = load_workbook(
        BytesIO(render_commercial_daily_xlsx(
            metric="reactivaciones",
            cutoff=cutoff,
            mtd_series={"VILLA_VERDE": values},
        )),
        data_only=True,
    )
    weekly = report["Semanal"]
    assert weekly["D2"].value == "09/08-15/08"
    assert weekly["D3"].value is None
    assert weekly["D4"].value is None  # no misleading subtotal
    assert weekly["E3"].value == 7
    assert weekly["L3"].value == 4
    assert any(
        "no representan cero" in str(weekly.cell(row, 1).value)
        for row in range(5, weekly.max_row + 1)
    )


def test_reactivaciones_three_month_window_rolls_at_calendar_year():
    from app.warehouse.services.commercial_daily_reports_service import (
        _three_month_window_start,
    )
    assert _three_month_window_start(date(2026, 10, 7)) == date(2026, 8, 1)
    assert _three_month_window_start(date(2027, 1, 7)) == date(2026, 11, 1)
    assert _three_month_window_start(date(2026, 1, 7)) == date(2026, 1, 1)


def test_reactivaciones_week_labels_always_span_sunday_to_saturday():
    import re
    start = date(2026, 7, 25)
    cutoff = date(2026, 10, 7)
    dates = (
        start + timedelta(days=i)
        for i in range((cutoff - start).days + 1)
    )
    data = {day: day.day for day in dates}
    report = load_workbook(
        BytesIO(render_commercial_daily_xlsx(
            metric="reactivaciones", cutoff=cutoff,
            mtd_series={"VILLA_VERDE": data},
        )),
        data_only=True,
    )
    sheet = report["Semanal"]
    assert sheet.max_column - 1 == 11
    for col in range(2, sheet.max_column + 1):
        label = sheet.cell(2, col).value
        match = re.match(r"^(\d{2})/(\d{2})-(\d{2})/(\d{2})", label)
        assert match is not None, label
        start_day, start_month, end_day, end_month = map(int, match.groups())
        sunday = date(2026, start_month, start_day)
        saturday = date(2026, end_month, end_day)
        assert sunday.weekday() == 6
        assert saturday.weekday() == 5
        assert saturday - sunday == timedelta(days=6)
    assert sheet["B3"].value == 7


def test_reactivaciones_missing_previous_day_does_not_invent_july_26():
    start = date(2026, 7, 26)  # July 25 not included: Sunday delta unknown
    cutoff = date(2026, 10, 7)
    series = {
        start + timedelta(days=i): (start + timedelta(days=i)).day
        for i in range((cutoff - start).days + 1)
    }
    report = load_workbook(
        BytesIO(render_commercial_daily_xlsx(
            metric="reactivaciones", cutoff=cutoff,
            mtd_series={"VILLA_VERDE": series},
        )),
        data_only=True,
    )
    weekly = report["Semanal"]
    assert weekly["B2"].value == "26/07-01/08"
    assert weekly["B3"].value is None
    assert weekly["B4"].value is None
    assert weekly["C3"].value == 7
