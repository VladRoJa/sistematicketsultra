from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch
import unittest

from openpyxl import Workbook, load_workbook

from app.services import mantenimiento_edificio_summary_service as summary


def _classification(name, parent=None, node_id=1):
    return SimpleNamespace(id=node_id, nombre=name, padre=parent)


def _building_ticket(ticket_id, branch, category=None):
    root = _classification("Mantenimiento", node_id=100)
    building = _classification("Edificio", root, 101)
    classification = (
        _classification(category, building, 102 + ticket_id)
        if category
        else building
    )
    branch_obj = SimpleNamespace(sucursal=branch)
    return SimpleNamespace(
        id=ticket_id,
        aparato_id=None,
        familia_equipo_id=None,
        clasificacion=classification,
        categoria=None,
        subcategoria=None,
        detalle=None,
        sucursal_destino=branch_obj,
        sucursal=None,
    )


def _equipment_ticket(ticket_id, branch):
    root = _classification("Mantenimiento", node_id=200)
    apparatus = _classification("Aparatos", root, 201)
    classification = _classification("Peso Libre", apparatus, 202 + ticket_id)
    branch_obj = SimpleNamespace(sucursal=branch)
    return SimpleNamespace(
        id=ticket_id,
        aparato_id=90,
        familia_equipo_id=6,
        clasificacion=classification,
        categoria=None,
        subcategoria=None,
        detalle=None,
        sucursal_destino=branch_obj,
        sucursal=None,
    )


def _base_report():
    workbook = Workbook()
    workbook.active.title = "Resumen ejecutivo"
    workbook.create_sheet("Tickets")
    workbook.create_sheet("Edificio")
    workbook.create_sheet("Equipos por familia")

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


class MantenimientoEdificioSummaryServiceTest(unittest.TestCase):
    def test_matriz_edificio_agrupa_sucursal_y_categoria(self):
        tickets = [
            _building_ticket(1, "SUCURSAL A", "Electricidad"),
            _building_ticket(2, "SUCURSAL A", "Electricidad"),
            _building_ticket(3, "SUCURSAL A", None),
            _building_ticket(4, "SUCURSAL B", "Plomería"),
            _equipment_ticket(5, "SUCURSAL A"),
        ]

        headers, rows, total = summary._building_category_matrix(tickets)

        self.assertEqual(
            headers,
            (
                "Sucursal",
                "Fallas",
                "Electricidad",
                "Plomería",
                "Sin clasificar",
            ),
        )
        self.assertEqual(
            rows,
            [
                ("SUCURSAL A", 3, 2, 0, 1),
                ("SUCURSAL B", 1, 0, 1, 0),
            ],
        )
        self.assertEqual(total, ("TOTAL", 4, 2, 1, 1))

        for row in rows:
            self.assertEqual(row[1], sum(row[2:]))
        self.assertEqual(total[1], sum(total[2:]))

    def test_xlsx_agrega_resumen_inmediatamente_despues_de_edificio(self):
        tickets = [
            _building_ticket(10, "SUCURSAL A", "Electricidad"),
            _building_ticket(11, "SUCURSAL B", "Plomería"),
        ]

        with patch.object(
            summary.report,
            "construir_reporte_xlsx",
            return_value=_base_report(),
        ) as base_builder:
            output = summary.construir_reporte_xlsx(
                tickets,
                historical_tickets=tickets,
            )

        workbook = load_workbook(output)
        self.assertEqual(
            workbook.sheetnames,
            [
                "Resumen ejecutivo",
                "Tickets",
                "Edificio",
                "Edificio por categoría",
                "Equipos por familia",
            ],
        )

        rows = list(
            workbook["Edificio por categoría"].iter_rows(values_only=True)
        )
        self.assertEqual(
            rows[0],
            (
                "Sucursal",
                "Fallas",
                "Electricidad",
                "Plomería",
                "Sin clasificar",
            ),
        )
        self.assertEqual(rows[-1], ("TOTAL", 2, 1, 1, 0))
        base_builder.assert_called_once()


if __name__ == "__main__":
    unittest.main()
