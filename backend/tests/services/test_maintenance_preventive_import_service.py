import io
import unittest

from openpyxl import Workbook, load_workbook

from app.services.maintenance_preventive_import_service import (
    TEMPLATE_HEADERS,
    build_preventive_template_xlsx,
    parse_preventive_import,
)
from app.services.maintenance_preventive_service import MaintenancePreventiveError


class MaintenancePreventiveImportServiceTest(unittest.TestCase):
    def test_official_template_has_catalogs_validations_and_date_format(self):
        content = build_preventive_template_xlsx(
            sucursales=["Paseo 2000", "Pabellon rosarito"],
            responsables=["SR_MANT_TIJ", "AUX_MANT_TIJ"],
            equipos=[
                {
                    "sucursal": "Paseo 2000",
                    "codigo_interno": "09CBRLF03",
                    "nombre": "Bicicleta Recumbente",
                    "familia": "Bicicleta Recumbente",
                }
            ],
            building_classifications=[
                "Baños > Mingitorios",
                "Instalaciones > Hidráulica",
            ],
        )

        workbook = load_workbook(io.BytesIO(content))
        try:
            self.assertEqual(
                workbook.sheetnames,
                [
                    "Programacion preventiva",
                    "Catálogos",
                    "Validaciones",
                ],
            )
            self.assertEqual(
                workbook["Validaciones"].sheet_state,
                "hidden",
            )

            sheet = workbook["Programacion preventiva"]
            self.assertEqual(
                tuple(cell.value for cell in sheet[1]),
                TEMPLATE_HEADERS,
            )
            self.assertIsNone(sheet["A2"].value)
            self.assertEqual(sheet["D2"].number_format, "dd/mm/yyyy")
            self.assertEqual(sheet.freeze_panes, "A2")

            validations = list(
                sheet.data_validations.dataValidation
            )
            validation_types = {
                validation.type
                for validation in validations
            }
            self.assertIn("list", validation_types)
            self.assertIn("date", validation_types)
            self.assertIn("custom", validation_types)
            self.assertIn("whole", validation_types)

            formulas = {
                validation.formula1
                for validation in validations
                if validation.type == "list"
            }
            self.assertIn("=CatalogoSucursales", formulas)
            self.assertIn("=CatalogoResponsables", formulas)
            self.assertIn('"No,Sí"', formulas)
            self.assertIn('"Equipo,Edificio"', formulas)
            self.assertIn("=CatalogoEdificio", formulas)
            self.assertIn(
                '=INDIRECT(IFERROR(VLOOKUP($A2,MapaFamilias,2,FALSE),"ListaVacia"))',
                formulas,
            )
            self.assertIn(
                '=INDIRECT(IFERROR(VLOOKUP($A2&"|"&$B2,MapaCodigos,2,FALSE),"ListaVacia"))',
                formulas,
            )

            activity_validation = next(
                validation
                for validation in validations
                if validation.type == "custom"
            )
            self.assertEqual(
                activity_validation.formula1,
                "=LEN(TRIM(E2))>0",
            )
            self.assertEqual(
                activity_validation.promptTitle,
                "Actividad preventiva",
            )
            self.assertIn(
                "Obligatoria",
                activity_validation.prompt,
            )

            catalogs = workbook["Catálogos"]
            self.assertEqual(catalogs["A2"].value, "Paseo 2000")
            self.assertEqual(catalogs["B2"].value, "SR_MANT_TIJ")
            self.assertEqual(catalogs["D2"].value, "Paseo 2000")
            self.assertEqual(
                catalogs["E2"].value,
                "Bicicleta Recumbente",
            )
            self.assertEqual(catalogs["F2"].value, "09CBRLF03")
            self.assertEqual(
                catalogs["G2"].value,
                "Bicicleta Recumbente",
            )
            self.assertEqual(
                catalogs["H2"].value,
                "Baños > Mingitorios",
            )
            self.assertIn(
                "CatalogoSucursales",
                workbook.defined_names,
            )
            self.assertIn(
                "CatalogoResponsables",
                workbook.defined_names,
            )
            self.assertIn("MapaFamilias", workbook.defined_names)
            self.assertIn("MapaCodigos", workbook.defined_names)
            self.assertIn("CatalogoEdificio", workbook.defined_names)
            self.assertTrue(
                any(
                    name.startswith("Familias_")
                    for name in workbook.defined_names
                )
            )
            self.assertTrue(
                any(
                    name.startswith("Codigos_")
                    for name in workbook.defined_names
                )
            )

            self.assertEqual(
                sheet["G1"].value,
                "Se repite",
            )
            self.assertEqual(
                sheet["H1"].value,
                "Cada N días hábiles",
            )
            self.assertEqual(sheet["J1"].value, "Tipo")
            self.assertEqual(
                sheet["K1"].value,
                "Clasificación edificio",
            )
            self.assertEqual(
                sheet["L1"].value,
                "Duración estimada (min)",
            )
            self.assertEqual(
                sheet["M1"].value,
                "Validación",
            )
            self.assertIn(
                "COUNTIFS",
                sheet["M2"].value,
            )
            self.assertIn(
                "DUPLICADO",
                sheet["M2"].value,
            )

            conditional_ranges = {
                str(rule_range.sqref)
                for rule_range in sheet.conditional_formatting
            }
            self.assertIn("A2:M1001", conditional_ranges)

            duplicate_rules = list(
                sheet.conditional_formatting["A2:M1001"]
            )
            self.assertEqual(len(duplicate_rules), 1)
            self.assertEqual(
                duplicate_rules[0].formula[0],
                '$M2="⚠ DUPLICADO"',
            )
        finally:
            workbook.close()

    def test_generated_template_roundtrips_dd_mm_yyyy_date(self):
        content = build_preventive_template_xlsx(
            sucursales=["Paseo 2000"],
            responsables=["SR_MANT_TIJ"],
        )
        workbook = load_workbook(io.BytesIO(content))
        sheet = workbook["Programacion preventiva"]
        values = [
            "Paseo 2000",
            "Bicicleta Recumbente",
            "09CBRLF03",
            "25/09/2026",
            "Mantenimiento general",
            "SR_MANT_TIJ",
            "Sí",
            5,
            "",
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row=2, column=column, value=value)

        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()

        parsed = parse_preventive_import(
            filename="plantilla_programacion_preventiva.xlsx",
            content=stream.getvalue(),
        )

        self.assertEqual(len(parsed["rows"]), 1)
        self.assertEqual(
            parsed["rows"][0]["fecha_programada"],
            "2026-09-25",
        )
        self.assertEqual(
            parsed["rows"][0]["repeat_enabled"],
            "Sí",
        )
        self.assertEqual(
            parsed["rows"][0]["repeat_interval_workdays"],
            "5",
        )
        self.assertIsNone(
            parsed["rows"][0]["estimated_duration_minutes"]
        )

    def test_csv_accepts_accented_official_headers(self):
        csv_content = (
            "Sucursal,Código equipo,Fecha programada,Actividad,"
            "Responsable,Observaciones\n"
            "VILLAS DEL REY,04CC01,20/09/2026,"
            "Mantenimiento general,TECNICO_PM,Sin novedad\n"
        ).encode("utf-8")

        parsed = parse_preventive_import(
            filename="preventivos.csv",
            content=csv_content,
        )

        self.assertEqual(len(parsed["rows"]), 1)
        self.assertEqual(
            parsed["rows"][0]["observaciones"],
            "Sin novedad",
        )
        self.assertEqual(
            parsed["rows"][0]["fecha_programada"],
            "2026-09-20",
        )

    def test_csv_accepts_optional_estimated_duration_header(self):
        csv_content = (
            "Sucursal,Código equipo,Fecha programada,Actividad,"
            "Responsable,Duración estimada (min)\n"
            "VILLAS DEL REY,04CC01,21/09/2026,"
            "Mantenimiento general,TECNICO_PM,45\n"
        ).encode("utf-8")

        parsed = parse_preventive_import(
            filename="preventivos_duracion.csv",
            content=csv_content,
        )

        self.assertEqual(
            parsed["rows"][0]["estimated_duration_minutes"],
            "45",
        )

    def test_csv_accepts_optional_recurrence_headers(self):
        csv_content = (
            "Sucursal,Código equipo,Fecha programada,Actividad,"
            "Responsable,Se repite,Cada N días hábiles,Observaciones\n"
            "VILLAS DEL REY,04CC01,21/09/2026,"
            "Mantenimiento general,TECNICO_PM,Sí,10,Recurrente\n"
        ).encode("utf-8")

        parsed = parse_preventive_import(
            filename="preventivos_recurrentes.csv",
            content=csv_content,
        )

        row = parsed["rows"][0]
        self.assertEqual(row["repeat_enabled"], "Sí")
        self.assertEqual(row["repeat_interval_workdays"], "10")
        self.assertEqual(row["fecha_programada"], "2026-09-21")

    def test_csv_accepts_building_target_without_equipment_code(self):
        csv_content = (
            "Sucursal,Fecha programada,Actividad,Responsable,"
            "Tipo,Clasificación edificio,Se repite,"
            "Cada N días hábiles\n"
            "VILLAS DEL REY,21/09/2026,Revisión de baños,"
            "TECNICO_PM,Edificio,Baños > Mingitorios,Sí,20\n"
        ).encode("utf-8")

        parsed = parse_preventive_import(
            filename="preventivos_edificio.csv",
            content=csv_content,
        )

        row = parsed["rows"][0]
        self.assertIsNone(row["codigo_equipo"])
        self.assertEqual(row["target_type"], "Edificio")
        self.assertEqual(
            row["building_classification"],
            "Baños > Mingitorios",
        )

    def test_missing_required_header_is_rejected(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(
            [
                "Sucursal",
                "Código equipo",
                "Fecha programada",
                "Actividad",
                # Responsable omitido a propósito.
                "Observaciones",
            ]
        )
        sheet.append(
            [
                "VILLAS DEL REY",
                "04CC01",
                "2026-09-20",
                "Mantenimiento",
                "",
            ]
        )

        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()

        with self.assertRaisesRegex(
            MaintenancePreventiveError,
            "Responsable",
        ):
            parse_preventive_import(
                filename="incompleta.xlsx",
                content=stream.getvalue(),
            )

    def test_unsupported_extension_is_rejected(self):
        with self.assertRaisesRegex(
            MaintenancePreventiveError,
            "Formato no soportado",
        ):
            parse_preventive_import(
                filename="preventivos.xls",
                content=b"legacy",
            )

    def test_template_contract_has_exact_column_order(self):
        self.assertEqual(
            TEMPLATE_HEADERS,
            (
                "Sucursal",
                "Familia",
                "Código equipo",
                "Fecha programada",
                "Actividad",
                "Responsable",
                "Se repite",
                "Cada N días hábiles",
                "Observaciones",
                "Tipo",
                "Clasificación edificio",
                "Duración estimada (min)",
                "Validación",
            ),
        )


if __name__ == "__main__":
    unittest.main()
