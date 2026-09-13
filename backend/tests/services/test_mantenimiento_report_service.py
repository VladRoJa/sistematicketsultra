from datetime import datetime, timezone
from types import SimpleNamespace
from unittest.mock import patch
import unittest

from openpyxl import load_workbook

from app.services import mantenimiento_report_service as report


class _Expression:
    def __init__(self, predicate):
        self.predicate = predicate


class _Column:
    def __init__(self, name):
        self.name = name

    def __eq__(self, value):
        return _Expression(lambda row: getattr(row, self.name) == value)

    def in_(self, values):
        return _Expression(lambda row: getattr(row, self.name) in values)

    def asc(self):
        return self


class _Query:
    def __init__(self, rows):
        self.rows = list(rows)

    def options(self, *_args):
        return self

    def filter(self, *expressions):
        self.rows = [
            row
            for row in self.rows
            if all(expression.predicate(row) for expression in expressions)
        ]
        return self

    def order_by(self, *_args):
        return self

    def all(self):
        return self.rows


def _classification(name, parent=None, node_id=1):
    return SimpleNamespace(
        id=node_id,
        nombre=name,
        padre=parent,
    )


def _ticket(
    ticket_id,
    *,
    building=False,
    state="en progreso",
    branch="SUCURSAL UNO",
):
    branch_obj = SimpleNamespace(sucursal=branch)
    root = _classification("Mantenimiento", node_id=100)

    if building:
        building_node = _classification("Edificio", root, 101)
        category = _classification("Electricidad", building_node, 102)
        classification = _classification("Iluminación", category, 103)
        inventory = None
        apparatus_id = None
        family_id = None
        family = None
        failure = None
        equipment = "Luminaria vestidores"
        location = "Vestidores"
        condition = None
    else:
        apparatus_node = _classification("Aparatos", root, 201)
        classification = _classification("Peso Libre", apparatus_node, 202)
        family = SimpleNamespace(key="PESO_LIBRE", nombre="Peso Libre")
        inventory = SimpleNamespace(
            id=90,
            nombre="Rack de mancuernas",
            codigo_interno="07PLJW36",
            familia_equipo=family,
            familia_equipo_id=6,
        )
        apparatus_id = 90
        family_id = 6
        failure = SimpleNamespace(nombre="Amortiguador dañado")
        equipment = None
        location = None
        condition = "NO_TRABAJA"

    ticket = SimpleNamespace(
        id=ticket_id,
        descripcion="Prueba de mantenimiento",
        estado=state,
        departamento_id=1,
        criticidad=5 if building else 4,
        aparato_id=apparatus_id,
        familia_equipo_id=family_id,
        fecha_creacion=datetime(2026, 9, 10, 18, 0, tzinfo=timezone.utc),
        fecha_solucion=datetime(2026, 9, 11, 14, 0, tzinfo=timezone.utc),
        fecha_finalizado=None,
        sucursal_id=10,
        sucursal_id_destino=10,
        sucursal_destino=branch_obj,
        sucursal=None,
        inventario=inventory,
        familia_equipo=family,
        falla_mantenimiento=failure,
        condicion_operativa=condition,
        historial_fechas=[
            {
                "fechaCambio": "2026-09-10T18:00:00Z",
                "motivo": "Plan de trabajo",
            }
        ],
        clasificacion=classification,
        categoria=None,
        subcategoria=None,
        detalle=None,
        equipo=equipment,
        ubicacion=location,
        necesita_refaccion=building,
        descripcion_refaccion="Balastro" if building else None,
    )
    return ticket


def _fake_ticket_model(rows):
    return SimpleNamespace(
        query=_Query(rows),
        inventario="inventario",
        familia_equipo="familia_equipo",
        falla_mantenimiento="falla_mantenimiento",
        sucursal="sucursal",
        sucursal_destino="sucursal_destino",
        clasificacion="clasificacion",
        departamento_id=_Column("departamento_id"),
        estado=_Column("estado"),
        fecha_creacion=_Column("fecha_creacion"),
        sucursal_id_destino=_Column("sucursal_id_destino"),
        id=_Column("id"),
    )


class MantenimientoReportServiceTest(unittest.TestCase):
    def test_clasifica_edificio_por_jerarquia_real(self):
        building = _ticket(1, building=True)
        equipment = _ticket(2, building=False)

        self.assertEqual(report._maintenance_type(building), report.TYPE_BUILDING)
        self.assertEqual(report._maintenance_type(equipment), report.TYPE_EQUIPMENT)
        self.assertEqual(
            report._building_classification(building),
            ("Electricidad", "Iluminación", None),
        )

    def test_consulta_activa_incluye_edificio_sin_aparato_ni_familia(self):
        building = _ticket(3, building=True)
        equipment = _ticket(4, building=False)
        other_department = _ticket(5, building=True)
        other_department.departamento_id = 7
        finalized = _ticket(6, building=True, state="finalizado")

        fake_ticket_model = _fake_ticket_model(
            [building, equipment, other_department, finalized]
        )

        with (
            patch.object(report, "Ticket", fake_ticket_model),
            patch.object(report, "joinedload", side_effect=lambda value: value),
        ):
            selected = report.obtener_tickets_reporte()

        self.assertEqual([ticket.id for ticket in selected], [3, 4])

    def test_xlsx_combina_equipo_y_edificio_con_caratula_primero(self):
        equipment = _ticket(10, building=False)
        building = _ticket(11, building=True)

        output = report.construir_reporte_xlsx(
            [equipment, building],
            historical_tickets=[equipment, building],
            now=datetime(2026, 9, 12, 18, 0, tzinfo=timezone.utc),
        )
        workbook = load_workbook(output)

        self.assertEqual(workbook.sheetnames[0], "Resumen ejecutivo")
        self.assertIn("Tickets", workbook.sheetnames)
        self.assertIn("Edificio", workbook.sheetnames)
        self.assertIn("Equipos por familia", workbook.sheetnames)

        ticket_rows = list(workbook["Tickets"].iter_rows(values_only=True))
        self.assertEqual(ticket_rows[0][0], "Tipo mantenimiento")
        self.assertEqual(
            {row[0] for row in ticket_rows[1:]},
            {report.TYPE_EQUIPMENT, report.TYPE_BUILDING},
        )

        building_rows = list(workbook["Edificio"].iter_rows(values_only=True))
        self.assertEqual(len(building_rows), 2)
        self.assertEqual(building_rows[1][2], "Electricidad")
        self.assertEqual(building_rows[1][3], "Iluminación")
        self.assertEqual(building_rows[1][5], "Luminaria vestidores")
        self.assertEqual(building_rows[1][6], "Vestidores")
        self.assertEqual(building_rows[1][8], 5)

        cover = workbook["Resumen ejecutivo"]
        self.assertEqual(cover["A1"].value, "REPORTE EJECUTIVO DE MANTENIMIENTO")
        self.assertEqual(cover["K12"].value, report.TYPE_EQUIPMENT)
        self.assertEqual(cover["L12"].value, 1)
        self.assertEqual(cover["K13"].value, report.TYPE_BUILDING)
        self.assertEqual(cover["L13"].value, 1)


if __name__ == "__main__":
    unittest.main()
