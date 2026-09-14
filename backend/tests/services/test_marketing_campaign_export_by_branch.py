from io import BytesIO
from types import SimpleNamespace as NS
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest

from app.services import marketing_campaign_export_service as export_service
from app.services.marketing_reactivation_service import (
    MarketingReactivationValidationError,
)


def _recipient(phone, branch, name=None):
    return {
        "phone_mx10": phone,
        "sucursal": branch,
        "member_name": name,
    }


def _mock_branch_labels(monkeypatch, labels):
    normalized = {
        export_service.normalize_socios_vencidos_branch_key(branch): label
        for branch, label in labels.items()
    }
    monkeypatch.setattr(
        export_service,
        "_branch_message_labels",
        lambda **kwargs: normalized,
    )


def test_single_branch_exports_direct_xlsx_with_message_branch(monkeypatch):
    monkeypatch.setattr(
        export_service,
        "get_marketing_reactivation_campaign",
        lambda **kwargs: {
            "id": 7,
            "name": "Campaña septiembre",
            "recipients": [
                _recipient(
                    "6861000002",
                    "SEND MXL",
                    "ROSA MARIA GONZALEZ VELASCO",
                ),
                _recipient(
                    "6861000001",
                    "SEND MXL",
                    "MARÍA DE JESÚS MARISCAL ALBA",
                ),
            ],
        },
    )
    _mock_branch_labels(
        monkeypatch,
        {"SEND MXL": "Sendero Mexicali"},
    )
    guarded_calls = []
    monkeypatch.setattr(
        export_service,
        "_validate_and_mark_exported",
        lambda **kwargs: guarded_calls.append(kwargs) or (b"ignored", "ignored.xlsx"),
    )

    data, filename = export_service.export_marketing_reactivation_campaign(
        campaign_id=7,
        allowed_sucursal_keys=("SEND MXL",),
        session=NS(),
    )

    assert filename == "CAMPANA_SEPTIEMBRE__SEND_MXL.xlsx"
    assert export_service.campaign_export_mimetype(filename).endswith("sheet")
    assert list(load_workbook(BytesIO(data)).active.values) == [
        (
            "telefono",
            "nombre",
            "nombre_completo",
            "sucursal",
            "sucursal_mensaje",
        ),
        (
            "6861000001",
            "María",
            "MARÍA DE JESÚS MARISCAL ALBA",
            "SEND MXL",
            "Sendero Mexicali",
        ),
        (
            "6861000002",
            "Rosa",
            "ROSA MARIA GONZALEZ VELASCO",
            "SEND MXL",
            "Sendero Mexicali",
        ),
    ]
    assert guarded_calls[0]["campaign_id"] == 7


def test_multiple_branches_export_zip_with_summary(monkeypatch):
    monkeypatch.setattr(
        export_service,
        "get_marketing_reactivation_campaign",
        lambda **kwargs: {
            "id": 9,
            "name": "Vencidos agosto 2026",
            "recipients": [
                _recipient("6861000001", "VILLAS DEL REY", "ANA LOPEZ"),
                _recipient("6861000002", "VILLALTA", "JUAN PEREZ"),
                _recipient("6861000003", "VILLALTA", "LUIS GARCIA"),
                _recipient("6861000004", "METEPEC", "SOFIA RAMIREZ"),
            ],
        },
    )
    _mock_branch_labels(
        monkeypatch,
        {
            "VILLAS DEL REY": "Villas del Rey",
            "VILLALTA": "Villalta",
            "METEPEC": "Metepec",
        },
    )
    monkeypatch.setattr(
        export_service,
        "_validate_and_mark_exported",
        lambda **kwargs: (b"ignored", "ignored.xlsx"),
    )

    data, filename = export_service.export_marketing_reactivation_campaign(
        campaign_id=9,
        session=NS(),
    )

    assert filename == "VENCIDOS_AGOSTO_2026.zip"
    assert export_service.campaign_export_mimetype(filename) == "application/zip"
    with ZipFile(BytesIO(data)) as archive:
        assert set(archive.namelist()) == {
            "VENCIDOS_AGOSTO_2026__METEPEC.xlsx",
            "VENCIDOS_AGOSTO_2026__VILLALTA.xlsx",
            "VENCIDOS_AGOSTO_2026__VILLAS_DEL_REY.xlsx",
            "RESUMEN.xlsx",
        }
        villalta = load_workbook(BytesIO(archive.read(
            "VENCIDOS_AGOSTO_2026__VILLALTA.xlsx"
        )))
        assert list(villalta.active.values) == [
            (
                "telefono",
                "nombre",
                "nombre_completo",
                "sucursal",
                "sucursal_mensaje",
            ),
            ("6861000002", "Juan", "JUAN PEREZ", "VILLALTA", "Villalta"),
            ("6861000003", "Luis", "LUIS GARCIA", "VILLALTA", "Villalta"),
        ]
        summary = load_workbook(BytesIO(archive.read("RESUMEN.xlsx")))
        assert list(summary.active.values) == [
            ("sucursal", "contactos"),
            ("METEPEC", 1),
            ("VILLALTA", 2),
            ("VILLAS DEL REY", 1),
            ("TOTAL", 4),
        ]


def test_short_name_normalizes_whitespace_and_keeps_accents():
    assert export_service._short_name("  MARÍA   DE JESÚS  ") == "María"
    assert export_service._short_name(None) == ""


def test_message_branch_label_is_customer_facing():
    assert export_service._message_branch_label("SENDERO MEXICALI") == "Sendero Mexicali"
    assert export_service._message_branch_label("VILLAS DEL REY") == "Villas del Rey"


def test_missing_message_branch_label_fails_closed():
    with pytest.raises(MarketingReactivationValidationError) as exc_info:
        export_service._resolve_branch_message_label("SEND MXL", {})

    assert "catálogo de sucursales" in str(exc_info.value)


def test_export_package_never_rebuilds_campaign_audience(monkeypatch):
    monkeypatch.setattr(
        export_service,
        "get_marketing_reactivation_campaign",
        lambda **kwargs: {
            "id": 11,
            "name": "Prueba",
            "recipients": [
                _recipient("6861000001", "CENTRO", "ANA LOPEZ"),
            ],
        },
    )
    _mock_branch_labels(monkeypatch, {"CENTRO": "Centro"})
    called = []
    monkeypatch.setattr(
        export_service,
        "_validate_and_mark_exported",
        lambda **kwargs: called.append(kwargs) or (b"ignored", "ignored.xlsx"),
    )

    export_service.export_marketing_reactivation_campaign(
        campaign_id=11,
        session=NS(),
    )

    assert len(called) == 1
    assert called[0]["campaign_id"] == 11
