"""Exportación XLSX de Salud de Sistemas, usando exclusivamente servicios BI existentes."""

from __future__ import annotations

from datetime import date, datetime, timezone
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.system_daily_check_bi_service import (
    build_system_daily_check_bi_summary,
    get_system_daily_check_bi_detail,
    list_system_daily_check_bi_history,
    list_system_daily_check_bi_issues,
)
from app.services.system_daily_check_service import (
    SystemDailyCheckValidationError,
)

LOCAL_TZ = ZoneInfo("America/Tijuana")
HISTORY_PAGE_SIZE = 100
ISSUES_PAGE_SIZE = 200
MAX_EXPORT_DAYS = 93
MAX_EXPORT_CHECKS = 2500
NAVY = "172033"
BLUE = "243B59"
LIGHT = "EDF2F7"
WHITE = "FFFFFF"
DARK = "344054"
STATUS_COLORS = {
    "NORMAL": "DCFCE7",
    "MINOR_FAILURE": "FEF3C7",
    "OPERATIONAL_IMPACT": "FEE2E2",
    "YES": "DCFCE7",
    "NO": "FEE2E2",
    "NA": "F1F5F9",
}


def _safe(value):
    """Evitar fórmulas inyectadas y caracteres inválidos en texto de usuarios."""
    if not isinstance(value, str):
        return value
    cleaned = ILLEGAL_CHARACTERS_RE.sub("", value)
    if cleaned.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + cleaned
    return cleaned


def _date(value):
    return date.fromisoformat(value) if value else None


def _local_datetime(value):
    if not value:
        return None
    timestamp = datetime.fromisoformat(value.replace("Z", "+00:00"))
    if timestamp.tzinfo is None:
        timestamp = timestamp.replace(tzinfo=timezone.utc)
    return timestamp.astimezone(LOCAL_TZ).replace(tzinfo=None)


def _status(value):
    return {
        "NORMAL": "Sin fallas",
        "MINOR_FAILURE": "Falla menor",
        "OPERATIONAL_IMPACT": "Afecta la operación",
        "YES": "Sí",
        "NO": "No",
        "NA": "No aplica",
    }.get(value, value)


def _table_sheet(
    book,
    title: str,
    headers: tuple[str, ...],
    rows: list[tuple],
    widths: tuple[int, ...],
    note: str,
    *,
    color_columns: tuple[int, ...] = (),
):
    ws = book.create_sheet(title)
    size = len(headers)
    final_column = get_column_letter(size)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=size)
    title_cell = ws.cell(1, 1, title.upper())
    title_cell.fill = PatternFill("solid", fgColor=NAVY)
    title_cell.font = Font(name="Aptos Display", size=16, bold=True, color=WHITE)
    title_cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=size)
    ws.cell(2, 1, _safe(note)).font = Font(
        name="Aptos", size=10, italic=True, color=DARK,
    )
    ws.row_dimensions[2].height = 29

    for col, label in enumerate(headers, 1):
        cell = ws.cell(4, col, label)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 30

    for row_idx, record in enumerate(rows, 5):
        for col_idx, value in enumerate(record, 1):
            cell = ws.cell(row_idx, col_idx, _safe(value))
            cell.font = Font(name="Aptos", size=10, color=DARK)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if row_idx % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FAFC")
            if isinstance(value, date) and not isinstance(value, datetime):
                cell.number_format = "dd/mm/yyyy"
            elif isinstance(value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm"
            if col_idx in color_columns and value in STATUS_COLORS:
                cell.fill = PatternFill(
                    "solid", fgColor=STATUS_COLORS[value],
                )
        ws.row_dimensions[row_idx].height = 23

    for col_idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{final_column}{max(4, ws.max_row)}"
    ws.sheet_view.showGridLines = False
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    return ws


def _collect_pages(fetch_page, *, size: int) -> list[dict]:
    page = 1
    items: list[dict] = []
    while True:
        payload = fetch_page(page, size)
        batch = payload["items"]
        items.extend(batch)
        if len(items) >= payload["total"] or not batch:
            return items
        page += 1


def export_system_daily_check_bi_workbook(
    actor,
    *,
    date_from: date,
    date_to: date,
    branch_id: int | None = None,
) -> BytesIO:
    """Genera un reporte completo; no recalcula métricas de negocio."""
    if not isinstance(date_from, date) or not isinstance(date_to, date):
        raise SystemDailyCheckValidationError("Rango de fechas inválido.")
    if date_from > date_to:
        raise SystemDailyCheckValidationError(
            "date_from no puede ser posterior a date_to."
        )
    if (date_to - date_from).days >= MAX_EXPORT_DAYS:
        raise SystemDailyCheckValidationError(
            "La exportación admite un máximo de 93 días por archivo."
        )

    options = {
        "date_from": date_from,
        "date_to": date_to,
        "branch_id": branch_id,
    }
    # Estos servicios aplican permisos, exclusiones y canonicalidad del rollout.
    summary = build_system_daily_check_bi_summary(actor, **options)
    first_history = list_system_daily_check_bi_history(
        actor, **options, page=1, page_size=HISTORY_PAGE_SIZE,
    )
    if first_history["total"] > MAX_EXPORT_CHECKS:
        raise SystemDailyCheckValidationError(
            "El rango supera 2500 checklists; acota fechas o sucursal."
        )

    history = list(first_history["items"])
    if len(history) < first_history["total"]:
        history.extend(
            _collect_pages(
                lambda page, size: list_system_daily_check_bi_history(
                    actor, **options, page=page, page_size=size,
                ),
                size=HISTORY_PAGE_SIZE,
            )[len(history):]
        )

    issues = _collect_pages(
        lambda page, size: list_system_daily_check_bi_issues(
            actor, **options, page=page, page_size=size,
        ),
        size=ISSUES_PAGE_SIZE,
    )
    # No se vuelven a construir respuestas o incidencias: se toma el detalle
    # exacto del módulo (incluye snapshot de pregunta y evidencias).
    details = [
        get_system_daily_check_bi_detail(actor, check_id=item["id"])
        for item in history
    ]
    scope = (
        f"Sucursal ID {branch_id}" if branch_id is not None
        else "Todas las sucursales participantes"
    )
    dates = f"Periodo: {date_from:%d/%m/%Y} al {date_to:%d/%m/%Y} · {scope}"

    book = Workbook()
    book.remove(book.active)

    metrics = [
        ("Sucursales esperadas", summary["universe"]["expected_branches"],
         "Sucursales habilitadas en el periodo."),
        ("Checklists esperados", summary["universe"]["expected_checklists"],
         "Días-sucursal con rollout habilitado."),
        ("Checklists completados", summary["summary"]["completed"],
         "Capturas dentro del rollout."),
        ("Pendientes", summary["summary"]["pending"],
         "Esperados aún sin captura."),
        ("Cumplimiento (%)", summary["summary"]["compliance_pct"],
         "Porcentaje; sin universo se deja vacío."),
        ("Sin fallas", summary["summary"]["normal"], ""),
        ("Falla menor", summary["summary"]["minor_failure"], ""),
        ("Afecta la operación", summary["summary"]["operational_impact"], ""),
        ("Checklists con falla", summary["failures"]["checks_with_failure"], ""),
        ("Respuestas negativas", summary["failures"]["no_answers"], ""),
        ("Incidencias", summary["failures"]["issues_total"], ""),
        ("Reportadas a soporte", summary["failures"]["reported"], ""),
        ("No reportadas a soporte", summary["failures"]["unreported"], ""),
        ("Incidencias reportadas (%)", summary["failures"]["reported_pct"], ""),
        ("Completados sin posponer",
         summary["postponements"]["completed_without_postpone"], ""),
        ("Completados tras 1 posposición",
         summary["postponements"]["completed_after_1"], ""),
        ("Completados tras 2+ posposiciones",
         summary["postponements"]["completed_after_2"], ""),
        ("Llegaron a obligatoriedad",
         summary["postponements"]["reached_mandatory"], ""),
        ("Capturas fuera del rollout",
         summary["universe"]["out_of_rollout_submissions"],
         "Pueden figurar en Historial, pero no suman al cumplimiento."),
        ("Capturas visibles en historial", first_history["total"],
         "Incluye capturas fuera del rollout (igual que la pantalla BI)."),
    ]
    _table_sheet(
        book, "Resumen", ("Indicador", "Valor", "Interpretación"),
        metrics, (39, 23, 75), dates,
    )
    for i, (label, _, _) in enumerate(metrics, 5):
        if label.endswith("(%)"):
            book["Resumen"].cell(i, 2).number_format = '0.0"%"'

    _table_sheet(
        book, "Checklists",
        ("ID", "Fecha operativa", "Sucursal", "Responsable",
         "Estado", "Enviado (Tijuana)", "Posposiciones", "Obligatorio"),
        [
            (
                item["id"], _date(item["business_date"]),
                item["sucursal"], item["performed_by_username"],
                _status(item["general_status"]),
                _local_datetime(item["submitted_at"]),
                item["postpone_count"],
                "Sí" if item["reached_mandatory"] else "No",
            )
            for item in history
        ],
        (12, 18, 32, 27, 26, 24, 17, 17),
        dates + " · Historial de capturas, incluidas fuera del rollout.",
    )
    _table_sheet(
        book, "Incidencias",
        ("ID incidencia", "ID checklist", "Fecha operativa", "Sucursal",
         "Pregunta", "Alcance afectado", "Descripción",
         "Reportada a soporte", "Evidencias"),
        [
            (
                item["issue_id"], item["check_id"],
                _date(item["business_date"]), item["sucursal"],
                item["question_label"], item["affected_scope"],
                item["description"],
                "Sí" if item["reported_to_support"] else "No",
                item["attachment_count"],
            )
            for item in issues
        ],
        (17, 17, 20, 29, 42, 32, 65, 23, 16),
        dates + " · Incidencias dentro del rollout habilitado.",
    )
    answers = []
    for check in details:
        for answer in check["answers"]:
            issue = answer.get("issue") or {}
            answers.append((
                check["id"], _date(check["business_date"]),
                check["sucursal"], check["performed_by_username"],
                answer["category_key"], answer["question_label"],
                _status(answer["answer"]),
                issue.get("description"),
                issue.get("affected_scope"),
                (
                    "Sí" if issue.get("reported_to_support") else "No"
                ) if issue else None,
                len(issue.get("attachments") or []),
            ))
    _table_sheet(
        book, "Respuestas",
        ("ID checklist", "Fecha operativa", "Sucursal", "Responsable",
         "Categoría", "Pregunta", "Respuesta", "Descripción de falla",
         "Alcance afectado", "Reportada a soporte", "Evidencias"),
        answers, (17, 20, 29, 25, 21, 46, 17, 64, 30, 24, 15),
        dates + " · Todas las respuestas de capturas visibles.",
    )
    book.active = 0
    buffer = BytesIO()
    book.save(buffer)
    buffer.seek(0)
    return buffer
