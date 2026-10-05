"""Exports frozen Campaign V2 recipients for manual delivery.

The export is a read-only projection of the frozen Campaign V2 cohort.
Recipients are split by branch and audience family (commercial segment).
One branch/segment produces one XLSX; multiple combinations produce a ZIP
with one XLSX per combination plus a summary workbook.
"""

from __future__ import annotations

from collections import defaultdict
from io import BytesIO
import re
import unicodedata
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook

from app.extensions import db
from app.models.marketing import MarketingCampaignV2RecipientORM
from app.services.marketing_campaign_v2_query_service import get_campaign_v2


XLSX_MIMETYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
ZIP_MIMETYPE = "application/zip"

_FAMILY_LABELS = {
    "DOMICILIADO": "DOMICILIADO",
    "TRIMESTRAL": "TRIMESTRAL",
    "CONVENIO": "CONVENIO",
    "SEMESTRE": "SEMESTRE",
    "ESTUDIANTE": "ESTUDIANTE",
    "MES": "MES",
    "OUT_OF_SEGMENT": "FUERA_DE_SEGMENTO",
}


def campaign_v2_delivery_export_mimetype(filename: str) -> str:
    return ZIP_MIMETYPE if str(filename).lower().endswith(".zip") else XLSX_MIMETYPE


def export_campaign_v2_delivery_package(
    *,
    campaign_id: int,
    allowed_sucursal_keys,
    session: Any | None = None,
) -> tuple[bytes, str]:
    """Build a manual-delivery package strictly from frozen recipients."""

    active_session = session if session is not None else db.session
    campaign = get_campaign_v2(
        campaign_id=campaign_id,
        allowed_sucursal_keys=allowed_sucursal_keys,
        session=active_session,
    )

    recipients = (
        active_session.query(MarketingCampaignV2RecipientORM)
        .filter(MarketingCampaignV2RecipientORM.campaign_id == int(campaign_id))
        .order_by(
            MarketingCampaignV2RecipientORM.sucursal.asc(),
            MarketingCampaignV2RecipientORM.audience_family.asc(),
            MarketingCampaignV2RecipientORM.phone_mx10.asc(),
            MarketingCampaignV2RecipientORM.id.asc(),
        )
        .all()
    )

    groups: dict[tuple[str, str], list[MarketingCampaignV2RecipientORM]] = defaultdict(list)
    for recipient in recipients:
        branch = _clean_text(recipient.sucursal) or "SIN SUCURSAL"
        family = _clean_text(recipient.audience_family) or "SIN_CLASIFICAR"
        groups[(branch, family)].append(recipient)

    campaign_part = _filename_part(
        campaign.get("name"),
        fallback=f"CAMPANA_V2_{int(campaign_id)}",
    )
    return _build_delivery_package(
        campaign_part=campaign_part,
        groups=groups,
    )


def _build_delivery_package(
    *,
    campaign_part: str,
    groups: dict[tuple[str, str], list[Any]],
) -> tuple[bytes, str]:
    ordered_groups = sorted(
        groups.items(),
        key=lambda item: (
            item[0][0].casefold(),
            item[0][1].casefold(),
        ),
    )

    if len(ordered_groups) == 1:
        (branch, family), recipients = ordered_groups[0]
        branch_part = _filename_part(branch, fallback="SIN_SUCURSAL")
        family_part = _family_filename_part(family)
        return (
            _recipients_workbook(recipients),
            f"{campaign_part}__{branch_part}__{family_part}.xlsx",
        )

    archive_output = BytesIO()
    used_names: set[str] = set()
    with ZipFile(archive_output, "w", compression=ZIP_DEFLATED) as archive:
        for (branch, family), recipients in ordered_groups:
            branch_part = _filename_part(branch, fallback="SIN_SUCURSAL")
            family_part = _family_filename_part(family)
            entry_name = _unique_archive_name(
                f"{campaign_part}__{branch_part}__{family_part}.xlsx",
                used_names,
            )
            archive.writestr(entry_name, _recipients_workbook(recipients))
        archive.writestr("RESUMEN.xlsx", _summary_workbook(ordered_groups))

    return archive_output.getvalue(), f"{campaign_part}.zip"


def _recipients_workbook(recipients: list[Any]) -> bytes:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Destinatarios")
    sheet.append(["telefono", "nombre", "sucursal", "tarifa"])

    ordered = sorted(
        recipients,
        key=lambda row: (
            str(getattr(row, "phone_mx10", "") or ""),
            str(getattr(row, "member_name", "") or "").casefold(),
        ),
    )
    for recipient in ordered:
        sheet.append(
            [
                _safe_excel_text(getattr(recipient, "phone_mx10", None)),
                _safe_excel_text(getattr(recipient, "member_name", None)),
                _safe_excel_text(getattr(recipient, "sucursal", None)),
                _safe_excel_text(getattr(recipient, "tarifa_raw", None)),
            ]
        )

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _summary_workbook(
    groups: list[tuple[tuple[str, str], list[Any]]],
) -> bytes:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Resumen")
    sheet.append(["sucursal", "segmento", "contactos"])

    total = 0
    for (branch, family), recipients in groups:
        count = len(recipients)
        total += count
        sheet.append([branch, _family_display_label(family), count])
    sheet.append(["TOTAL", "", total])

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _family_display_label(value: Any) -> str:
    normalized = _clean_text(value).upper()
    if normalized == "OUT_OF_SEGMENT":
        return "Fuera de segmento"
    if normalized == "SIN_CLASIFICAR":
        return "Sin clasificar"
    return {
        "DOMICILIADO": "Domiciliado",
        "TRIMESTRAL": "Trimestral",
        "CONVENIO": "Convenio",
        "SEMESTRE": "Semestre",
        "ESTUDIANTE": "Estudiante",
        "MES": "Mes",
    }.get(normalized, _clean_text(value) or "Sin clasificar")


def _family_filename_part(value: Any) -> str:
    normalized = _clean_text(value).upper()
    return _FAMILY_LABELS.get(
        normalized,
        _filename_part(normalized, fallback="SIN_CLASIFICAR"),
    )


def _clean_text(value: Any) -> str:
    return " ".join(str(value or "").split())


def _safe_excel_text(value: Any) -> str:
    text = _clean_text(value)
    probe = text.lstrip(" \t\r\n\x00")
    if probe.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _filename_part(value: Any, *, fallback: str, max_length: int = 80) -> str:
    normalized = unicodedata.normalize("NFKD", str(value or ""))
    ascii_value = normalized.encode("ascii", "ignore").decode("ascii").upper()
    cleaned = re.sub(r"[^A-Z0-9]+", "_", ascii_value).strip("_")
    cleaned = re.sub(r"_+", "_", cleaned)
    return (cleaned[:max_length].rstrip("_") or fallback)


def _unique_archive_name(candidate: str, used_names: set[str]) -> str:
    if candidate not in used_names:
        used_names.add(candidate)
        return candidate

    stem, suffix = candidate.rsplit(".", 1)
    index = 2
    while True:
        resolved = f"{stem}_{index}.{suffix}"
        if resolved not in used_names:
            used_names.add(resolved)
            return resolved
        index += 1
