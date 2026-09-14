from __future__ import annotations

from calendar import monthrange
from dataclasses import dataclass
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.warehouse import VentaTotalSnapshotRowORM
from app.services.marketing_access import MarketingAccess
from app.services.marketing_attribution import SaleRecord, reconcile_visit_sales
from app.services.marketing_dashboard_service import (
    _load_sales as _load_attribution_sales,
    _load_visit_events as _load_attribution_visits,
    load_visible_marketing_branches,
)
from app.services.marketing_inputs_service import parse_month
from app.services.marketing_sales_funnel_detail_service import (
    MarketingSalesFunnelDetailValidationError,
)
from app.services.marketing_sales_funnel_service import (
    MATCH_WINDOW_DAYS,
    ORIGIN_IVENTAS_META,
    ORIGIN_IVENTAS_OTHER,
    ORIGIN_LABELS,
    _load_branch_alias_map,
    _load_iventas_data,
    _load_visits,
    _match_iventas,
    _select_venta_total_snapshot,
)


VISIT_CONVERSION_METRICS = frozenset(
    {
        "visits_iventas_bought",
        "visits_iventas_not_bought",
        "visits_not_iventas_bought",
        "visits_not_iventas_not_bought",
    }
)

VISIT_CONVERSION_TITLES = {
    "visits_iventas_bought": "Visitas iVentas que compraron",
    "visits_iventas_not_bought": "Visitas iVentas que no compraron",
    "visits_not_iventas_bought": "Visitas sin trazabilidad que compraron",
    "visits_not_iventas_not_bought": "Visitas sin trazabilidad que no compraron",
}

DETAIL_EXPORT_MIMETYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
DEFAULT_PAGE_SIZE = 50
MAX_PAGE_SIZE = 100


@dataclass(frozen=True)
class VisitConversionRow:
    event_key: str
    branch_id: int
    visit_date: date
    phone: str | None
    origin: str | None
    sale: SaleRecord | None

    @property
    def is_iventas(self) -> bool:
        return self.origin in {ORIGIN_IVENTAS_META, ORIGIN_IVENTAS_OTHER}

    @property
    def bought(self) -> bool:
        return self.sale is not None


@dataclass(frozen=True)
class VisitConversionBundle:
    rows: tuple[VisitConversionRow, ...]
    sales_snapshot_ids: tuple[int, ...]
    cohort_complete: bool


def _month_end(month_start: date) -> date:
    return date(
        month_start.year,
        month_start.month,
        monthrange(month_start.year, month_start.month)[1],
    )


def _normalize_optional_int(value: Any, *, field_name: str) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(value)
    except (TypeError, ValueError) as exc:
        raise MarketingSalesFunnelDetailValidationError(
            f"{field_name} inválido."
        ) from exc


def _normalize_page(value: Any, default: int) -> int:
    if value in (None, ""):
        return default
    try:
        return int(value)
    except (TypeError, ValueError) as exc:
        raise MarketingSalesFunnelDetailValidationError(
            "Paginación inválida."
        ) from exc


def _normalize_sort(
    sort_by: Any,
    sort_dir: Any,
) -> tuple[str | None, str]:
    normalized_sort_by = str(sort_by or "").strip() or None
    normalized_sort_dir = str(sort_dir or "asc").strip().lower() or "asc"
    allowed = {
        "branch",
        "date",
        "phone",
        "origin",
        "conversion_status",
        "sale_date",
        "sale_member_id",
        "sale_revenue",
    }
    if normalized_sort_dir not in {"asc", "desc"}:
        raise MarketingSalesFunnelDetailValidationError(
            "sort_dir debe ser asc o desc."
        )
    if normalized_sort_by is not None and normalized_sort_by not in allowed:
        raise MarketingSalesFunnelDetailValidationError(
            "sort_by no soportado para detalle de conversión de visitas."
        )
    return normalized_sort_by, normalized_sort_dir


def _sortable_value(value: Any) -> tuple[int, Any]:
    if value is None or str(value).strip() == "":
        return 1, ""
    if isinstance(value, (int, float, Decimal)):
        return 0, Decimal(str(value))
    text = str(value).strip()
    try:
        return 0, Decimal(text.replace(",", ""))
    except (InvalidOperation, ValueError):
        return 0, text.casefold()


def _metric_matches(row: VisitConversionRow, metric: str) -> bool:
    if metric == "visits_iventas_bought":
        return row.is_iventas and row.bought
    if metric == "visits_iventas_not_bought":
        return row.is_iventas and not row.bought
    if metric == "visits_not_iventas_bought":
        return not row.is_iventas and row.bought
    if metric == "visits_not_iventas_not_bought":
        return not row.is_iventas and not row.bought
    return False


def _build_bundle(
    *,
    month_start: date,
    branch_ids: tuple[int, ...],
    today: date | None = None,
) -> VisitConversionBundle:
    snapshot = _select_venta_total_snapshot(month_start)
    if snapshot is None:
        return VisitConversionBundle(
            rows=(),
            sales_snapshot_ids=(),
            cohort_complete=False,
        )

    venta_total_rows = (
        VentaTotalSnapshotRowORM.query.filter_by(snapshot_id=snapshot.id)
        .order_by(VentaTotalSnapshotRowORM.row_index.asc())
        .all()
    )
    visits = _load_visits(
        venta_total_rows,
        month_start,
        branch_ids,
        _load_branch_alias_map(),
    )

    evidence, _, _ = _load_iventas_data(month_start, branch_ids)
    attribution_visit_result = _load_attribution_visits(
        month_start=month_start,
        branch_ids=branch_ids,
    )
    attribution_window_end = _month_end(month_start) + timedelta(
        days=MATCH_WINDOW_DAYS
    )
    attribution_sales_result = _load_attribution_sales(
        window_start=month_start,
        window_end=attribution_window_end,
        branch_ids=branch_ids,
    )
    attributions = reconcile_visit_sales(
        visits=attribution_visit_result.events,
        sales=attribution_sales_result.sales,
    )

    first_sale_by_identity: dict[tuple[int, str], SaleRecord] = {}
    for attribution in sorted(
        attributions,
        key=lambda item: (
            item.sale.payment_date,
            item.sale.sale_key,
        ),
    ):
        identity = (
            int(attribution.visit.branch_id),
            str(attribution.visit.phone),
        )
        first_sale_by_identity.setdefault(identity, attribution.sale)

    result: list[VisitConversionRow] = []
    for visit in visits:
        origin = _match_iventas(
            evidence,
            visit.branch_id,
            visit.phone,
            visit.visit_date,
        )
        sale = (
            first_sale_by_identity.get((visit.branch_id, visit.phone))
            if visit.phone is not None
            else None
        )
        result.append(
            VisitConversionRow(
                event_key=visit.event_key,
                branch_id=visit.branch_id,
                visit_date=visit.visit_date,
                phone=visit.phone,
                origin=origin,
                sale=sale,
            )
        )

    normalized_today = today or date.today()
    return VisitConversionBundle(
        rows=tuple(result),
        sales_snapshot_ids=tuple(
            int(snapshot_id)
            for snapshot_id in attribution_sales_result.snapshot_ids
        ),
        cohort_complete=normalized_today >= attribution_window_end,
    )


def _empty_metrics() -> dict[str, Any]:
    return {
        "visits_iventas_bought": 0,
        "visits_iventas_not_bought": 0,
        "visits_not_iventas_bought": 0,
        "visits_not_iventas_not_bought": 0,
        "iventas_visit_conversion_rate": None,
        "not_iventas_visit_conversion_rate": None,
    }


def _serialize_metrics(rows: list[VisitConversionRow]) -> dict[str, Any]:
    metrics = _empty_metrics()
    for row in rows:
        if row.is_iventas:
            key = (
                "visits_iventas_bought"
                if row.bought
                else "visits_iventas_not_bought"
            )
        else:
            key = (
                "visits_not_iventas_bought"
                if row.bought
                else "visits_not_iventas_not_bought"
            )
        metrics[key] += 1

    iventas_total = (
        metrics["visits_iventas_bought"]
        + metrics["visits_iventas_not_bought"]
    )
    not_iventas_total = (
        metrics["visits_not_iventas_bought"]
        + metrics["visits_not_iventas_not_bought"]
    )
    metrics["iventas_visit_conversion_rate"] = (
        metrics["visits_iventas_bought"] / iventas_total
        if iventas_total > 0
        else None
    )
    metrics["not_iventas_visit_conversion_rate"] = (
        metrics["visits_not_iventas_bought"] / not_iventas_total
        if not_iventas_total > 0
        else None
    )
    return metrics


def build_visit_conversion_summary(
    *,
    month_start: date,
    branch_ids: tuple[int, ...],
) -> dict[str, Any]:
    bundle = _build_bundle(
        month_start=month_start,
        branch_ids=branch_ids,
    )
    rows_by_branch: dict[int, list[VisitConversionRow]] = {
        branch_id: [] for branch_id in branch_ids
    }
    for row in bundle.rows:
        rows_by_branch.setdefault(row.branch_id, []).append(row)

    return {
        "summary": _serialize_metrics(list(bundle.rows)),
        "branches": {
            branch_id: _serialize_metrics(rows_by_branch.get(branch_id, []))
            for branch_id in branch_ids
        },
        "data_quality": {
            "visit_conversion_mode": "exact_phone_same_branch_30d",
            "visit_conversion_cohort_complete": bundle.cohort_complete,
            "visit_conversion_sales_snapshot_ids": list(
                bundle.sales_snapshot_ids
            ),
        },
    }


def _branch_names(access: MarketingAccess) -> tuple[dict[int, str], tuple[int, ...], dict[str, object]]:
    branches, branch_ids, scope = load_visible_marketing_branches(access)
    return (
        {int(branch.sucursal_id): str(branch.name) for branch in branches},
        branch_ids,
        scope,
    )


def _serialize_detail_row(
    row: VisitConversionRow,
    branch_names: dict[int, str],
) -> dict[str, Any]:
    sale = row.sale
    return {
        "branch_id": row.branch_id,
        "branch": branch_names.get(row.branch_id, ""),
        "date": row.visit_date.isoformat(),
        "phone": row.phone,
        "origin_key": row.origin,
        "origin": (
            ORIGIN_LABELS.get(row.origin, row.origin)
            if row.origin is not None
            else "Sin match iVentas"
        ),
        "source": "Pase comercial en Venta Total",
        "conversion_status": "Compró" if sale is not None else "No compró",
        "sale_date": sale.payment_date.isoformat() if sale is not None else None,
        "sale_member_id": sale.member_id if sale is not None else None,
        "sale_revenue": float(sale.revenue) if sale is not None else None,
    }


def build_visit_conversion_detail(
    *,
    month: str,
    access: MarketingAccess,
    metric: str,
    branch_id: Any = None,
    page: Any = None,
    page_size: Any = None,
    sort_by: Any = None,
    sort_dir: Any = None,
) -> dict[str, Any]:
    if metric not in VISIT_CONVERSION_METRICS:
        raise MarketingSalesFunnelDetailValidationError(
            "metric no soportado para conversión de visitas."
        )

    month_start = parse_month(month)
    branch_names, branch_ids, scope = _branch_names(access)
    normalized_branch_id = _normalize_optional_int(
        branch_id,
        field_name="branch_id",
    )
    if normalized_branch_id is not None and normalized_branch_id not in set(branch_ids):
        raise MarketingSalesFunnelDetailValidationError(
            "La sucursal solicitada no pertenece al alcance del usuario."
        )

    normalized_page = _normalize_page(page, 1)
    normalized_page_size = _normalize_page(page_size, DEFAULT_PAGE_SIZE)
    if normalized_page < 1:
        raise MarketingSalesFunnelDetailValidationError(
            "page debe ser mayor o igual a 1."
        )
    if normalized_page_size < 1 or normalized_page_size > MAX_PAGE_SIZE:
        raise MarketingSalesFunnelDetailValidationError(
            f"page_size debe estar entre 1 y {MAX_PAGE_SIZE}."
        )
    normalized_sort_by, normalized_sort_dir = _normalize_sort(
        sort_by,
        sort_dir,
    )

    bundle = _build_bundle(
        month_start=month_start,
        branch_ids=branch_ids,
    )
    selected = [
        row
        for row in bundle.rows
        if _metric_matches(row, metric)
        and (
            normalized_branch_id is None
            or row.branch_id == normalized_branch_id
        )
    ]
    serialized = [
        _serialize_detail_row(row, branch_names)
        for row in selected
    ]
    if normalized_sort_by is not None:
        populated = [
            row
            for row in serialized
            if row.get(normalized_sort_by) not in (None, "")
        ]
        empty = [
            row
            for row in serialized
            if row.get(normalized_sort_by) in (None, "")
        ]
        populated.sort(
            key=lambda row: _sortable_value(row.get(normalized_sort_by)),
            reverse=normalized_sort_dir == "desc",
        )
        serialized = populated + empty

    count = len(serialized)
    start = (normalized_page - 1) * normalized_page_size
    end = start + normalized_page_size
    total_pages = max(1, (count + normalized_page_size - 1) // normalized_page_size)
    title = VISIT_CONVERSION_TITLES[metric]
    if normalized_branch_id is not None:
        title = f"{title} · {branch_names.get(normalized_branch_id, normalized_branch_id)}"

    return {
        "month": month_start.strftime("%Y-%m"),
        "scope": scope,
        "metric": metric,
        "origin": None,
        "kind": "visits",
        "title": title,
        "branch_id": normalized_branch_id,
        "count": count,
        "revenue_total": 0,
        "page": normalized_page,
        "page_size": normalized_page_size,
        "total_pages": total_pages,
        "sort_by": normalized_sort_by,
        "sort_dir": normalized_sort_dir,
        "rows": serialized[start:end],
    }


def build_visit_conversion_export(
    *,
    month: str,
    access: MarketingAccess,
    metric: str,
    branch_id: Any = None,
    sort_by: Any = None,
    sort_dir: Any = None,
) -> tuple[BytesIO, str]:
    detail = build_visit_conversion_detail(
        month=month,
        access=access,
        metric=metric,
        branch_id=branch_id,
        page=1,
        page_size=MAX_PAGE_SIZE,
        sort_by=sort_by,
        sort_dir=sort_dir,
    )

    # Rebuild without the API page limit for export.
    month_start = parse_month(month)
    branch_names, branch_ids, _ = _branch_names(access)
    normalized_branch_id = _normalize_optional_int(
        branch_id,
        field_name="branch_id",
    )
    bundle = _build_bundle(month_start=month_start, branch_ids=branch_ids)
    rows = [
        _serialize_detail_row(row, branch_names)
        for row in bundle.rows
        if _metric_matches(row, metric)
        and (
            normalized_branch_id is None
            or row.branch_id == normalized_branch_id
        )
    ]
    normalized_sort_by, normalized_sort_dir = _normalize_sort(sort_by, sort_dir)
    if normalized_sort_by is not None:
        rows.sort(
            key=lambda row: _sortable_value(row.get(normalized_sort_by)),
            reverse=normalized_sort_dir == "desc",
        )

    columns = (
        ("branch", "Sucursal KPI"),
        ("date", "Fecha visita"),
        ("phone", "Teléfono"),
        ("origin", "Trazabilidad"),
        ("conversion_status", "Conversión"),
        ("sale_date", "Fecha venta"),
        ("sale_member_id", "ID socio"),
        ("sale_revenue", "Ingreso venta"),
        ("source", "Fuente"),
    )
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Detalle"
    worksheet.append([label for _, label in columns])
    header_fill = PatternFill("solid", fgColor="211F1E")
    header_font = Font(color="FFFFFF", bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    for row in rows:
        worksheet.append([row.get(field) for field, _ in columns])

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        f"A1:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
    )
    for index, (field, label) in enumerate(columns, start=1):
        max_length = len(label)
        for cell in worksheet[get_column_letter(index)][1:]:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        worksheet.column_dimensions[get_column_letter(index)].width = min(
            max(max_length + 2, 12),
            42,
        )
        if field == "sale_revenue":
            for cell in worksheet[get_column_letter(index)][1:]:
                cell.number_format = '$' + '#,##0.00'

    worksheet.sheet_view.showGridLines = False
    workbook.properties.title = detail["title"]
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    filename = f"funnel_venta_nueva_{month}_{metric}"
    if normalized_branch_id is not None:
        filename += f"_sucursal_{normalized_branch_id}"
    return output, filename + ".xlsx"
