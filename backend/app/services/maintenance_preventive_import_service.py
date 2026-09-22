# backend/app/services/maintenance_preventive_import_service.py

from __future__ import annotations

import csv
import hashlib
import io
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

from app.services.maintenance_preventive_service import MaintenancePreventiveError


TEMPLATE_HEADERS = (
    "Sucursal",
    "Familia",
    "Código equipo",
    "Fecha programada",
    "Actividad",
    "Responsable",
    "Se repite",
    "Cada N días hábiles",
    "Observaciones",
    "Tipo",
    "Clasificación edificio",
    "Validación",
)

REQUIRED_IMPORT_HEADERS = (
    "Sucursal",
    "Fecha programada",
    "Actividad",
    "Responsable",
)

MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 1000


def _normalize_header(value) -> str:
    raw = str(value or "").strip().casefold()
    replacements = {
        "á": "a",
        "é": "e",
        "í": "i",
        "ó": "o",
        "ú": "u",
        "ü": "u",
    }
    for source, target in replacements.items():
        raw = raw.replace(source, target)
    return " ".join(raw.replace("_", " ").split())


HEADER_MAP = {
    "sucursal": "sucursal",
    "familia": "familia",
    "codigo equipo": "codigo_equipo",
    "fecha programada": "fecha_programada",
    "actividad": "actividad",
    "responsable": "responsable",
    "se repite": "repeat_enabled",
    "repetir": "repeat_enabled",
    "cada n dias habiles": "repeat_interval_workdays",
    "dias habiles": "repeat_interval_workdays",
    "observaciones": "observaciones",
    "tipo": "target_type",
    "tipo objetivo": "target_type",
    "clasificacion edificio": "building_classification",
    "clasificacion de edificio": "building_classification",
}


def _date_input(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if value is None:
        return None

    raw = str(value).strip()
    if not raw:
        return None

    for date_format in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw, date_format).date().isoformat()
        except ValueError:
            continue

    return raw


def _row_is_empty(values) -> bool:
    return not any(
        str(value).strip()
        for value in values
        if value is not None
    )


def _resolve_columns(headers) -> dict[int, str]:
    resolved: dict[int, str] = {}

    for index, value in enumerate(headers):
        key = _normalize_header(value)
        field = HEADER_MAP.get(key)
        if field:
            resolved[index] = field

    missing = [
        header
        for header in REQUIRED_IMPORT_HEADERS
        if HEADER_MAP[_normalize_header(header)] not in resolved.values()
    ]
    if missing:
        raise MaintenancePreventiveError(
            "La plantilla no contiene columnas obligatorias: "
            + ", ".join(missing)
        )

    return resolved


def _build_rows(raw_rows, columns: dict[int, str], *, start_row: int):
    result: list[dict] = []

    for row_number, values in enumerate(raw_rows, start=start_row):
        values = list(values)

        recognized_values = [
            values[column_index]
            for column_index in columns
            if column_index < len(values)
        ]
        if _row_is_empty(recognized_values):
            continue

        payload = {
            "source_row_number": row_number,
            "sucursal": None,
            "familia": None,
            "codigo_equipo": None,
            "fecha_programada": None,
            "actividad": None,
            "responsable": None,
            "repeat_enabled": None,
            "repeat_interval_workdays": None,
            "observaciones": None,
            "target_type": None,
            "building_classification": None,
        }

        for column_index, field in columns.items():
            value = values[column_index] if column_index < len(values) else None
            if field == "fecha_programada":
                payload[field] = _date_input(value)
            else:
                payload[field] = (
                    str(value).strip()
                    if value is not None and str(value).strip()
                    else None
                )

        result.append(payload)

        if len(result) > MAX_IMPORT_ROWS:
            raise MaintenancePreventiveError(
                f"La carga excede el máximo de {MAX_IMPORT_ROWS} renglones."
            )

    if not result:
        raise MaintenancePreventiveError(
            "La plantilla no contiene renglones de programación."
        )

    return result


def _parse_xlsx(content: bytes) -> list[dict]:
    try:
        workbook = load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise MaintenancePreventiveError(
            "No se pudo leer el archivo Excel."
        ) from exc

    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = next(iterator, None)
        if headers is None:
            raise MaintenancePreventiveError(
                "La plantilla Excel está vacía."
            )

        columns = _resolve_columns(headers)
        return _build_rows(iterator, columns, start_row=2)
    finally:
        workbook.close()


def _parse_csv(content: bytes) -> list[dict]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise MaintenancePreventiveError(
            "El CSV debe estar codificado en UTF-8."
        ) from exc

    reader = csv.reader(io.StringIO(text))
    headers = next(reader, None)
    if headers is None:
        raise MaintenancePreventiveError("El CSV está vacío.")

    columns = _resolve_columns(headers)
    return _build_rows(reader, columns, start_row=2)


def parse_preventive_import(
    *,
    filename: str,
    content: bytes,
) -> dict:
    if not filename:
        raise MaintenancePreventiveError(
            "El archivo debe tener nombre."
        )

    if not content:
        raise MaintenancePreventiveError(
            "El archivo está vacío."
        )

    if len(content) > MAX_IMPORT_BYTES:
        raise MaintenancePreventiveError(
            "El archivo excede el límite de 5 MB."
        )

    extension = Path(filename).suffix.casefold()
    if extension == ".xlsx":
        rows = _parse_xlsx(content)
    elif extension == ".csv":
        rows = _parse_csv(content)
    else:
        raise MaintenancePreventiveError(
            "Formato no soportado; usa .xlsx o .csv."
        )

    return {
        "filename": Path(filename).name,
        "sha256": hashlib.sha256(content).hexdigest(),
        "rows": rows,
    }


def _clean_catalog_values(values) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()

    for value in values or []:
        text = str(value or "").strip()
        if not text:
            continue

        key = text.casefold()
        if key in seen:
            continue

        seen.add(key)
        result.append(text)

    return result


def _add_named_list(
    workbook,
    *,
    name: str,
    sheet_name: str,
    column: str,
    start_row: int,
    end_row: int,
) -> None:
    if end_row < start_row:
        return

    sheet_ref = sheet_name.replace("'", "''")
    attr_text = (
        "'" + sheet_ref + "'!$"
        + column + "$" + str(start_row)
        + ":$" + column + "$" + str(end_row)
    )
    workbook.defined_names.add(
        DefinedName(name, attr_text=attr_text)
    )


def _equipment_catalog_tree(equipos: list[dict] | None) -> dict[str, dict[str, list[str]]]:
    tree: dict[str, dict[str, list[str]]] = {}

    for equipment in equipos or []:
        branch = str((equipment or {}).get("sucursal") or "").strip()
        code = str((equipment or {}).get("codigo_interno") or "").strip()
        family = str((equipment or {}).get("familia") or "").strip() or "Sin familia"

        if not branch or not code:
            continue

        branch_families = tree.setdefault(branch, {})
        codes = branch_families.setdefault(family, [])
        if code not in codes:
            codes.append(code)

    for branch_families in tree.values():
        for codes in branch_families.values():
            codes.sort(key=str.casefold)

    return tree


def build_preventive_template_xlsx(
    *,
    sucursales: list[str] | None = None,
    responsables: list[str] | None = None,
    equipos: list[dict] | None = None,
    building_classifications: list[str] | None = None,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Programacion preventiva"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:L1001"

    header_fill = PatternFill("solid", fgColor="E54525")
    header_font = Font(color="FFFFFF", bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="D8DEE8")
    input_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for column_index, header in enumerate(TEMPLATE_HEADERS, start=1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = input_border

    comments = {
        "A1": "Selecciona una sucursal del catálogo disponible.",
        "B1": "Selecciona una familia. La lista depende de la sucursal elegida.",
        "C1": (
            "Selecciona el código interno. La lista depende de sucursal y familia. "
            "Un mismo equipo no puede programarse dos veces para la misma fecha."
        ),
        "D1": "Captura la fecha en formato dd/mm/aaaa.",
        "E1": "Describe el mantenimiento preventivo que se realizará.",
        "F1": "Selecciona un responsable activo del catálogo de Mantenimiento.",
        "G1": "Selecciona Sí si este mantenimiento debe repetirse.",
        "H1": (
            "Si Se repite es Sí, captura cada cuántos días hábiles "
            "(lunes a viernes) debe repetirse."
        ),
        "I1": "Campo opcional para notas de programación.",
        "J1": (
            "Equipo o Edificio. Si queda vacío, se interpreta como Equipo "
            "para conservar compatibilidad con plantillas anteriores."
        ),
        "K1": (
            "Solo para Edificio. Selecciona una clasificación del árbol "
            "oficial de Mantenimiento > Edificio."
        ),
        "L1": (
            "Columna automática. Si aparece ⚠ DUPLICADO, corrige "
            "el objetivo o la fecha antes de cargar el archivo."
        ),
    }
    for cell_ref, message in comments.items():
        sheet[cell_ref].comment = Comment(message, "Suite Ultra")

    widths = {
        "A": 26,
        "B": 28,
        "C": 22,
        "D": 18,
        "E": 42,
        "F": 26,
        "G": 14,
        "H": 22,
        "I": 42,
        "J": 16,
        "K": 34,
        "L": 18,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    for row in range(2, MAX_IMPORT_ROWS + 2):
        for column in range(1, len(TEMPLATE_HEADERS) + 1):
            cell = sheet.cell(row=row, column=column)
            cell.border = input_border
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=(column in {5, 9, 11}),
            )
        sheet.cell(row=row, column=4).number_format = "dd/mm/yyyy"
        sheet.cell(
            row=row,
            column=12,
            value=(
                '=IF(OR(A' + str(row) + '="",D' + str(row) + '=""),"",'
                + 'IF(UPPER(J' + str(row) + ')="EDIFICIO",'
                + 'IF(K' + str(row) + '="","",'
                + 'IF(COUNTIFS($A$2:$A$1001,A' + str(row)
                + ',$K$2:$K$1001,K' + str(row)
                + ',$D$2:$D$1001,D' + str(row)
                + ',$J$2:$J$1001,"Edificio")>1,"⚠ DUPLICADO","")),'
                + 'IF(C' + str(row) + '="","",'
                + 'IF(COUNTIFS($A$2:$A$1001,A' + str(row)
                + ',$C$2:$C$1001,C' + str(row)
                + ',$D$2:$D$1001,D' + str(row)
                + ')>1,"⚠ DUPLICADO",""))))'
            ),
        )
        sheet.cell(row=row, column=12).font = Font(
            color="B91C1C",
            bold=True,
        )
        sheet.cell(row=row, column=12).alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    branch_values = _clean_catalog_values(sucursales)
    responsible_values = _clean_catalog_values(responsables)
    building_values = _clean_catalog_values(building_classifications)
    equipment_rows = list(equipos or [])
    equipment_tree = _equipment_catalog_tree(equipment_rows)

    catalogs = workbook.create_sheet("Catálogos")
    catalogs.sheet_view.showGridLines = False
    catalogs.freeze_panes = "A2"
    catalogs.append([
        "Sucursales",
        "Responsables",
        "",
        "Sucursal equipo",
        "Familia",
        "Código equipo",
        "Equipo",
        "Clasificaciones edificio",
    ])

    for cell in catalogs[1]:
        if cell.column == 3:
            continue
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = input_border

    max_catalog_rows = max(
        len(branch_values),
        len(responsible_values),
        len(equipment_rows),
        len(building_values),
        1,
    )
    for index in range(max_catalog_rows):
        row_number = index + 2
        if index < len(branch_values):
            catalogs.cell(row=row_number, column=1, value=branch_values[index])
        if index < len(responsible_values):
            catalogs.cell(row=row_number, column=2, value=responsible_values[index])
        if index < len(equipment_rows):
            equipment = equipment_rows[index] or {}
            catalogs.cell(row=row_number, column=4, value=str(equipment.get("sucursal") or ""))
            catalogs.cell(row=row_number, column=5, value=str(equipment.get("familia") or "") or "Sin familia")
            catalogs.cell(row=row_number, column=6, value=str(equipment.get("codigo_interno") or ""))
            catalogs.cell(row=row_number, column=7, value=str(equipment.get("nombre") or ""))
        if index < len(building_values):
            catalogs.cell(
                row=row_number,
                column=8,
                value=building_values[index],
            )

    catalog_widths = {
        "A": 28,
        "B": 28,
        "C": 3,
        "D": 28,
        "E": 28,
        "F": 22,
        "G": 34,
        "H": 38,
    }
    for column, width in catalog_widths.items():
        catalogs.column_dimensions[column].width = width

    equipment_end = max(len(equipment_rows) + 1, 2)
    catalogs.auto_filter.ref = "D1:G" + str(equipment_end)

    validations = workbook.create_sheet("Validaciones")
    validations.sheet_state = "hidden"
    validations["A1"] = "Sucursal"
    validations["B1"] = "Rango familias"
    validations["D1"] = "Sucursal|Familia"
    validations["E1"] = "Rango códigos"
    validations["G1"] = "Lista vacía"
    validations["G2"] = ""

    _add_named_list(
        workbook,
        name="ListaVacia",
        sheet_name=validations.title,
        column="G",
        start_row=2,
        end_row=2,
    )

    family_value_row = 2
    code_value_row = 2
    family_map_row = 2
    code_map_row = 2
    family_range_index = 1
    code_range_index = 1

    for branch in branch_values:
        branch_families = equipment_tree.get(branch, {})
        families = sorted(branch_families, key=str.casefold)
        if not families:
            continue

        family_range_name = "Familias_" + str(family_range_index).zfill(4)
        family_start = family_value_row
        for family in families:
            validations.cell(row=family_value_row, column=9, value=family)
            family_value_row += 1
        _add_named_list(
            workbook,
            name=family_range_name,
            sheet_name=validations.title,
            column="I",
            start_row=family_start,
            end_row=family_value_row - 1,
        )
        validations.cell(row=family_map_row, column=1, value=branch)
        validations.cell(row=family_map_row, column=2, value=family_range_name)
        family_map_row += 1
        family_range_index += 1

        for family in families:
            codes = branch_families.get(family, [])
            if not codes:
                continue

            code_range_name = "Codigos_" + str(code_range_index).zfill(4)
            code_start = code_value_row
            for code in codes:
                validations.cell(row=code_value_row, column=11, value=code)
                code_value_row += 1
            _add_named_list(
                workbook,
                name=code_range_name,
                sheet_name=validations.title,
                column="K",
                start_row=code_start,
                end_row=code_value_row - 1,
            )
            validations.cell(
                row=code_map_row,
                column=4,
                value=branch + "|" + family,
            )
            validations.cell(row=code_map_row, column=5, value=code_range_name)
            code_map_row += 1
            code_range_index += 1

    if branch_values:
        _add_named_list(
            workbook,
            name="CatalogoSucursales",
            sheet_name=catalogs.title,
            column="A",
            start_row=2,
            end_row=len(branch_values) + 1,
        )
        branch_validation = DataValidation(
            type="list",
            formula1="=CatalogoSucursales",
            allow_blank=False,
        )
        branch_validation.error = "Selecciona una sucursal incluida en el catálogo."
        branch_validation.errorTitle = "Sucursal inválida"
        branch_validation.showErrorMessage = True
        sheet.add_data_validation(branch_validation)
        branch_validation.add("A2:A1001")

    if family_map_row > 2:
        sheet_ref = validations.title.replace("'", "''")
        workbook.defined_names.add(
            DefinedName(
                "MapaFamilias",
                attr_text=(
                    "'" + sheet_ref
                    + "'!$A$2:$B$"
                    + str(family_map_row - 1)
                ),
            )
        )
        family_validation = DataValidation(
            type="list",
            formula1=(
                '=INDIRECT(IFERROR(VLOOKUP($A2,MapaFamilias,2,FALSE),"ListaVacia"))'
            ),
            allow_blank=True,
        )
        family_validation.error = "Selecciona primero una sucursal y luego una familia válida."
        family_validation.errorTitle = "Familia inválida"
        family_validation.showErrorMessage = True
        sheet.add_data_validation(family_validation)
        family_validation.add("B2:B1001")

    if code_map_row > 2:
        sheet_ref = validations.title.replace("'", "''")
        workbook.defined_names.add(
            DefinedName(
                "MapaCodigos",
                attr_text=("'" + sheet_ref + "'!$D$2:$E$" + str(code_map_row - 1)),
            )
        )
        code_validation = DataValidation(
            type="list",
            formula1=(
                '=INDIRECT(IFERROR(VLOOKUP($A2&"|"&$B2,MapaCodigos,2,FALSE),"ListaVacia"))'
            ),
            allow_blank=True,
        )
        code_validation.error = "Selecciona una familia para cargar los códigos disponibles."
        code_validation.errorTitle = "Código inválido"
        code_validation.showErrorMessage = True
        sheet.add_data_validation(code_validation)
        code_validation.add("C2:C1001")

    date_validation = DataValidation(
        type="date",
        operator="between",
        formula1="DATE(2020,1,1)",
        formula2="DATE(2100,12,31)",
        allow_blank=False,
    )
    date_validation.error = "Captura una fecha válida en formato dd/mm/aaaa."
    date_validation.errorTitle = "Fecha inválida"
    date_validation.prompt = "Formato esperado: dd/mm/aaaa"
    date_validation.promptTitle = "Fecha programada"
    date_validation.showErrorMessage = True
    date_validation.showInputMessage = True
    sheet.add_data_validation(date_validation)
    date_validation.add("D2:D1001")

    activity_validation = DataValidation(
        type="custom",
        formula1="=LEN(TRIM(E2))>0",
        allow_blank=False,
    )
    activity_validation.error = (
        "La actividad es obligatoria. Describe el trabajo preventivo."
    )
    activity_validation.errorTitle = "Actividad obligatoria"
    activity_validation.prompt = (
        "Obligatoria. Describe el trabajo a realizar. "
        "Ej.: Mantenimiento preventivo general."
    )
    activity_validation.promptTitle = "Actividad preventiva"
    activity_validation.showErrorMessage = True
    activity_validation.showInputMessage = True
    sheet.add_data_validation(activity_validation)
    activity_validation.add("E2:E1001")

    if responsible_values:
        _add_named_list(
            workbook,
            name="CatalogoResponsables",
            sheet_name=catalogs.title,
            column="B",
            start_row=2,
            end_row=len(responsible_values) + 1,
        )
        responsible_validation = DataValidation(
            type="list",
            formula1="=CatalogoResponsables",
            allow_blank=False,
        )
        responsible_validation.error = "Selecciona un responsable incluido en el catálogo."
        responsible_validation.errorTitle = "Responsable inválido"
        responsible_validation.showErrorMessage = True
        sheet.add_data_validation(responsible_validation)
        responsible_validation.add("F2:F1001")

    target_validation = DataValidation(
        type="list",
        formula1='"Equipo,Edificio"',
        allow_blank=True,
    )
    target_validation.error = "Selecciona Equipo o Edificio."
    target_validation.errorTitle = "Tipo inválido"
    target_validation.prompt = (
        "Vacío se interpreta como Equipo. Para Edificio deja Familia y "
        "Código equipo vacíos y selecciona una clasificación."
    )
    target_validation.promptTitle = "Tipo de preventivo"
    target_validation.showErrorMessage = True
    target_validation.showInputMessage = True
    sheet.add_data_validation(target_validation)
    target_validation.add("J2:J1001")

    if building_values:
        _add_named_list(
            workbook,
            name="CatalogoEdificio",
            sheet_name=catalogs.title,
            column="H",
            start_row=2,
            end_row=len(building_values) + 1,
        )
        building_validation = DataValidation(
            type="list",
            formula1="=CatalogoEdificio",
            allow_blank=True,
        )
        building_validation.error = (
            "Selecciona una clasificación de Edificio incluida en el catálogo."
        )
        building_validation.errorTitle = "Clasificación inválida"
        building_validation.showErrorMessage = True
        sheet.add_data_validation(building_validation)
        building_validation.add("K2:K1001")

    repeat_validation = DataValidation(
        type="list",
        formula1='"No,Sí"',
        allow_blank=True,
    )
    repeat_validation.error = "Selecciona Sí o No."
    repeat_validation.errorTitle = "Repetición inválida"
    repeat_validation.prompt = (
        "Sí = genera una programación recurrente. "
        "No o vacío = preventivo de una sola fecha."
    )
    repeat_validation.promptTitle = "Se repite"
    repeat_validation.showErrorMessage = True
    repeat_validation.showInputMessage = True
    sheet.add_data_validation(repeat_validation)
    repeat_validation.add("G2:G1001")

    interval_validation = DataValidation(
        type="whole",
        operator="greaterThan",
        formula1="0",
        allow_blank=True,
    )
    interval_validation.error = (
        "Captura un número entero mayor a cero."
    )
    interval_validation.errorTitle = "Intervalo inválido"
    interval_validation.prompt = (
        "Solo aplica cuando Se repite = Sí. "
        "Se cuentan únicamente lunes a viernes."
    )
    interval_validation.promptTitle = "Días hábiles"
    interval_validation.showErrorMessage = True
    interval_validation.showInputMessage = True
    sheet.add_data_validation(interval_validation)
    interval_validation.add("H2:H1001")

    duplicate_fill = PatternFill("solid", fgColor="FFC7CE")
    duplicate_font = Font(color="9C0006", bold=True)
    sheet.conditional_formatting.add(
        "A2:L1001",
        FormulaRule(
            formula=['$L2="⚠ DUPLICADO"'],
            fill=duplicate_fill,
            font=duplicate_font,
            stopIfTrue=True,
        ),
    )

    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    sheet.row_dimensions[1].height = 26

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
