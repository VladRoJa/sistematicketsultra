"""Frozen Excel/CSV audiences for Campaign V2; no provider calls.

Every request re-parses uploaded bytes, resolves authoritative branch/catalog state,
enforces permissions and template/channel compatibility, and fingerprints the
result. There is deliberately no arbitrary recipient-ID JSON or bypass of the
existing governed dispatch / blacklist pipeline.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
from collections import Counter
from datetime import date, datetime, timezone
from typing import Any, Iterable

from openpyxl import load_workbook
from sqlalchemy.exc import SQLAlchemyError

from app.extensions import db
from app.models.marketing import (
    MarketingCampaignV2ORM,
    MarketingCampaignV2RecipientEvidenceORM,
    MarketingCampaignV2RecipientORM,
)
from app.services.marketing_campaign_v2_blacklist_service import get_blacklisted_phones
from app.services.marketing_campaign_v2_dispatch_branch_service import resolve_dispatch_branch_key
from app.services.marketing_campaign_v2_dispatch_config_service import (
    resolve_dispatch_channel_binding,
    resolve_dispatch_template,
    template_allows_channel,
)
from app.services.marketing_phone import normalize_phone
from app.warehouse.services.socios_vencidos_current_status_resolver import (
    normalize_socios_vencidos_branch_key,
)

SOURCE = "EXCEL_IMPORT"
VERSION = "campaign-v2-excel-import-v1"
MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 100000
PURPOSES = frozenset({"REACTIVATION", "ACTIVE_MEMBERS", "NEW_SALE", "UNCLASSIFIED"})
FIELDS = (
    "sucursal", "telefono", "nombre", "pin", "id_socio",
    "tarifa_ultimo_plan", "tarifa_actual", "fecha_vencimiento",
    "fecha_vencimiento_2026",
)


class ExcelCampaignImportValidationError(ValueError):
    pass


class ExcelCampaignImportConflictError(ValueError):
    pass


def _text(value: Any, *, limit: int = 255) -> str:
    result = str(value if value is not None else "").strip()
    if len(result) > limit:
        raise ExcelCampaignImportValidationError("Texto del Excel excede la longitud permitida.")
    return result


def _purpose(value: Any) -> str:
    purpose = _text(value, limit=30).upper()
    if purpose not in PURPOSES:
        raise ExcelCampaignImportValidationError("Propósito de campaña inválido.")
    return purpose


def _parse_rows(file_bytes: bytes, filename: str) -> list[dict[str, Any]]:
    name = _text(filename, limit=255)
    if not file_bytes or len(file_bytes) > MAX_BYTES:
        raise ExcelCampaignImportValidationError("Archivo vacío o mayor a 5 MB.")
    try:
        if name.lower().endswith(".csv"):
            reader = csv.DictReader(io.StringIO(file_bytes.decode("utf-8-sig"), newline=""))
            headers = [str(h or "").strip().lower() for h in (reader.fieldnames or [])]
            if len(headers) != len(set(headers)):
                raise ExcelCampaignImportValidationError("Columnas repetidas en CSV.")
            rows = [dict(zip(headers, values.values())) for values in reader]
        elif name.lower().endswith(".xlsx"):
            workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
            try:
                rows = []
                matched = False
                for sheet in workbook.worksheets:
                    iterator = sheet.iter_rows(values_only=True)
                    first = next(iterator, None)
                    headers = [str(h or "").strip().lower() for h in (first or ())]
                    if "sucursal" not in headers or "telefono" not in headers:
                        continue
                    if len(headers) != len(set(headers)):
                        raise ExcelCampaignImportValidationError("Columnas repetidas en Excel.")
                    matched = True
                    rows = [dict(zip(headers, values)) for values in iterator]
                    break
                if not matched:
                    raise ExcelCampaignImportValidationError("No existe una hoja con sucursal y telefono.")
            finally:
                workbook.close()
        else:
            raise ExcelCampaignImportValidationError("Solo se permiten archivos .xlsx y .csv.")
    except ExcelCampaignImportValidationError:
        raise
    except (UnicodeError, csv.Error, ValueError, KeyError, OSError, TypeError, EOFError) as exc:
        raise ExcelCampaignImportValidationError("No fue posible leer el archivo enviado.") from exc
    except Exception as exc:
        raise ExcelCampaignImportValidationError("Archivo XLSX/CSV corrupto o inválido.") from exc

    if not rows or len(rows) > MAX_ROWS:
        raise ExcelCampaignImportValidationError("La audiencia debe tener entre 1 y 100000 filas.")
    if not {"telefono", "sucursal"}.issubset(rows[0]):
        raise ExcelCampaignImportValidationError("Columnas obligatorias: telefono y sucursal.")
    return rows


def _expiration(raw: Any) -> date | None:
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = _text(raw, limit=32)
    if not text:
        return None
    try:
        return date.fromisoformat(text[:10])
    except ValueError as exc:
        raise ExcelCampaignImportValidationError("La fecha de vencimiento debe ser AAAA-MM-DD.") from exc


def _compile(
    *, file_bytes: bytes, filename: str, name: str, purpose: str,
    template_id: int, allowed_sucursal_keys: Iterable[str] | None, session: Any,
) -> dict[str, Any]:
    campaign_name = _text(name)
    if not campaign_name:
        raise ExcelCampaignImportValidationError("Nombre de campaña obligatorio.")
    normalized_purpose = _purpose(purpose)
    template = resolve_dispatch_template(
        template_id=template_id, purpose=normalized_purpose, session=session,
    )
    rows = _parse_rows(file_bytes, filename)
    normalized_scope = (
        {normalize_socios_vencidos_branch_key(k) for k in allowed_sucursal_keys}
        if allowed_sucursal_keys is not None else None
    )
    if normalized_scope is not None:
        normalized_scope.discard(None)
    recipients: list[dict[str, Any]] = []
    seen: set[str] = set()
    branches: Counter[str] = Counter()
    branch_cache: dict[str, Any] = {}
    channel_cache: dict[int, Any] = {}
    for line, row in enumerate(rows, start=2):
        branch_raw = _text(row.get("sucursal"))
        if branch_raw not in branch_cache:
            branch_cache[branch_raw] = resolve_dispatch_branch_key(branch_raw, session=session)
        branch = branch_cache[branch_raw]
        if branch is None:
            raise ExcelCampaignImportValidationError(
                f"Fila {line}: sucursal sin identidad canónica única."
            )
        canon_key = normalize_socios_vencidos_branch_key(branch.sucursal_canon)
        if normalized_scope is not None and canon_key not in normalized_scope:
            raise ExcelCampaignImportValidationError(
                f"Fila {line}: sucursal fuera del alcance autorizado."
            )
        if branch.sucursal_id not in channel_cache:
            channel_cache[branch.sucursal_id] = resolve_dispatch_channel_binding(
                sucursal_id=branch.sucursal_id, session=session,
            )
        channel = channel_cache[branch.sucursal_id]
        if channel is None or not template_allows_channel(
            template, str(channel.provider_channel_id)
        ):
            raise ExcelCampaignImportValidationError(
                f"Fila {line}: sucursal no autorizada para la plantilla."
            )
        phone = normalize_phone(_text(row.get("telefono"), limit=64))
        if phone is None:
            raise ExcelCampaignImportValidationError(f"Fila {line}: teléfono inválido.")
        if phone in seen:
            raise ExcelCampaignImportValidationError(
                f"Fila {line}: teléfono duplicado dentro del archivo."
            )
        seen.add(phone)
        expiration = _expiration(
            row.get("fecha_vencimiento") or row.get("fecha_vencimiento_2026")
        )
        recipients.append({
            "phone": phone,
            "member_name": _text(row.get("nombre")),
            "member_pin": _text(row.get("pin"), limit=64),
            "member_id": _text(row.get("id_socio"), limit=64),
            "sucursal": branch.track_label,
            "sucursal_key": branch.sucursal_canon,
            "tarifa": _text(row.get("tarifa_ultimo_plan") or row.get("tarifa_actual")),
            "expiration": expiration.isoformat() if expiration else None,
        })
        branches[branch.sucursal_canon] += 1
    recipients.sort(key=lambda x: x["phone"])
    blocked = get_blacklisted_phones(
        phones=[r["phone"] for r in recipients], session=session,
    )
    sendable = [r for r in recipients if r["phone"] not in blocked]
    if not sendable:
        raise ExcelCampaignImportValidationError("No hay destinatarios fuera de lista negra.")
    # Fingerprint depends on the actual cohort and current blacklist state.
    material = {
        "version": VERSION, "name": campaign_name, "purpose": normalized_purpose,
        "template_id": int(template.id), "template_name": template.template_name,
        "file_sha256": hashlib.sha256(file_bytes).hexdigest(),
        "recipients": sendable,
        "blacklist_count": len(blocked),
    }
    fingerprint = hashlib.sha256(json.dumps(
        material, sort_keys=True, ensure_ascii=False, separators=(",", ":"),
    ).encode("utf-8")).hexdigest()
    return {
        "name": campaign_name, "purpose": normalized_purpose,
        "template_id": int(template.id), "template_name": template.template_name,
        "sha256": material["file_sha256"],
        "rows_read": len(rows), "unique_recipient_count": len(sendable),
        "blacklist_excluded_count": len(blocked),
        "branch_counts": dict(sorted(branches.items())),
        "preview_fingerprint_version": VERSION, "preview_fingerprint": fingerprint,
        "recipients": sendable,
    }


def preview_excel_campaign(
    *, file_bytes: bytes, filename: str, name: str, purpose: str,
    template_id: int, allowed_sucursal_keys: Iterable[str] | None,
    session: Any | None = None,
) -> dict[str, Any]:
    plan = _compile(
        file_bytes=file_bytes, filename=filename, name=name, purpose=purpose,
        template_id=template_id, allowed_sucursal_keys=allowed_sucursal_keys,
        session=session if session is not None else db.session,
    )
    return {key: value for key, value in plan.items() if key != "recipients"}


def freeze_excel_campaign(
    *, file_bytes: bytes, filename: str, name: str, purpose: str,
    template_id: int, allowed_sucursal_keys: Iterable[str] | None,
    expected_preview_fingerprint: str, created_by_user_id: int,
    session: Any | None = None,
) -> dict[str, Any]:
    active = session if session is not None else db.session
    plan = _compile(
        file_bytes=file_bytes, filename=filename, name=name, purpose=purpose,
        template_id=template_id, allowed_sucursal_keys=allowed_sucursal_keys,
        session=active,
    )
    if plan["preview_fingerprint"] != str(expected_preview_fingerprint or ""):
        raise ExcelCampaignImportConflictError(
            "Cambió el Excel, catálogo o lista negra; vuelve a generar Preview."
        )
    if active.query(MarketingCampaignV2ORM.id).filter(
        MarketingCampaignV2ORM.name == plan["name"],
        MarketingCampaignV2ORM.source == SOURCE,
    ).first():
        raise ExcelCampaignImportConflictError("Ya existe una campaña importada con ese nombre.")
    frozen_at = datetime.now(timezone.utc)
    branch_keys = sorted(plan["branch_counts"])
    campaign = MarketingCampaignV2ORM(
        name=plan["name"], purpose=plan["purpose"], source=SOURCE,
        audience_definition_json={
            "schema_version": 1,
            "filters": {
                "source": SOURCE,
                "allowed_sucursal_keys": branch_keys,
            },
            "source_metadata": {
                "upload_filename": _text(filename),
                "upload_sha256": plan["sha256"],
                "template_id_at_freeze": plan["template_id"],
                "template_name_at_freeze": plan["template_name"],
                "rows_read": plan["rows_read"],
                "blacklist_excluded_count": plan["blacklist_excluded_count"],
                "branch_counts": plan["branch_counts"],
            },
            "preview": {
                "fingerprint_version": VERSION,
                "fingerprint": plan["preview_fingerprint"],
                "summary": {
                    "unique_recipient_count": plan["unique_recipient_count"],
                    "blacklist_excluded_count": plan["blacklist_excluded_count"],
                },
            },
        },
        created_by_user_id=int(created_by_user_id),
        frozen_at=frozen_at, created_at=frozen_at, updated_at=frozen_at,
    )
    for row in plan["recipients"]:
        exp = date.fromisoformat(row["expiration"]) if row["expiration"] else None
        recipient = MarketingCampaignV2RecipientORM(
            phone_mx10=row["phone"], source=SOURCE,
            member_name=row["member_name"] or None,
            member_pin=row["member_pin"] or None,
            member_id=row["member_id"] or None,
            sucursal=row["sucursal"],
            tarifa_raw=row["tarifa"] or None,
            fecha_vencimiento_date=exp,
            inclusion_reason="VERIFIED_EXCEL_IMPORT",
            conflict_fields_json=[],
            created_at=frozen_at,
        )
        recipient.evidence_rows.append(MarketingCampaignV2RecipientEvidenceORM(
            evidence_order=0, source=SOURCE,
            phone_raw=row["phone"], phone_mx10=row["phone"],
            member_name=row["member_name"] or None,
            member_pin=row["member_pin"] or None,
            member_id=row["member_id"] or None,
            sucursal=row["sucursal"], sucursal_key=row["sucursal_key"],
            tarifa_raw=row["tarifa"] or None,
            fecha_vencimiento_date=exp,
            evidence_json=["EXCEL_IMPORT", f"FILE_SHA256:{plan['sha256']}"],
            created_at=frozen_at,
        ))
        campaign.recipients.append(recipient)
    try:
        active.add(campaign)
        active.flush()
        active.commit()
    except SQLAlchemyError:
        active.rollback()
        raise
    return {
        "campaign_id": int(campaign.id),
        "name": campaign.name, "purpose": campaign.purpose, "source": SOURCE,
        "recipient_count": plan["unique_recipient_count"],
        "blacklist_excluded_count": plan["blacklist_excluded_count"],
        "upload_sha256": plan["sha256"],
        "frozen_at": frozen_at.isoformat(),
    }
