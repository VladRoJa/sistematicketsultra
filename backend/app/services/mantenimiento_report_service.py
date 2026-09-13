from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime, timezone
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import joinedload

from app.models.ticket_model import Ticket
from app.services import mantenimiento_equipos_report_service as legacy
from app.utils.ticket_filters import filtrar_tickets_por_usuario


BUSINESS_TIMEZONE = legacy.BUSINESS_TIMEZONE
RegionReporteNoEncontradaError = legacy.RegionReporteNoEncontradaError
listar_regiones_reporte = legacy.listar_regiones_reporte
obtener_region_reporte = legacy.obtener_region_reporte

ACTIVE_TICKET_STATUSES = legacy.ACTIVE_TICKET_STATUSES
TYPE_EQUIPMENT = "Equipo"
TYPE_BUILDING = "Edificio"
TYPE_UNCLASSIFIED = "Sin clasificar"

TICKETS_HEADERS = (
    "Tipo mantenimiento",
    "Ticket ID",
    "Aparato / Elemento",
    "Código Interno",
    "Descripción",
    "Estado",
    "Criticidad",
    "Fecha Creación",
    "Sucursal",
    "Categoría",
    "Subcategoría",
    "Detalle",
    "Familia",
    "Falla detectada",
    "Condición",
    "Ubicación",
    "Reparación",
    "Refacción",
    "Observación / Plan de trabajo",
)

BUILDING_HEADERS = (
    "Sucursal",
    "Ticket ID",
    "Categoría",
    "Subcategoría",
    "Detalle",
    "Elemento / Equipo",
    "Ubicación",
    "Descripción",
    "Criticidad",
    "Estado",
    "Reparación",
    "Refacción",
    "Observación / Plan de trabajo",
)

VALIDATION_HEADERS = (
    "Tipo mantenimiento",
    "Sucursal",
    "Ticket ID",
    "Aparato / Elemento",
    "Descripción",
    "Categoría / Familia",
    "Desde por validar",
    "Días por validar",
    "Reparación",
    "Refacción",
    "Observación / Plan de trabajo",
)

HISTORY_HEADERS = (
    "Mes",
    "Tickets creados",
    "Equipo",
    "Edificio",
    "Sin clasificar",
    "Finalizados",
)

SPANISH_MONTH_ABBREVIATIONS = (
    "ene",
    "feb",
    "mar",
    "abr",
    "may",
    "jun",
    "jul",
    "ago",
    "sep",
    "oct",
    "nov",
    "dic",
)


# ---------------------------------------------------------------------------
# Query / clasificación
# ---------------------------------------------------------------------------


def _query_tickets_mantenimiento(user=None):
    query = (
        filtrar_tickets_por_usuario(user)
        if user is not None
        else Ticket.query
    )

    return (
        query
        .options(
            joinedload(Ticket.inventario),
            joinedload(Ticket.familia_equipo),
            joinedload(Ticket.falla_mantenimiento),
            joinedload(Ticket.sucursal),
            joinedload(Ticket.sucursal_destino),
            joinedload(Ticket.clasificacion),
        )
        .filter(Ticket.departamento_id == 1)
    )


def obtener_tickets_reporte(user=None, region_id=None):
    query = _query_tickets_mantenimiento(user).filter(
        Ticket.estado.in_(ACTIVE_TICKET_STATUSES)
    )
    query = legacy._aplicar_filtro_region(query, region_id)
    return (
        query
        .order_by(Ticket.sucursal_id_destino.asc(), Ticket.id.asc())
        .all()
    )


def obtener_tickets_historico_reporte(user=None, region_id=None):
    query = _query_tickets_mantenimiento(user)
    query = legacy._aplicar_filtro_region(query, region_id)
    return (
        query
        .order_by(Ticket.fecha_creacion.asc(), Ticket.id.asc())
        .all()
    )


def _classification_path(ticket):
    resolver = getattr(ticket, "_obtener_jerarquia_clasificacion", None)
    if callable(resolver):
        try:
            path = resolver() or []
            return [str(value).strip() for value in path if str(value).strip()]
        except Exception:
            pass

    path = []
    node = getattr(ticket, "clasificacion", None)
    visited = set()
    while node is not None:
        node_id = getattr(node, "id", None)
        marker = (node_id if node_id is not None else id(node))
        if marker in visited:
            break
        visited.add(marker)

        name = str(getattr(node, "nombre", "") or "").strip()
        if name:
            path.insert(0, name)
        node = getattr(node, "padre", None)
    return path


def _maintenance_type(ticket):
    path = _classification_path(ticket)
    normalized = [value.casefold() for value in path]

    if len(normalized) >= 2 and normalized[0] == "mantenimiento":
        if normalized[1] == "edificio":
            return TYPE_BUILDING
        if normalized[1] == "aparatos":
            return TYPE_EQUIPMENT

    if "edificio" in normalized[:3]:
        return TYPE_BUILDING

    if (
        getattr(ticket, "aparato_id", None) is not None
        or getattr(ticket, "familia_equipo_id", None) is not None
    ):
        return TYPE_EQUIPMENT

    return TYPE_UNCLASSIFIED


def _building_classification(ticket):
    path = _classification_path(ticket)
    normalized = [value.casefold() for value in path]

    if (
        len(path) >= 2
        and normalized[0] == "mantenimiento"
        and normalized[1] == "edificio"
    ):
        values = path[2:5]
    else:
        values = [
            getattr(ticket, "categoria", None),
            getattr(ticket, "subcategoria", None),
            getattr(ticket, "detalle", None),
        ]

    cleaned = [
        str(value).strip() if value not in (None, "") else None
        for value in values
    ]
    return tuple((cleaned + [None, None, None])[:3])


# ---------------------------------------------------------------------------
# Filas combinadas
# ---------------------------------------------------------------------------


def _subject_name(ticket):
    if _maintenance_type(ticket) == TYPE_EQUIPMENT:
        equipment_name = legacy._equipment_name(ticket)
        if equipment_name:
            return equipment_name

    free_name = str(getattr(ticket, "equipo", "") or "").strip()
    if free_name:
        return free_name

    category, subcategory, detail = _building_classification(ticket)
    return detail or subcategory or category or "Sin elemento"


def _equipment_code(ticket):
    if _maintenance_type(ticket) != TYPE_EQUIPMENT:
        return ""
    return legacy._equipment_code(ticket)


def _spare_part_text(ticket):
    if not bool(getattr(ticket, "necesita_refaccion", False)):
        return "No"
    description = str(
        getattr(ticket, "descripcion_refaccion", "") or ""
    ).strip()
    return description or "Sí"


def _ticket_row(ticket):
    maintenance_type = _maintenance_type(ticket)
    category, subcategory, detail = _building_classification(ticket)

    if maintenance_type == TYPE_EQUIPMENT:
        family = legacy._snapshot_family_name(ticket)
        failure = legacy._failure_name(ticket)
        condition = legacy._condition_name(ticket)
        category = subcategory = detail = ""
    else:
        family = ""
        failure = ""
        condition = str(getattr(ticket, "condicion_operativa", "") or "").strip()

    return (
        maintenance_type,
        getattr(ticket, "id", None),
        _subject_name(ticket),
        _equipment_code(ticket),
        str(getattr(ticket, "descripcion", "") or "").strip(),
        str(getattr(ticket, "estado", "") or "").strip(),
        int(getattr(ticket, "criticidad", 1) or 1),
        legacy._format_created_at(getattr(ticket, "fecha_creacion", None)),
        legacy._safe_branch_name(ticket),
        category or "",
        subcategory or "",
        detail or "",
        family,
        failure,
        condition,
        str(getattr(ticket, "ubicacion", "") or "").strip(),
        legacy._format_commitment(getattr(ticket, "fecha_solucion", None)),
        _spare_part_text(ticket),
        legacy.plan_trabajo_desde_historial(
            getattr(ticket, "historial_fechas", None)
        ),
    )


def _building_row(ticket):
    category, subcategory, detail = _building_classification(ticket)
    return (
        legacy._safe_branch_name(ticket),
        getattr(ticket, "id", None),
        category or "",
        subcategory or "",
        detail or "",
        _subject_name(ticket),
        str(getattr(ticket, "ubicacion", "") or "").strip(),
        str(getattr(ticket, "descripcion", "") or "").strip(),
        int(getattr(ticket, "criticidad", 1) or 1),
        str(getattr(ticket, "estado", "") or "").strip(),
        legacy._format_commitment(getattr(ticket, "fecha_solucion", None)),
        _spare_part_text(ticket),
        legacy.plan_trabajo_desde_historial(
            getattr(ticket, "historial_fechas", None)
        ),
    )


def _validation_row(ticket, now=None):
    maintenance_type = _maintenance_type(ticket)
    category, _, _ = _building_classification(ticket)
    classification = (
        legacy._snapshot_family_name(ticket)
        if maintenance_type == TYPE_EQUIPMENT
        else (category or TYPE_UNCLASSIFIED)
    )

    return (
        maintenance_type,
        legacy._safe_branch_name(ticket),
        getattr(ticket, "id", None),
        _subject_name(ticket),
        str(getattr(ticket, "descripcion", "") or "").strip(),
        classification,
        legacy._format_created_at(getattr(ticket, "fecha_finalizado", None)),
        legacy._validation_wait_days(ticket, now=now),
        legacy._format_commitment(getattr(ticket, "fecha_solucion", None)),
        _spare_part_text(ticket),
        legacy.plan_trabajo_desde_historial(
            getattr(ticket, "historial_fechas", None)
        ),
    )


def _monthly_history_rows(tickets, now=None):
    created_by_month = defaultdict(Counter)
    finalized_by_month = Counter()
    first_month = None

    for ticket in tickets:
        created_month = legacy._month_key(getattr(ticket, "fecha_creacion", None))
        if created_month is not None:
            if first_month is None or created_month < first_month:
                first_month = created_month
            created_by_month[created_month]["total"] += 1
            created_by_month[created_month][_maintenance_type(ticket)] += 1

        if str(getattr(ticket, "estado", "") or "").strip().lower() == "finalizado":
            finalized_month = legacy._month_key(
                getattr(ticket, "fecha_finalizado", None)
            )
            if finalized_month is not None:
                finalized_by_month[finalized_month] += 1

    if first_month is None:
        return []

    now_local = (
        legacy._to_business_datetime(now)
        if now is not None
        else datetime.now(timezone.utc).astimezone(BUSINESS_TIMEZONE)
    )
    last_month = (now_local.year, now_local.month)

    rows = []
    month_key = first_month
    while month_key <= last_month:
        counts = created_by_month.get(month_key, Counter())
        year, month = month_key
        rows.append(
            (
                datetime(year, month, 1),
                counts.get("total", 0),
                counts.get(TYPE_EQUIPMENT, 0),
                counts.get(TYPE_BUILDING, 0),
                counts.get(TYPE_UNCLASSIFIED, 0),
                finalized_by_month.get(month_key, 0),
            )
        )
        month_key = legacy._next_month_key(month_key)
    return rows


# ---------------------------------------------------------------------------
# Carátula ejecutiva
# ---------------------------------------------------------------------------


def _is_overdue(ticket, today):
    due = legacy._to_business_datetime(getattr(ticket, "fecha_solucion", None))
    return bool(due and due.date() < today)


def _cover_metrics(tickets, historical_tickets, now_local):
    today = now_local.date()
    overdue = sum(1 for ticket in tickets if _is_overdue(ticket, today))
    no_date = sum(
        1 for ticket in tickets
        if getattr(ticket, "fecha_solucion", None) is None
    )
    no_work = sum(
        1 for ticket in tickets
        if str(getattr(ticket, "condicion_operativa", "") or "").strip().upper()
        == "NO_TRABAJA"
    )
    validating = sum(
        1 for ticket in tickets
        if str(getattr(ticket, "estado", "") or "").strip().lower()
        == "por_validar"
    )
    finalized_month = sum(
        1
        for ticket in historical_tickets
        if (
            str(getattr(ticket, "estado", "") or "").strip().lower()
            == "finalizado"
            and legacy._month_key(getattr(ticket, "fecha_finalizado", None))
            == (now_local.year, now_local.month)
        )
    )
    type_counts = Counter(_maintenance_type(ticket) for ticket in tickets)

    branch_stats = defaultdict(
        lambda: {"active": 0, "overdue": 0, "no_work": 0}
    )
    for ticket in tickets:
        branch = legacy._safe_branch_name(ticket)
        branch_stats[branch]["active"] += 1
        if _is_overdue(ticket, today):
            branch_stats[branch]["overdue"] += 1
        if (
            str(getattr(ticket, "condicion_operativa", "") or "").strip().upper()
            == "NO_TRABAJA"
        ):
            branch_stats[branch]["no_work"] += 1

    top_branches = sorted(
        branch_stats.items(),
        key=lambda item: (
            -item[1]["active"],
            -item[1]["overdue"],
            item[0].casefold(),
        ),
    )[:6]

    return {
        "active": len(tickets),
        "overdue": overdue,
        "no_date": no_date,
        "no_work": no_work,
        "validating": validating,
        "finalized_month": finalized_month,
        "type_counts": type_counts,
        "top_branches": top_branches,
        "in_time": max(0, len(tickets) - overdue - no_date),
    }


def _scope_label(region_id):
    if region_id is None:
        return "Todas las sucursales"
    region = obtener_region_reporte(region_id)
    return f"Región: {region.region_label}"


def _apply_cover_style(sheet, metrics, now_local, region_id):
    black = "111827"
    charcoal = "1F2937"
    orange = "E54525"
    orange_soft = "FFF1EC"
    red = "C62828"
    red_soft = "FDECEC"
    amber = "A16207"
    amber_soft = "FFF7D6"
    blue = "1D4ED8"
    blue_soft = "EEF4FF"
    green = "15803D"
    green_soft = "ECFDF3"
    gray_50 = "F8FAFC"
    gray_100 = "F1F5F9"
    gray_200 = "E2E8F0"
    gray_500 = "64748B"
    white = "FFFFFF"
    thin = Side(style="thin", color=gray_200)

    sheet.sheet_view.showGridLines = False
    for column in range(1, 15):
        sheet.column_dimensions[get_column_letter(column)].width = 13

    sheet.merge_cells("A1:N2")
    sheet["A1"] = "REPORTE EJECUTIVO DE MANTENIMIENTO"
    sheet["A1"].font = Font(color=white, bold=True, size=20)
    sheet["A1"].fill = PatternFill("solid", fgColor=black)
    sheet["A1"].alignment = Alignment(vertical="center")
    for row in (1, 2):
        sheet.row_dimensions[row].height = 28
        for cell in sheet[row]:
            cell.fill = PatternFill("solid", fgColor=black)

    sheet.merge_cells("A3:H3")
    sheet["A3"] = (
        f"Corte: {now_local.day:02d} "
        f"{SPANISH_MONTH_ABBREVIATIONS[now_local.month - 1]} {now_local.year}"
        f"  ·  Alcance: {_scope_label(region_id)}"
    )
    sheet["A3"].font = Font(color=gray_500, bold=True, size=10)
    sheet["A3"].fill = PatternFill("solid", fgColor=gray_50)

    sheet.merge_cells("I3:N3")
    sheet["I3"] = "Cobertura: EQUIPO + EDIFICIO"
    sheet["I3"].font = Font(color=orange, bold=True, size=10)
    sheet["I3"].fill = PatternFill("solid", fgColor=orange_soft)
    sheet["I3"].alignment = Alignment(horizontal="right")

    cards = (
        ("A5:B8", "ACTIVOS", metrics["active"], black, gray_50, ""),
        (
            "C5:D8",
            "VENCIDOS",
            metrics["overdue"],
            red,
            red_soft,
            (
                f"{metrics['overdue'] / metrics['active']:.1%} del backlog"
                if metrics["active"]
                else "0.0% del backlog"
            ),
        ),
        (
            "E5:F8",
            "SIN FECHA",
            metrics["no_date"],
            amber,
            amber_soft,
            "sin compromiso",
        ),
        (
            "G5:H8",
            "NO TRABAJA",
            metrics["no_work"],
            red,
            red_soft,
            "impacto operativo",
        ),
        (
            "I5:J8",
            "POR VALIDAR",
            metrics["validating"],
            blue,
            blue_soft,
            "pendientes de cierre",
        ),
        (
            "K5:N8",
            f"FINALIZADOS {SPANISH_MONTH_ABBREVIATIONS[now_local.month - 1].upper()}",
            metrics["finalized_month"],
            green,
            green_soft,
            "avance del mes",
        ),
    )

    for cell_range, label, value, color, fill, note in cards:
        sheet.merge_cells(cell_range)
        cell = sheet[cell_range.split(":")[0]]
        cell.value = f"{label}\n{value}\n{note}"
        cell.font = Font(color=color, bold=True, size=12)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)

    sheet.merge_cells("A10:G10")
    sheet["A10"] = "DÓNDE ESTÁ LA PRESIÓN"
    sheet.merge_cells("H10:N10")
    sheet["H10"] = "QUÉ ESTÁ PASANDO"
    for address in ("A10", "H10"):
        sheet[address].font = Font(color=white, bold=True, size=11)
        sheet[address].fill = PatternFill("solid", fgColor=charcoal)

    branch_headers = ("Sucursal", "Activos", "Vencidos", "No trabaja", "% vencido")
    for column, value in enumerate(branch_headers, start=1):
        cell = sheet.cell(row=11, column=column, value=value)
        cell.font = Font(color=charcoal, bold=True, size=9)
        cell.fill = PatternFill("solid", fgColor=gray_100)
        cell.alignment = Alignment(horizontal="center")

    for row_index, (branch, values) in enumerate(metrics["top_branches"], start=12):
        percent = (
            values["overdue"] / values["active"]
            if values["active"]
            else 0
        )
        row = (
            branch,
            values["active"],
            values["overdue"],
            values["no_work"],
            percent,
        )
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=column, value=value)
            cell.font = Font(bold=(column == 1), size=9)
            cell.border = Border(bottom=thin)
        sheet.cell(row=row_index, column=5).number_format = "0.0%"

    sheet.merge_cells("H11:J11")
    sheet["H11"] = "BACKLOG"
    sheet["H11"].font = Font(color=charcoal, bold=True, size=9)
    sheet["H11"].fill = PatternFill("solid", fgColor=gray_100)
    backlog_rows = (
        ("Vencidos", metrics["overdue"]),
        ("Sin fecha", metrics["no_date"]),
        ("En tiempo / futuro", metrics["in_time"]),
    )
    for row_index, (label, value) in enumerate(backlog_rows, start=12):
        sheet.cell(row=row_index, column=8, value=label).font = Font(bold=True, size=9)
        sheet.cell(row=row_index, column=9, value=value).font = Font(bold=True, size=11)
        ratio = value / metrics["active"] if metrics["active"] else 0
        sheet.cell(row=row_index, column=10, value=ratio).number_format = "0.0%"

    sheet.merge_cells("K11:N11")
    sheet["K11"] = "TIPO DE MANTENIMIENTO"
    sheet["K11"].font = Font(color=charcoal, bold=True, size=9)
    sheet["K11"].fill = PatternFill("solid", fgColor=gray_100)
    type_rows = (
        (TYPE_EQUIPMENT, metrics["type_counts"].get(TYPE_EQUIPMENT, 0)),
        (TYPE_BUILDING, metrics["type_counts"].get(TYPE_BUILDING, 0)),
        (
            TYPE_UNCLASSIFIED,
            metrics["type_counts"].get(TYPE_UNCLASSIFIED, 0),
        ),
    )
    for row_index, (label, value) in enumerate(type_rows, start=12):
        sheet.cell(row=row_index, column=11, value=label).font = Font(bold=True, size=9)
        sheet.cell(row=row_index, column=12, value=value).font = Font(bold=True, size=11)

    sheet.merge_cells("A19:N19")
    sheet["A19"] = "LECTURA EJECUTIVA"
    sheet["A19"].font = Font(color=white, bold=True, size=11)
    sheet["A19"].fill = PatternFill("solid", fgColor=orange)

    active = metrics["active"]
    overdue_pct = metrics["overdue"] / active if active else 0
    top_branch = metrics["top_branches"][0] if metrics["top_branches"] else None
    insights = [
        (
            f"1. El principal frente es el atraso: {metrics['overdue']} de "
            f"{active} tickets activos están vencidos ({overdue_pct:.1%})."
        ),
        (
            f"2. {top_branch[0]} concentra {top_branch[1]['active']} activos y "
            f"{top_branch[1]['overdue']} vencidos."
            if top_branch
            else "2. No hay sucursales con tickets activos en el alcance actual."
        ),
        (
            f"3. El backlog se compone de "
            f"{metrics['type_counts'].get(TYPE_EQUIPMENT, 0)} de Equipo, "
            f"{metrics['type_counts'].get(TYPE_BUILDING, 0)} de Edificio y "
            f"{metrics['no_date']} sin fecha compromiso."
        ),
    ]
    for row_index, text in enumerate(insights, start=20):
        sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=14)
        cell = sheet.cell(row=row_index, column=1, value=text)
        cell.font = Font(bold=(row_index == 20), size=10)
        cell.fill = PatternFill("solid", fgColor=(white if row_index % 2 == 0 else gray_50))
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
        sheet.row_dimensions[row_index].height = 25

    sheet.merge_cells("A24:N24")
    sheet["A24"] = "FOCO DE LA SEMANA"
    sheet["A24"].font = Font(color=white, bold=True, size=11)
    sheet["A24"].fill = PatternFill("solid", fgColor=charcoal)

    priorities = sorted(
        metrics["top_branches"],
        key=lambda item: (
            -item[1]["overdue"],
            -item[1]["no_work"],
            -item[1]["active"],
        ),
    )[:4]
    priority_names = " → ".join(item[0] for item in priorities)
    focus = (
        f"Atacar primero el backlog vencido y los elementos fuera de operación. "
        f"Prioridad sugerida: {priority_names}. "
        f"En paralelo, programar los {metrics['no_date']} tickets sin fecha."
        if priorities
        else "No hay backlog activo en el alcance actual."
    )
    sheet.merge_cells("A25:N28")
    sheet["A25"] = focus
    sheet["A25"].font = Font(bold=True, size=11)
    sheet["A25"].fill = PatternFill("solid", fgColor=orange_soft)
    sheet["A25"].alignment = Alignment(vertical="center", wrap_text=True)
    sheet["A25"].border = Border(
        top=Side(style="medium", color=orange),
        bottom=Side(style="medium", color=orange),
        left=Side(style="medium", color=orange),
        right=Side(style="medium", color=orange),
    )

    sheet.merge_cells("A30:N30")
    sheet["A30"] = "Fuente: Suite Ultra · Maintenance Planner"
    sheet["A30"].font = Font(color=gray_500, size=8)
    sheet["A30"].alignment = Alignment(horizontal="right")

    sheet.freeze_panes = "A4"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_area = "A1:N30"
    sheet.sheet_view.zoomScale = 90


def _rebuild_generic_sheets(workbook, tickets, historical_tickets, now_local, region_id):
    for name in ("Tickets", "Por validar", "Histórico mensual"):
        if name in workbook.sheetnames:
            workbook.remove(workbook[name])

    if "Fallas por familia" in workbook.sheetnames:
        workbook["Fallas por familia"].title = "Equipos por familia"

    cover = workbook.create_sheet("Resumen ejecutivo", 0)
    _apply_cover_style(
        cover,
        _cover_metrics(tickets, historical_tickets, now_local),
        now_local,
        region_id,
    )

    tickets_sheet = workbook.create_sheet("Tickets", 1)
    tickets_sheet.append(TICKETS_HEADERS)
    for ticket in tickets:
        tickets_sheet.append(_ticket_row(ticket))
    legacy._style_worksheet(
        tickets_sheet,
        TICKETS_HEADERS,
        (18, 11, 28, 18, 42, 16, 11, 20, 22, 20, 20, 20, 24, 28, 18, 24, 18, 30, 52),
    )

    building_sheet = workbook.create_sheet("Edificio", 2)
    building_sheet.append(BUILDING_HEADERS)
    for ticket in tickets:
        if _maintenance_type(ticket) == TYPE_BUILDING:
            building_sheet.append(_building_row(ticket))
    legacy._style_worksheet(
        building_sheet,
        BUILDING_HEADERS,
        (24, 11, 22, 22, 22, 28, 24, 42, 11, 16, 18, 30, 52),
    )

    summary_index = workbook.sheetnames.index("Equipos por familia")
    validation_sheet = workbook.create_sheet("Por validar", summary_index + 1)
    validation_sheet.append(VALIDATION_HEADERS)
    validation_tickets = sorted(
        (
            ticket
            for ticket in tickets
            if str(getattr(ticket, "estado", "") or "").strip().lower()
            == "por_validar"
        ),
        key=legacy._por_validar_sort_key,
    )
    for ticket in validation_tickets:
        validation_sheet.append(_validation_row(ticket, now=now_local))
    legacy._style_worksheet(
        validation_sheet,
        VALIDATION_HEADERS,
        (18, 24, 11, 28, 42, 24, 20, 16, 18, 30, 52),
    )

    history_sheet = workbook.create_sheet(
        "Histórico mensual",
        workbook.sheetnames.index("Por validar") + 1,
    )
    history_sheet.append(HISTORY_HEADERS)
    for row in _monthly_history_rows(historical_tickets, now=now_local):
        history_sheet.append(row)
    legacy._style_worksheet(
        history_sheet,
        HISTORY_HEADERS,
        (18, 16, 14, 14, 16, 14),
    )
    for cell in history_sheet["A"][1:]:
        cell.number_format = '[$-es-MX]mmmm yyyy'


def construir_reporte_xlsx(
    tickets=None,
    *,
    historical_tickets=None,
    user=None,
    region_id=None,
    now=None,
):
    tickets_were_provided = tickets is not None

    tickets = list(
        tickets
        if tickets is not None
        else obtener_tickets_reporte(user=user, region_id=region_id)
    )
    historical_tickets = list(
        historical_tickets
        if historical_tickets is not None
        else (
            tickets
            if tickets_were_provided
            else obtener_tickets_historico_reporte(
                user=user,
                region_id=region_id,
            )
        )
    )

    equipment_tickets = [
        ticket
        for ticket in tickets
        if _maintenance_type(ticket) == TYPE_EQUIPMENT
    ]
    equipment_historical = [
        ticket
        for ticket in historical_tickets
        if _maintenance_type(ticket) == TYPE_EQUIPMENT
    ]

    legacy_output = legacy.construir_reporte_xlsx(
        equipment_tickets,
        historical_tickets=equipment_historical,
    )
    workbook = load_workbook(legacy_output)

    now_local = (
        legacy._to_business_datetime(now)
        if now is not None
        else datetime.now(timezone.utc).astimezone(BUSINESS_TIMEZONE)
    )
    _rebuild_generic_sheets(
        workbook,
        tickets,
        historical_tickets,
        now_local,
        region_id,
    )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
