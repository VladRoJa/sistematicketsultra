from __future__ import annotations

from collections import Counter, defaultdict
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from app.services import mantenimiento_equipos_report_service as legacy
from app.services import mantenimiento_report_service as report


BUILDING_SUMMARY_SHEET = "Edificio por categoría"
UNCLASSIFIED_CATEGORY = "Sin clasificar"


def _building_category(ticket):
    if report._maintenance_type(ticket) != report.TYPE_BUILDING:
        return None

    category, _, _ = report._building_classification(ticket)
    value = str(category or "").strip()
    return value or UNCLASSIFIED_CATEGORY


def _building_category_matrix(tickets):
    building_tickets = [
        ticket
        for ticket in tickets
        if report._maintenance_type(ticket) == report.TYPE_BUILDING
    ]

    categories = sorted(
        {
            category
            for ticket in building_tickets
            if (category := _building_category(ticket))
            != UNCLASSIFIED_CATEGORY
        },
        key=str.casefold,
    )

    headers = (
        "Sucursal",
        "Fallas",
        *categories,
        UNCLASSIFIED_CATEGORY,
    )

    counts_by_branch = defaultdict(Counter)
    for ticket in building_tickets:
        branch = legacy._safe_branch_name(ticket)
        category = _building_category(ticket) or UNCLASSIFIED_CATEGORY
        counts_by_branch[branch]["total"] += 1
        counts_by_branch[branch][category] += 1

    rows = []
    for branch in sorted(counts_by_branch, key=str.casefold):
        counts = counts_by_branch[branch]
        rows.append(
            (
                branch,
                counts["total"],
                *(counts[category] for category in categories),
                counts[UNCLASSIFIED_CATEGORY],
            )
        )

    total_row = (
        "TOTAL",
        sum(row[1] or 0 for row in rows),
        *(
            sum(row[column_index] or 0 for row in rows)
            for column_index in range(2, len(headers))
        ),
    )

    return headers, rows, total_row


def _summary_widths(headers):
    widths = [24, 12]
    for header in headers[2:]:
        widths.append(max(14, min(28, len(str(header)) + 4)))
    return tuple(widths)


def _add_building_category_summary(workbook, tickets):
    if BUILDING_SUMMARY_SHEET in workbook.sheetnames:
        workbook.remove(workbook[BUILDING_SUMMARY_SHEET])

    insert_index = (
        workbook.sheetnames.index("Edificio") + 1
        if "Edificio" in workbook.sheetnames
        else len(workbook.sheetnames)
    )
    sheet = workbook.create_sheet(BUILDING_SUMMARY_SHEET, insert_index)

    headers, rows, total_row = _building_category_matrix(tickets)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    sheet.append(total_row)

    legacy._style_worksheet(
        sheet,
        headers,
        _summary_widths(headers),
    )

    total_row_index = sheet.max_row
    for cell in sheet[total_row_index]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FFF1EC")

    return sheet


def construir_reporte_xlsx(
    tickets=None,
    *,
    historical_tickets=None,
    user=None,
    region_id=None,
    now=None,
):
    tickets_were_provided = tickets is not None

    tickets = list(
        tickets
        if tickets is not None
        else report.obtener_tickets_reporte(user=user, region_id=region_id)
    )
    historical_tickets = list(
        historical_tickets
        if historical_tickets is not None
        else (
            tickets
            if tickets_were_provided
            else report.obtener_tickets_historico_reporte(
                user=user,
                region_id=region_id,
            )
        )
    )

    output = report.construir_reporte_xlsx(
        tickets,
        historical_tickets=historical_tickets,
        region_id=region_id,
        now=now,
    )
    workbook = load_workbook(output)
    _add_building_category_summary(workbook, tickets)

    result = BytesIO()
    workbook.save(result)
    result.seek(0)
    return result
