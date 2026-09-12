from __future__ import annotations

import re
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.marketing_access import MarketingAccess
from app.services.marketing_dashboard_service import load_visible_marketing_branches
from app.services.marketing_inputs_service import parse_month
from app.services.marketing_sales_funnel_detail_service import (
    LEAD_METRICS,
    METRIC_TITLES,
    SALES_METRICS,
    VISIT_METRICS,
    MarketingSalesFunnelDetailValidationError,
    _branch_name_map,
    _leads_meta_detail,
    _normalize_metric,
    _normalize_optional_int,
    _normalize_pagination,
    _sales_detail,
    _visits_detail,
)
from app.services.marketing_sales_funnel_service import ORIGIN_LABELS


FULL_DETAIL_PAGE_SIZE = 1_000_000
DETAIL_EXPORT_MIMETYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

SORTABLE_FIELDS_BY_KIND: dict[str, frozenset[str]] = {
    "sales": frozenset(
        {
            "branch",
            "date",
            "name",
            "pin",
            "phone",
            "tariff",
            "revenue",
            "origin",
            "survey",
            "transaction_branch",
        }
    ),
    "visits": frozenset(
        {
            "branch",
            "date",
            "phone",
            "origin",
            "source",
        }
    ),
    "leads": frozenset(
        {
            "branch",
            "date",
            "name",
            "phone",
            "channel",
            "contact_id",
        }
    ),
}

EXPORT_COLUMNS_BY_KIND: dict[str, tuple[tuple[str, str], ...]] = {
    "sales": (
        ("branch", "Sucursal KPI"),
        ("date", "Fecha"),
        ("name", "Nombre"),
        ("member_id", "ID socio"),
        ("pin", "PIN"),
        ("phone", "Teléfono"),
        ("folio", "Folio"),
        ("membership_type", "Tipo membresía"),
        ("tariff", "Tarifa"),
        ("revenue", "Ingreso"),
        ("origin", "Origen"),
        ("survey", "Encuesta"),
        ("transaction_branch", "Sucursal cobro"),
        ("payment_method", "Forma de pago"),
        ("id_order", "ID orden"),
        ("payment_place", "Lugar de pago"),
    ),
    "visits": (
        ("branch", "Sucursal KPI"),
        ("date", "Fecha"),
        ("phone", "Teléfono"),
        ("origin", "Origen"),
        ("source", "Fuente"),
    ),
    "leads": (
        ("branch", "Sucursal KPI"),
        ("date", "Fecha primer mensaje"),
        ("name", "Nombre"),
        ("phone", "Teléfono"),
        ("channel", "Canal"),
        ("contact_id", "ID contacto"),
    ),
}


def _normalize_sort(
    sort_by: Any,
    sort_dir: Any,
    kind: str,
) -> tuple[str | None, str]:
    normalized_sort_by = str(sort_by or "").strip() or None
    normalized_sort_dir = str(sort_dir or "asc").strip().lower() or "asc"

    if normalized_sort_dir not in {"asc", "desc"}:
        raise MarketingSalesFunnelDetailValidationError(
            "sort_dir debe ser asc o desc."
        )

    if normalized_sort_by is not None:
        allowed = SORTABLE_FIELDS_BY_KIND.get(kind, frozenset())
        if normalized_sort_by not in allowed:
            raise MarketingSalesFunnelDetailValidationError(
                f"sort_by no soportado para detalle {kind}."
            )

    return normalized_sort_by, normalized_sort_dir


def _sortable_value(value: Any) -> tuple[int, Any]:
    if isinstance(value, (int, float, Decimal)):
        return 0, Decimal(str(value))

    text = str(value).strip()
    numeric_text = text.replace(",", "")
    try:
        return 0, Decimal(numeric_text)
    except (InvalidOperation, ValueError):
        return 1, text.casefold()


def _sort_rows(
    rows: list[dict[str, Any]],
    sort_by: str | None,
    sort_dir: str,
) -> list[dict[str, Any]]:
    if sort_by is None:
        return list(rows)

    populated: list[dict[str, Any]] = []
    empty: list[dict[str, Any]] = []

    for row in rows:
        value = row.get(sort_by)
        if value is None or str(value).strip() == "":
            empty.append(row)
        else:
            populated.append(row)

    populated.sort(
        key=lambda row: _sortable_value(row.get(sort_by)),
        reverse=sort_dir == "desc",
    )
    return populated + empty


def _detail_title(
    *,
    metric: str,
    origin: str | None,
    branch_id_filter: int | None,
    branch_names: dict[int, str],
) -> str:
    title = (
        ORIGIN_LABELS[origin]
        if metric == "origin" and origin is not None
        else METRIC_TITLES[metric]
    )
    if branch_id_filter is not None:
        title = f"{title} · {branch_names.get(branch_id_filter, branch_id_filter)}"
    return title


def _load_unpaged_detail(
    *,
    month: str,
    access: MarketingAccess,
    metric: str,
    branch_id: Any = None,
    origin: str | None = None,
) -> dict[str, Any]:
    month_start = parse_month(month)
    metric, origin = _normalize_metric(metric, origin)
    branch_id_filter = _normalize_optional_int(
        branch_id,
        field_name="branch_id",
    )

    branches, branch_ids, scope = load_visible_marketing_branches(access)
    allowed = set(branch_ids)
    if branch_id_filter is not None and branch_id_filter not in allowed:
        raise MarketingSalesFunnelDetailValidationError(
            "La sucursal solicitada no pertenece al alcance del usuario."
        )

    branch_names = _branch_name_map(branches)

    if metric in SALES_METRICS or metric == "origin":
        kind = "sales"
        rows, count, revenue_total = _sales_detail(
            month_start=month_start,
            metric=metric,
            origin_filter=origin,
            branch_ids=branch_ids,
            branch_names=branch_names,
            branch_id_filter=branch_id_filter,
            page=1,
            page_size=FULL_DETAIL_PAGE_SIZE,
        )
    elif metric in VISIT_METRICS:
        kind = "visits"
        rows, count = _visits_detail(
            month_start=month_start,
            metric=metric,
            branch_ids=branch_ids,
            branch_names=branch_names,
            branch_id_filter=branch_id_filter,
            page=1,
            page_size=FULL_DETAIL_PAGE_SIZE,
        )
        revenue_total = Decimal("0")
    elif metric in LEAD_METRICS:
        kind = "leads"
        rows, count = _leads_meta_detail(
            month_start=month_start,
            branch_ids=branch_ids,
            branch_names=branch_names,
            branch_id_filter=branch_id_filter,
            page=1,
            page_size=FULL_DETAIL_PAGE_SIZE,
        )
        revenue_total = Decimal("0")
    else:
        raise MarketingSalesFunnelDetailValidationError(
            "metric no soportado para detalle del funnel."
        )

    return {
        "month": month_start.strftime("%Y-%m"),
        "scope": scope,
        "metric": metric,
        "origin": origin,
        "kind": kind,
        "title": _detail_title(
            metric=metric,
            origin=origin,
            branch_id_filter=branch_id_filter,
            branch_names=branch_names,
        ),
        "branch_id": branch_id_filter,
        "count": count,
        "revenue_total": revenue_total,
        "rows": rows,
    }


def build_marketing_sales_funnel_drilldown(
    *,
    month: str,
    access: MarketingAccess,
    metric: str,
    branch_id: Any = None,
    origin: str | None = None,
    page: Any = None,
    page_size: Any = None,
    sort_by: Any = None,
    sort_dir: Any = None,
) -> dict[str, Any]:
    normalized_page, normalized_page_size = _normalize_pagination(
        page,
        page_size,
    )
    detail = _load_unpaged_detail(
        month=month,
        access=access,
        metric=metric,
        branch_id=branch_id,
        origin=origin,
    )
    normalized_sort_by, normalized_sort_dir = _normalize_sort(
        sort_by,
        sort_dir,
        detail["kind"],
    )

    rows = _sort_rows(
        detail["rows"],
        normalized_sort_by,
        normalized_sort_dir,
    )
    count = int(detail["count"])
    page_start = (normalized_page - 1) * normalized_page_size
    page_end = page_start + normalized_page_size
    total_pages = max(
        1,
        (count + normalized_page_size - 1) // normalized_page_size,
    )

    return {
        "month": detail["month"],
        "scope": detail["scope"],
        "metric": detail["metric"],
        "origin": detail["origin"],
        "kind": detail["kind"],
        "title": detail["title"],
        "branch_id": detail["branch_id"],
        "count": count,
        "revenue_total": float(detail["revenue_total"]),
        "page": normalized_page,
        "page_size": normalized_page_size,
        "total_pages": total_pages,
        "sort_by": normalized_sort_by,
        "sort_dir": normalized_sort_dir,
        "rows": rows[page_start:page_end],
    }


def _excel_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    return value


def _safe_filename_segment(value: Any) -> str:
    normalized = re.sub(r"[^A-Za-z0-9_-]+", "_", str(value or "").strip())
    return normalized.strip("_") or "detalle"


def _build_excel_workbook(
    *,
    title: str,
    kind: str,
    rows: list[dict[str, Any]],
) -> BytesIO:
    columns = EXPORT_COLUMNS_BY_KIND[kind]
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Detalle"

    header_fill = PatternFill("solid", fgColor="211F1E")
    header_font = Font(color="FFFFFF", bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center")

    worksheet.append([label for _, label in columns])
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    for row in rows:
        worksheet.append(
            [_excel_value(row.get(field)) for field, _ in columns]
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        f"A1:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
    )

    for index, (field, label) in enumerate(columns, start=1):
        max_length = len(label)
        if worksheet.max_row >= 2:
            for cell in worksheet.iter_cols(
                min_col=index,
                max_col=index,
                min_row=2,
                max_row=worksheet.max_row,
            ):
                for item in cell:
                    value = item.value
                    if value is not None:
                        max_length = max(max_length, len(str(value)))
        worksheet.column_dimensions[get_column_letter(index)].width = min(
            max(max_length + 2, 12),
            42,
        )

        if field == "revenue" and worksheet.max_row >= 2:
            for cell in worksheet[get_column_letter(index)][1:]:
                cell.number_format = '$' + '#,##0.00'

    worksheet.sheet_view.showGridLines = False
    workbook.properties.title = title

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_marketing_sales_funnel_drilldown_export(
    *,
    month: str,
    access: MarketingAccess,
    metric: str,
    branch_id: Any = None,
    origin: str | None = None,
    sort_by: Any = None,
    sort_dir: Any = None,
) -> tuple[BytesIO, str]:
    detail = _load_unpaged_detail(
        month=month,
        access=access,
        metric=metric,
        branch_id=branch_id,
        origin=origin,
    )
    normalized_sort_by, normalized_sort_dir = _normalize_sort(
        sort_by,
        sort_dir,
        detail["kind"],
    )
    rows = _sort_rows(
        detail["rows"],
        normalized_sort_by,
        normalized_sort_dir,
    )

    output = _build_excel_workbook(
        title=detail["title"],
        kind=detail["kind"],
        rows=rows,
    )

    filename_parts = [
        "funnel_venta_nueva",
        _safe_filename_segment(detail["month"]),
        _safe_filename_segment(detail["metric"]),
    ]
    if detail["origin"]:
        filename_parts.append(_safe_filename_segment(detail["origin"]))
    if detail["branch_id"] is not None:
        filename_parts.append(f"sucursal_{detail['branch_id']}")

    return output, "_".join(filename_parts) + ".xlsx"
