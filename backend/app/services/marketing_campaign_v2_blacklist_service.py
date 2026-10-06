"""Global incremental blacklist for Campaign V2 phone recipients."""

from __future__ import annotations

from io import BytesIO
import math
import unicodedata
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from sqlalchemy import func

from app.extensions import db
from app.models.marketing import MarketingCampaignV2BlacklistORM
from app.services.marketing_phone import normalize_phone


XLSX_MIMETYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
MAX_BLACKLIST_IMPORT_BYTES = 20 * 1024 * 1024
_HEADER_ALIASES = frozenset(
    {
        "TELEFONO",
        "PHONE",
        "CELULAR",
        "NUMERO",
    }
)


class MarketingCampaignV2BlacklistValidationError(ValueError):
    pass


class MarketingCampaignV2BlacklistPersistenceError(RuntimeError):
    pass


def get_campaign_v2_blacklist_summary(
    *,
    session: Any | None = None,
) -> dict[str, Any]:
    active_session = session if session is not None else db.session
    total, latest = (
        active_session.query(
            func.count(MarketingCampaignV2BlacklistORM.id),
            func.max(MarketingCampaignV2BlacklistORM.created_at),
        )
        .one()
    )
    return {
        "total": int(total or 0),
        "latest_created_at": _iso_datetime(latest),
    }


def get_blacklisted_phones(
    *,
    phones: Iterable[str],
    session: Any | None = None,
) -> set[str]:
    active_session = session if session is not None else db.session
    normalized = sorted(
        {
            phone
            for value in phones
            if (phone := normalize_phone(value)) is not None
        }
    )
    if not normalized:
        return set()

    found: set[str] = set()
    for chunk in _chunks(normalized, 5000):
        rows = (
            active_session.query(MarketingCampaignV2BlacklistORM.phone_mx10)
            .filter(MarketingCampaignV2BlacklistORM.phone_mx10.in_(tuple(chunk)))
            .all()
        )
        for row in rows:
            value = row[0] if isinstance(row, tuple) else getattr(row, "phone_mx10", row)
            phone = normalize_phone(value)
            if phone is not None:
                found.add(phone)
    return found


def import_campaign_v2_blacklist_xlsx(
    *,
    file_bytes: bytes,
    filename: str,
    created_by_user_id: int | None,
    session: Any | None = None,
) -> dict[str, Any]:
    active_session = session if session is not None else db.session
    normalized_filename = str(filename or "").strip()

    if not normalized_filename.lower().endswith(".xlsx"):
        raise MarketingCampaignV2BlacklistValidationError(
            "La lista negra debe importarse desde un archivo .xlsx."
        )
    if not file_bytes:
        raise MarketingCampaignV2BlacklistValidationError(
            "El archivo de lista negra está vacío."
        )
    if len(file_bytes) > MAX_BLACKLIST_IMPORT_BYTES:
        raise MarketingCampaignV2BlacklistValidationError(
            "El archivo de lista negra excede el límite de 20 MB."
        )

    parsed = _parse_blacklist_workbook(file_bytes)
    valid_phones = parsed["phones"]
    existing = get_blacklisted_phones(
        phones=valid_phones,
        session=active_session,
    )
    to_add = sorted(set(valid_phones) - existing)

    for phone in to_add:
        active_session.add(
            MarketingCampaignV2BlacklistORM(
                phone_mx10=phone,
                source_filename=normalized_filename[:255] or None,
                created_by_user_id=created_by_user_id,
            )
        )

    try:
        if to_add:
            active_session.commit()
    except Exception as exc:
        active_session.rollback()
        raise MarketingCampaignV2BlacklistPersistenceError(
            "No fue posible guardar la lista negra de Campaign V2."
        ) from exc

    summary = get_campaign_v2_blacklist_summary(session=active_session)
    return {
        "filename": normalized_filename,
        "rows_read": int(parsed["rows_read"]),
        "valid_unique": len(valid_phones),
        "added": len(to_add),
        "already_existing": len(existing),
        "duplicates_in_file": int(parsed["duplicates_in_file"]),
        "invalid": int(parsed["invalid"]),
        "blacklist_total": int(summary["total"]),
        "latest_created_at": summary["latest_created_at"],
    }


def export_campaign_v2_blacklist_xlsx(
    *,
    session: Any | None = None,
) -> tuple[bytes, str]:
    active_session = session if session is not None else db.session
    rows = (
        active_session.query(MarketingCampaignV2BlacklistORM)
        .order_by(
            MarketingCampaignV2BlacklistORM.created_at.asc(),
            MarketingCampaignV2BlacklistORM.id.asc(),
        )
        .all()
    )

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Lista negra")
    sheet.append(
        [
            "telefono",
            "fecha_alta",
            "archivo_origen",
            "usuario_id",
        ]
    )
    for row in rows:
        sheet.append(
            [
                str(row.phone_mx10 or ""),
                _iso_datetime(row.created_at) or "",
                _safe_excel_text(row.source_filename),
                row.created_by_user_id,
            ]
        )

    output = BytesIO()
    workbook.save(output)
    return output.getvalue(), "campaign_v2_lista_negra.xlsx"


def _parse_blacklist_workbook(file_bytes: bytes) -> dict[str, Any]:
    try:
        workbook = load_workbook(
            filename=BytesIO(file_bytes),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise MarketingCampaignV2BlacklistValidationError(
            "No fue posible leer el archivo Excel de lista negra."
        ) from exc

    try:
        if not workbook.worksheets:
            raise MarketingCampaignV2BlacklistValidationError(
                "El Excel de lista negra no contiene hojas."
            )
        sheet = workbook.worksheets[0]
        header_row, phone_column = _find_phone_header(sheet)
        if header_row is None or phone_column is None:
            raise MarketingCampaignV2BlacklistValidationError(
                "No se encontró una columna telefono en las primeras 10 filas."
            )

        phones: set[str] = set()
        rows_read = 0
        invalid = 0
        duplicates_in_file = 0

        for row in sheet.iter_rows(
            min_row=header_row + 1,
            min_col=phone_column,
            max_col=phone_column,
            values_only=True,
        ):
            raw_value = row[0] if row else None
            if _is_blank(raw_value):
                continue
            rows_read += 1
            phone = _normalize_excel_phone(raw_value)
            if phone is None:
                invalid += 1
                continue
            if phone in phones:
                duplicates_in_file += 1
                continue
            phones.add(phone)

        return {
            "phones": phones,
            "rows_read": rows_read,
            "invalid": invalid,
            "duplicates_in_file": duplicates_in_file,
        }
    finally:
        workbook.close()


def _find_phone_header(sheet) -> tuple[int | None, int | None]:
    for row_index, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=10, values_only=True),
        start=1,
    ):
        for column_index, value in enumerate(row, start=1):
            if _header_key(value) in _HEADER_ALIASES:
                return row_index, column_index
    return None, None


def _normalize_excel_phone(value: Any) -> str | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return normalize_phone(str(value))
    if isinstance(value, float):
        if not math.isfinite(value) or not value.is_integer():
            return None
        return normalize_phone(str(int(value)))
    return normalize_phone(value)


def _header_key(value: Any) -> str:
    normalized = unicodedata.normalize("NFKD", str(value or ""))
    ascii_value = normalized.encode("ascii", "ignore").decode("ascii")
    return "".join(character for character in ascii_value.upper() if character.isalnum())


def _is_blank(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def _chunks(values: list[str], size: int):
    for index in range(0, len(values), size):
        yield values[index : index + size]


def _safe_excel_text(value: Any) -> str:
    text = " ".join(str(value or "").split())
    if text.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _iso_datetime(value: Any) -> str | None:
    if value is None:
        return None
    isoformat = getattr(value, "isoformat", None)
    return isoformat() if callable(isoformat) else str(value)
