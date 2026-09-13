from __future__ import annotations

from datetime import date
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.extensions import db
from app.models import MarketingIventasContactORM
from app.services.marketing_access import MarketingAccess
from app.services.marketing_dashboard_service import load_visible_marketing_branches
from app.services.marketing_inputs_service import parse_month
from app.services.marketing_sales_funnel_detail_service import (
    MarketingSalesFunnelDetailValidationError,
    _normalize_optional_int,
    _normalize_pagination,
)
from app.services.marketing_sales_funnel_service import (
    _canonical_runs_for_window,
    _month_end,
)


LEAD_SORT_FIELDS = frozenset(
    {"branch", "date", "name", "phone", "channel", "contact_id"}
)
LEAD_EXPORT_COLUMNS = (
    ("branch", "Sucursal KPI"),
    ("date", "Fecha primer mensaje"),
    ("name", "Nombre"),
    ("phone", "Teléfono"),
    ("channel", "Canal"),
    ("contact_id", "ID contacto"),
)


def _current_month_run_ids(month_start: date) -> tuple[int, ...]:
    month_end = _month_end(month_start)
    period_key = f"IVENTAS-{month_start.strftime('%Y-%m')}"
    runs = _canonical_runs_for_window(month_start, month_end)
    return tuple(
        int(run.id)
        for run in runs
        if str(run.period_key) == period_key
    )


def count_monthly_iventas_leads(
    *,
    month_start: date,
    branch_ids: tuple[int, ...],
) -> int:
    """Count unique iVentas leads whose first message happened in the month."""
    if not branch_ids:
        return 0

    month_end = _month_end(month_start)
    run_ids = _current_month_run_ids(month_start)
    if not run_ids:
        return 0

    return int(
        db.session.query(
            MarketingIventasContactORM.sucursal_id,
            MarketingIventasContactORM.phone_mx10,
        )
        .filter(
            MarketingIventasContactORM.sync_run_id.in_(run_ids),
            MarketingIventasContactORM.sucursal_id.in_(branch_ids),
            MarketingIventasContactORM.first_message_at_utc.isnot(None),
            MarketingIventasContactORM.first_message_date_local.isnot(None),
            MarketingIventasContactORM.first_message_date_local >= month_start,
            MarketingIventasContactORM.first_message_date_local <= month_end,
            MarketingIventasContactORM.phone_mx10.isnot(None),
            MarketingIventasContactORM.phone_mx10 != "",
        )
        .distinct()
        .count()
    )


def _normalize_sort(sort_by: Any, sort_dir: Any) -> tuple[str | None, str]:
    normalized_sort_by = str(sort_by or "").strip() or None
    normalized_sort_dir = str(sort_dir or "asc").strip().lower() or "asc"

    if normalized_sort_dir not in {"asc", "desc"}:
        raise MarketingSalesFunnelDetailValidationError(
            "sort_dir debe ser asc o desc."
        )
    if (
        normalized_sort_by is not None
        and normalized_sort_by not in LEAD_SORT_FIELDS
    ):
        raise MarketingSalesFunnelDetailValidationError(
            "sort_by no soportado para Leads iVentas."
        )

    return normalized_sort_by, normalized_sort_dir


def _monthly_iventas_lead_rows(
    *,
    month_start: date,
    branch_ids: tuple[int, ...],
    branch_names: dict[int, str],
    branch_id_filter: int | None,
) -> list[dict[str, Any]]:
    if not branch_ids:
        return []

    month_end = _month_end(month_start)
    run_ids = _current_month_run_ids(month_start)
    if not run_ids:
        return []

    contacts = (
        MarketingIventasContactORM.query.filter(
            MarketingIventasContactORM.sync_run_id.in_(run_ids),
            MarketingIventasContactORM.sucursal_id.in_(branch_ids),
            MarketingIventasContactORM.first_message_at_utc.isnot(None),
            MarketingIventasContactORM.first_message_date_local.isnot(None),
            MarketingIventasContactORM.first_message_date_local >= month_start,
            MarketingIventasContactORM.first_message_date_local <= month_end,
            MarketingIventasContactORM.phone_mx10.isnot(None),
            MarketingIventasContactORM.phone_mx10 != "",
        )
        .order_by(
            MarketingIventasContactORM.first_message_date_local.asc(),
            MarketingIventasContactORM.id.asc(),
        )
        .all()
    )

    seen: set[tuple[int, str]] = set()
    rows: list[dict[str, Any]] = []
    for contact in contacts:
        branch_id = int(contact.sucursal_id)
        if branch_id_filter is not None and branch_id != branch_id_filter:
            continue

        phone = str(contact.phone_mx10 or "").strip()
        key = (branch_id, phone)
        if not phone or key in seen:
            continue
        seen.add(key)

        rows.append(
            {
                "branch_id": branch_id,
                "branch": branch_names.get(branch_id, ""),
                "date": (
                    contact.first_message_date_local.isoformat()
                    if contact.first_message_date_local is not None
                    else None
                ),
                "name": str(contact.name or "").strip() or None,
                "phone": phone,
                "contact_id": str(contact.contact_id or "").strip() or None,
                "channel": str(contact.channel_name or "").strip() or None,
                "origin_key": "IVENTAS",
                "origin": "iVentas",
            }
        )

    return rows


def _sort_rows(
    rows: list[dict[str, Any]],
    sort_by: str | None,
    sort_dir: str,
) -> list[dict[str, Any]]:
    if sort_by is None:
        return list(rows)

    populated: list[dict[str, Any]] = []
    empty: list[dict[str, Any]] = []
    for row in rows:
        value = row.get(sort_by)
        if value is None or str(value).strip() == "":
            empty.append(row)
        else:
            populated.append(row)

    populated.sort(
        key=lambda row: str(row.get(sort_by) or "").casefold(),
        reverse=sort_dir == "desc",
    )
    return populated + empty


def build_monthly_iventas_leads_detail(
    *,
    month: str,
    access: MarketingAccess,
    branch_id: Any = None,
    page: Any = None,
    page_size: Any = None,
    sort_by: Any = None,
    sort_dir: Any = None,
) -> dict[str, Any]:
    month_start = parse_month(month)
    normalized_page, normalized_page_size = _normalize_pagination(
        page,
        page_size,
    )
    branch_id_filter = _normalize_optional_int(
        branch_id,
        field_name="branch_id",
    )
    normalized_sort_by, normalized_sort_dir = _normalize_sort(
        sort_by,
        sort_dir,
    )

    branches, branch_ids, scope = load_visible_marketing_branches(access)
    allowed = set(branch_ids)
    if branch_id_filter is not None and branch_id_filter not in allowed:
        raise MarketingSalesFunnelDetailValidationError(
            "La sucursal solicitada no pertenece al alcance del usuario."
        )

    branch_names = {
        int(branch.sucursal_id): str(branch.name)
        for branch in branches
    }
    rows = _monthly_iventas_lead_rows(
        month_start=month_start,
        branch_ids=branch_ids,
        branch_names=branch_names,
        branch_id_filter=branch_id_filter,
    )
    rows = _sort_rows(rows, normalized_sort_by, normalized_sort_dir)

    count = len(rows)
    start = (normalized_page - 1) * normalized_page_size
    end = start + normalized_page_size
    total_pages = max(
        1,
        (count + normalized_page_size - 1) // normalized_page_size,
    )

    title = "Leads iVentas"
    if branch_id_filter is not None:
        title = f"{title} · {branch_names.get(branch_id_filter, branch_id_filter)}"

    return {
        "month": month_start.strftime("%Y-%m"),
        "scope": scope,
        "metric": "leads_iventas",
        "origin": None,
        "kind": "leads",
        "title": title,
        "branch_id": branch_id_filter,
        "count": count,
        "revenue_total": 0.0,
        "page": normalized_page,
        "page_size": normalized_page_size,
        "total_pages": total_pages,
        "sort_by": normalized_sort_by,
        "sort_dir": normalized_sort_dir,
        "rows": rows[start:end],
    }


def build_monthly_iventas_leads_export(
    *,
    month: str,
    access: MarketingAccess,
    branch_id: Any = None,
    sort_by: Any = None,
    sort_dir: Any = None,
) -> tuple[BytesIO, str]:
    month_start = parse_month(month)
    branch_id_filter = _normalize_optional_int(
        branch_id,
        field_name="branch_id",
    )
    normalized_sort_by, normalized_sort_dir = _normalize_sort(
        sort_by,
        sort_dir,
    )

    branches, branch_ids, _ = load_visible_marketing_branches(access)
    allowed = set(branch_ids)
    if branch_id_filter is not None and branch_id_filter not in allowed:
        raise MarketingSalesFunnelDetailValidationError(
            "La sucursal solicitada no pertenece al alcance del usuario."
        )

    branch_names = {
        int(branch.sucursal_id): str(branch.name)
        for branch in branches
    }
    rows = _monthly_iventas_lead_rows(
        month_start=month_start,
        branch_ids=branch_ids,
        branch_names=branch_names,
        branch_id_filter=branch_id_filter,
    )
    rows = _sort_rows(rows, normalized_sort_by, normalized_sort_dir)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Leads iVentas"

    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(color="FFFFFF", bold=True)
    for column_index, (_, label) in enumerate(LEAD_EXPORT_COLUMNS, start=1):
        cell = worksheet.cell(row=1, column=column_index, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_index, row in enumerate(rows, start=2):
        for column_index, (key, _) in enumerate(
            LEAD_EXPORT_COLUMNS,
            start=1,
        ):
            worksheet.cell(
                row=row_index,
                column=column_index,
                value=row.get(key),
            )

    for column_index, (key, label) in enumerate(
        LEAD_EXPORT_COLUMNS,
        start=1,
    ):
        max_length = len(label)
        for row in rows:
            max_length = max(max_length, len(str(row.get(key) or "")))
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(max_length + 2, 12),
            36,
        )

    worksheet.freeze_panes = "A2"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    filename = f"funnel_leads_iventas_{month_start.strftime('%Y-%m')}.xlsx"
    return output, filename
