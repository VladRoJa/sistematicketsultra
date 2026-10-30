"""E4 v0.1: lector de evidencia; no usa fórmulas ni porcentajes del Excel."""
from collections import Counter, defaultdict
from datetime import date, datetime
from io import BytesIO
import csv
import re
import unicodedata
from zipfile import ZipFile, BadZipFile
from openpyxl import load_workbook

BASELINE_COLUMNS = (
    "sucursal_id", "sucursal", "inventario_id", "codigo_interno",
    "categoria", "aparato", "marca", "stock_sistema_qa",
)
MAX_XLSX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 12000


class ArchivoE4Error(ValueError):
    pass


def norm(value):
    return "".join(x for x in unicodedata.normalize("NFKD", str(value or "").strip().casefold())
                   if not unicodedata.combining(x))


def text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def positive_int(value):
    if isinstance(value, bool):
        return None
    try:
        if isinstance(value, float) and not value.is_integer():
            return None
        number = int(str(value).strip()) if not isinstance(value, float) else int(value)
        return number if number > 0 else None
    except (ValueError, TypeError, OverflowError):
        return None


def parse_historical_csv(raw, sucursal_id):
    """Conservar la base exportada al distribuir; JAMÁS reconstruirla desde PostgreSQL actual."""
    if not raw or len(raw) > 5 * 1024 * 1024:
        raise ArchivoE4Error("CSV histórico vacío o demasiado grande")
    try:
        source = raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ArchivoE4Error("CSV histórico requiere UTF-8") from exc
    reader = csv.DictReader(source.splitlines())
    if reader.fieldnames is None or set(BASELINE_COLUMNS) - set(reader.fieldnames):
        raise ArchivoE4Error("CSV histórico carece de columnas obligatorias")
    matched = []
    excluded = []
    seen = set()
    for n, row in enumerate(reader, start=2):
        if n > MAX_ROWS:
            raise ArchivoE4Error("CSV histórico excede el máximo de filas")
        branch = positive_int(row.get("sucursal_id"))
        inv_id = positive_int(row.get("inventario_id"))
        stock_raw = text(row.get("stock_sistema_qa"))
        if branch is None or inv_id is None or not re.fullmatch(r"\d+", stock_raw):
            raise ArchivoE4Error(f"CSV histórico fila {n}: identidad o stock inválidos")
        if branch != sucursal_id:
            continue
        if inv_id in seen:
            raise ArchivoE4Error(f"CSV histórico: ID oficial repetido {inv_id}")
        seen.add(inv_id)
        snapshot = {key: text(row.get(key)) for key in BASELINE_COLUMNS}
        if int(stock_raw) == 0:
            excluded.append(snapshot)
            continue
        matched.append({
            "inventario_id_original": inv_id,
            "codigo_interno_original": text(row["codigo_interno"]),
            "stock_original": int(stock_raw),
            "datos_originales": snapshot,
        })
    if not matched:
        raise ArchivoE4Error("CSV histórico no contiene aparatos esperados para esta sucursal")
    return matched, excluded


def _safe_book(raw):
    if not raw or len(raw) > MAX_XLSX_BYTES:
        raise ArchivoE4Error("XLSX vacío o superior a 10 MB")
    try:
        with ZipFile(BytesIO(raw)) as z:
            entries = z.infolist()
            if len(entries) > 1200 or sum(e.file_size for e in entries) > 40 * 1024 * 1024:
                raise ArchivoE4Error("XLSX excede los límites de contenido")
            if any(e.file_size > 0 and e.compress_size > 0 and e.file_size / e.compress_size > 500 for e in entries):
                raise ArchivoE4Error("XLSX con compresión anómala")
            if any("vbaproject" in e.filename.casefold() or e.filename.startswith("xl/externalLinks/") for e in entries):
                raise ArchivoE4Error("XLSX con macros o enlaces externos no admitidos")
        return load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, OSError, KeyError, ValueError) as exc:
        if isinstance(exc, ArchivoE4Error):
            raise
        raise ArchivoE4Error("XLSX corrupto o no reconocible") from exc


def _cell(c, diag, hoja, fila, campo):
    if c.data_type == "f":
        diag.append({"severidad": "ERROR", "codigo": "FORMULA_CRITICA",
                     "hoja": hoja, "fila": fila, "campo": campo})
        return ""
    return text(c.value)


def _control(sheet, diag):
    fields = {
        "sucursal": "sucursal", "id sucursal": "sucursal_id",
        "gerente responsable": "responsable", "fecha del recorrido": "fecha_recorrido",
        "folio / corte e4": "folio_declarado", "version de entrega": "version_formato",
    }
    result = {}
    for cells in sheet.iter_rows(min_row=1, max_row=50, max_col=6):
        for i, c in enumerate(cells[:3]):
            key = fields.get(norm(c.value))
            if key:
                values = [x for x in cells[i + 1:]]
                nonempty = [x for x in values if x.value is not None and text(x.value)]
                result[key] = _cell(nonempty[-1], diag, "CONTROL", c.row, key) if nonempty else ""
                break
    return result


def _headers(sheet, definitions):
    for cells in sheet.iter_rows(min_row=1, max_row=30, max_col=10):
        values = [norm(c.value) for c in cells]
        # La plantilla v0.1 usa títulos sobre la fila de encabezados.
        if all(any(marker in v for v in values) for marker in definitions.values()):
            found = {}
            for key, marker in definitions.items():
                found[key] = next(i for i, v in enumerate(values) if marker in v)
            if len(set(found.values())) == len(definitions):
                return cells[0].row, found
    return None, None


def leer_identidad_entrega(raw):
    """CONTROL de un XLSX sin asumir que el contenido concede identidad oficial.

    Solo devuelve datos declarados. La identidad efectiva exige contrastar
    sucursal/ciclo contra E4 persistido; la BD actual nunca sustituye el corte.
    """
    book = _safe_book(raw)
    try:
        if "CONTROL" not in book.sheetnames:
            raise ArchivoE4Error("Falta hoja CONTROL")
        diagnostics = []
        control = _control(book["CONTROL"], diagnostics)
        if diagnostics:
            raise ArchivoE4Error("CONTROL contiene fórmulas en campos críticos")
        return control
    finally:
        book.close()


def parse_xlsx_v01(raw, sucursal_id, expected, modalidad):
    """Devuelve evidencia de la entrega y pendientes por ID original, sin escrituras."""
    book = _safe_book(raw)
    diagnostics = []
    rows = []
    expected_by_id = {int(row["inventario_id_original"]): row for row in expected}
    try:
        if not {"CONTROL", "APARATOS", "HALLAZGOS"}.issubset(book.sheetnames):
            raise ArchivoE4Error("Faltan hojas CONTROL, APARATOS o HALLAZGOS")
        ctl = _control(book["CONTROL"], diagnostics)
        # v0.1: plantilla original; v0.2: archivo efectivamente distribuido.
        # No aceptar por accidente versiones futuras con semánticas desconocidas.
        declared_version = norm(ctl.get("version_formato"))
        if not re.fullmatch(r"v0\.(?:1|2)(?:\s+.*)?", declared_version):
            diagnostics.append({
                "severidad": "ERROR", "codigo": "VERSION_FORMATO_NO_RECONOCIDA",
                "hoja": "CONTROL", "declarado": ctl.get("version_formato"),
                "versiones_admitidas": ["v0.1", "v0.2"],
            })
        declared = positive_int(ctl.get("sucursal_id"))
        if declared is None or declared != sucursal_id:
            diagnostics.append({"severidad": "ERROR", "codigo": "SUCURSAL_NO_COINCIDE",
                                "hoja": "CONTROL", "declarado": ctl.get("sucursal_id"), "esperado": sucursal_id})
        aheader, amap = _headers(book["APARATOS"], {
            "inventario_id": "id oficial", "codigo": "codigo interno",
            "categoria": "categoria", "aparato": "aparato", "marca": "marca",
            "respuesta": "esta? si / no", "observaciones": "observaciones",
        })
        hheader, hmap = _headers(book["HALLAZGOS"], {
            "codigo": "codigo en etiqueta", "aparato": "nombre / aparato",
            "categoria": "categoria", "marca": "marca",
            "ubicacion": "ubicacion / zona", "descripcion": "descripcion del hallazgo",
        })
        if aheader is None or hheader is None:
            raise ArchivoE4Error("Encabezados de APARATOS/HALLAZGOS incompatibles con v0.1")
        apparatus = defaultdict(list)
        for sheet_name, header, mapping in (("APARATOS", aheader, amap), ("HALLAZGOS", hheader, hmap)):
            sheet = book[sheet_name]
            # Los XLSX distribuidos pueden omitir <dimension>; en lectura
            # read_only openpyxl reporta max_row=None en ese caso.
            if sheet.max_row is not None and sheet.max_row > MAX_ROWS + header:
                raise ArchivoE4Error(f"{sheet_name}: demasiadas filas")
            for fila_excel, cells in enumerate(
                sheet.iter_rows(min_row=header + 1, max_col=max(mapping.values()) + 1),
                start=header + 1,
            ):
                if fila_excel > MAX_ROWS + header:
                    raise ArchivoE4Error(f"{sheet_name}: demasiadas filas")
                item = {key: _cell(cells[index], diagnostics, sheet_name, fila_excel, key)
                        for key, index in mapping.items()}
                if not any(item.values()):
                    continue
                issues = []
                inv = positive_int(item.get("inventario_id")) if sheet_name == "APARATOS" else None
                response = norm(item.get("respuesta")) if sheet_name == "APARATOS" else None
                answer = "SI" if response == "si" else "NO" if response == "no" else "SIN_REVISAR" if sheet_name == "APARATOS" else None
                if sheet_name == "APARATOS":
                    if modalidad == "INICIAL" and inv is None and item.get("aparato"):
                        # Equipo sin ID oficial: evidencia provisional, nunca alta canónica.
                        pass
                    elif inv is None or inv not in expected_by_id:
                        issues.append("REFERENCIA_NO_ESPERADA")
                    else:
                        if norm(item["codigo"]) != norm(expected_by_id[inv]["codigo_interno_original"]):
                            issues.append("CODIGO_DIFIERE_DEL_CORTE")
                        apparatus[inv].append((fila_excel, answer, issues))
                    if response not in ("si", "no", ""):
                        issues.append("RESPUESTA_INVALIDA")
                elif not item.get("aparato"):
                    issues.append("HALLAZGO_SIN_NOMBRE")
                for code in issues:
                    diagnostics.append({"severidad": "ERROR", "codigo": code,
                                        "hoja": sheet_name, "fila": fila_excel, "inventario_id": inv})
                rows.append({"hoja": sheet_name, "fila_excel": fila_excel,
                             "inventario_id_declarado": inv, "codigo_declarado": item.get("codigo"),
                             "respuesta": answer, "datos_crudos": item, "diagnosticos": issues})
        # Los códigos de la etiqueta declarada no son altas automáticas.
        # Un hallazgo que reutiliza el código de un aparato oficial requiere
        # resolver la identidad; no se puede conciliar como equipo nuevo.
        codigos_base = defaultdict(list)
        for identificador, referencia in expected_by_id.items():
            codigo = norm(referencia.get("codigo_interno_original"))
            if codigo:
                codigos_base[codigo].append(identificador)
            original = text(referencia.get("codigo_interno_original"))
            if re.match(r"^[0-9]{2}", original) and original[:2] != f"{sucursal_id:02d}":
                diagnostics.append({
                    "severidad": "AVISO", "codigo": "PREFIJO_SUCURSAL_INCONSISTENTE",
                    "hoja": "APARATOS", "inventario_id": identificador,
                    "codigo_de_corte": original, "prefijo_esperado": f"{sucursal_id:02d}",
                })
        ignorar_etiquetas = {"", "sin codigo", "pendiente", "s/c", "n/a", "na"}
        for fila in rows:
            if fila["hoja"] != "HALLAZGOS":
                continue
            codigo_hallazgo = norm(fila["codigo_declarado"])
            if codigo_hallazgo in ignorar_etiquetas or codigo_hallazgo not in codigos_base:
                continue
            diagnostics.append({
                "severidad": "ERROR", "codigo": "CODIGO_HALLAZGO_EN_CORTE",
                "hoja": "HALLAZGOS", "fila": fila["fila_excel"],
                "codigo_declarado": fila["codigo_declarado"],
                "inventario_ids_existentes": codigos_base[codigo_hallazgo],
            })
            fila["diagnosticos"].append("CODIGO_HALLAZGO_EN_CORTE")

        # Una etiqueta nueva no puede identificar dos hallazgos ni dos
        # referencias provisionales de alta inicial.
        etiquetas_nuevas = defaultdict(list)
        for fila in rows:
            provisional = (fila["hoja"] == "HALLAZGOS" or (
                modalidad == "INICIAL" and fila["hoja"] == "APARATOS"
                and fila["inventario_id_declarado"] is None))
            if not provisional:
                continue
            etiqueta = norm(fila["codigo_declarado"])
            if etiqueta not in ignorar_etiquetas:
                etiquetas_nuevas[etiqueta].append(fila)
        for etiqueta, filas_repetidas in etiquetas_nuevas.items():
            if len(filas_repetidas) < 2:
                continue
            diagnostics.append({
                "severidad": "ERROR", "codigo": "ETIQUETA_NUEVA_DUPLICADA",
                "codigo_declarado": etiqueta,
                "ubicaciones": [
                    {"hoja": row["hoja"], "fila": row["fila_excel"]}
                    for row in filas_repetidas
                ],
            })
            for row in filas_repetidas:
                row["diagnosticos"].append("ETIQUETA_NUEVA_DUPLICADA")

        pending = []
        yes = no = 0
        for inv, expected_row in expected_by_id.items():
            observations = apparatus.get(inv, [])
            if len(observations) > 1:
                diagnostics.append({"severidad": "ERROR", "codigo": "ID_DUPLICADO_EN_EXCEL",
                                    "hoja": "APARATOS", "inventario_id": inv,
                                    "filas": [z[0] for z in observations]})
            valid = len(observations) == 1 and not observations[0][2]
            answer = observations[0][1] if valid else "SIN_REVISAR"
            if answer == "SI":
                yes += 1
            elif answer == "NO":
                no += 1
            else:
                pending.append({"inventario_id": inv,
                                "codigo_original": expected_row["codigo_interno_original"],
                                "fila_excel": observations[0][0] if len(observations) == 1 else None,
                                "motivo": "FILA_AUSENTE" if not observations else "SIN_RESPUESTA_O_IDENTIDAD"})
        # Códigos repetidos del corte son incidencias de identidad, pero no invalidan IDs distintos.
        duplicates = [code for code, count in Counter(
            norm(ref["codigo_interno_original"]) for ref in expected if text(ref["codigo_interno_original"])
        ).items() if count > 1]
        for code in duplicates:
            diagnostics.append({"severidad": "AVISO", "codigo": "CODIGO_REPETIDO_EN_BASE_HISTORICA", "valor": code})
        invalid = any(x["severidad"] == "ERROR" for x in diagnostics)
        status = "INVALID" if invalid else "INCOMPLETE" if pending or modalidad == "INICIAL" else "VALIDATED"
        summary = {
            "referencias_esperadas": len(expected_by_id), "si": yes, "no": no,
            "sin_revisar": len(pending), "pendientes": pending,
            "hallazgos": sum(x["hoja"] == "HALLAZGOS" for x in rows),
            "provisionales_aparatos": sum(
                x["hoja"] == "APARATOS" and x["inventario_id_declarado"] is None
                for x in rows) if modalidad == "INICIAL" else 0,
            "errores": sum(x["severidad"] == "ERROR" for x in diagnostics),
            "avisos": sum(x["severidad"] == "AVISO" for x in diagnostics),
            "requiere_revision_inicial": modalidad == "INICIAL",
        }
        return {"resultado": status, "control": ctl, "resumen": summary,
                "diagnosticos": diagnostics, "filas": rows}
    finally:
        book.close()
