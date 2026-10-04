"""Proyección XLSX pura para Campaign V2 Reporting."""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


XLSX_MIMETYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
_REPORT_TYPES = frozenset({"INDIVIDUAL", "CONSOLIDATED"})
_FORMULA_PREFIXES = ("=", "+", "-", "@")
_HEADER_FILL = PatternFill("solid", fgColor="211F1E")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")
_PERCENT_FORMAT = "0.00%"


def build_campaign_v2_reporting_excel(
    *,
    report_type: str,
    report: Mapping[str, Any],
    evolution: Sequence[Mapping[str, Any]],
    scope: Mapping[str, Any],
    generated_at: datetime,
) -> tuple[BytesIO, str]:
    """Construye el artefacto XLSX sin consultar DB ni recalcular KPIs."""
    normalized_type = str(report_type or "").strip().upper()
    if normalized_type not in _REPORT_TYPES:
        raise ValueError("report_type debe ser INDIVIDUAL o CONSOLIDATED.")
    generated_at_iso = _generated_at_iso(generated_at)

    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)

    _build_summary_sheet(
        workbook,
        report_type=normalized_type,
        report=report,
        generated_at_iso=generated_at_iso,
    )
    _build_campaigns_sheet(
        workbook,
        report_type=normalized_type,
        report=report,
    )
    _build_kpis_sheet(
        workbook,
        report_type=normalized_type,
        report=report,
    )
    _build_evolution_sheet(workbook, evolution=evolution)
    _build_metadata_sheet(
        workbook,
        report_type=normalized_type,
        report=report,
        scope=scope,
        generated_at_iso=generated_at_iso,
    )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    stamp = _filename_timestamp(generated_at)
    if normalized_type == "INDIVIDUAL":
        campaign_id = int(report["campaign"]["id"])
        filename = f"campaign_v2_{campaign_id}_report_{stamp}.xlsx"
    else:
        filename = f"campaign_v2_reporting_{stamp}.xlsx"
    return output, filename


def _build_summary_sheet(
    workbook: Workbook,
    *,
    report_type: str,
    report: Mapping[str, Any],
    generated_at_iso: str,
) -> None:
    sheet = workbook.create_sheet("Resumen")
    rows: list[tuple[str, Any, str | None]] = [
        ("generated_at", generated_at_iso, None),
        ("report_type", report_type, None),
    ]

    if report_type == "CONSOLIDATED":
        summary = report["summary"]
        rows.extend(
            [
                (
                    "latest_observed_at",
                    _latest_observed_at(report["campaigns"]),
                    None,
                ),
                ("campaign_count", summary["campaign_count"], None),
                (
                    "campaigns_with_snapshot",
                    summary["campaigns_with_snapshot"],
                    None,
                ),
                (
                    "campaigns_without_snapshot",
                    summary["campaigns_without_snapshot"],
                    None,
                ),
                (
                    "Exposiciones de destinatarios",
                    summary["total_recipients"],
                    None,
                ),
            ]
        )
        rows.extend(_summary_metric_rows(summary))
        rows.extend(_filter_rows(report.get("filters") or {}))
    else:
        rows.extend(
            [
                ("campaign_id", report["campaign"]["id"], None),
                ("campaign_name", report["campaign"]["name"], None),
                (
                    "latest_observed_at",
                    report["observation"]["latest_observed_at"],
                    None,
                ),
                (
                    "total_frozen_recipients",
                    report["audience"]["total_recipients"],
                    None,
                ),
            ]
        )
        rows.extend(_individual_metric_rows(report))

    _write_key_value_sheet(sheet, rows)


def _summary_metric_rows(
    summary: Mapping[str, Any],
) -> list[tuple[str, Any, str | None]]:
    normalized = summary["normalized"]
    rates = summary["rates"]
    coverage = summary["coverage"]
    interactions = summary["interactions"]
    return [
        ("successful", normalized["successful"], None),
        ("failed", normalized["failed"], None),
        ("sent", normalized["sent"], None),
        ("delivered", normalized["delivered"], None),
        ("viewed", normalized["viewed"], None),
        ("reach_count", normalized["reach_count"], None),
        ("successful_rate", rates["successful_rate"], _PERCENT_FORMAT),
        ("reach_rate", rates["reach_rate"], _PERCENT_FORMAT),
        ("read_rate", rates["read_rate"], _PERCENT_FORMAT),
        ("failure_rate", rates["failure_rate"], _PERCENT_FORMAT),
        (
            "matched_recipient_count",
            coverage["matched_recipient_count"],
            None,
        ),
        (
            "unmatched_provider_count",
            coverage["unmatched_provider_count"],
            None,
        ),
        (
            "frozen_without_status",
            coverage["frozen_recipient_without_provider_status_count"],
            None,
        ),
        (
            "status_coverage_rate",
            coverage["status_coverage_rate"],
            _PERCENT_FORMAT,
        ),
        (
            "button_interaction_recipient_exposures",
            interactions["button_interaction_recipient_exposures"],
            None,
        ),
        (
            "responders_aggregate",
            interactions["responders_aggregate"],
            None,
        ),
        (
            "responders_campaigns_with_value",
            interactions["responders_campaigns_with_value"],
            None,
        ),
        (
            "free_text_aggregate",
            interactions["free_text_aggregate"],
            None,
        ),
        (
            "free_text_campaigns_with_value",
            interactions["free_text_campaigns_with_value"],
            None,
        ),
        ("cost_status", summary["cost"]["status"], None),
        ("currency", summary["cost"]["currency"], None),
        ("cost_total", summary["cost"]["total"], None),
    ]


def _individual_metric_rows(
    report: Mapping[str, Any],
) -> list[tuple[str, Any, str | None]]:
    normalized = report["normalized"]
    rates = report["rates"]
    coverage = report["coverage"]
    interactions = report["interactions"]
    return [
        ("successful", normalized["successful"], None),
        ("failed", normalized["failed"], None),
        ("sent", normalized["sent"], None),
        ("delivered", normalized["delivered"], None),
        ("viewed", normalized["viewed"], None),
        ("reach_count", normalized["reach_count"], None),
        ("successful_rate", rates["successful_rate"], _PERCENT_FORMAT),
        ("reach_rate", rates["reach_rate"], _PERCENT_FORMAT),
        ("read_rate", rates["read_rate"], _PERCENT_FORMAT),
        ("failure_rate", rates["failure_rate"], _PERCENT_FORMAT),
        (
            "matched_recipient_count",
            coverage["matched_recipient_count"],
            None,
        ),
        (
            "unmatched_provider_count",
            coverage["unmatched_provider_count"],
            None,
        ),
        (
            "frozen_without_status",
            coverage["frozen_recipient_without_provider_status_count"],
            None,
        ),
        (
            "status_coverage_rate",
            coverage["status_coverage_rate"],
            _PERCENT_FORMAT,
        ),
        (
            "button_interaction_recipient_exposures",
            interactions["unique_button_recipients"],
            None,
        ),
        (
            "responders_aggregate",
            interactions["responders_aggregate"],
            None,
        ),
        (
            "free_text_aggregate",
            interactions["free_text_aggregate"],
            None,
        ),
        ("cost_status", report["cost"]["status"], None),
        ("currency", report["cost"]["currency"], None),
        ("cost_total", report["cost"]["total"], None),
    ]


def _filter_rows(
    filters: Mapping[str, Any],
) -> list[tuple[str, Any, str | None]]:
    return [
        (f"filter.{key}", value, None)
        for key, value in filters.items()
    ]


_CAMPAIGN_COLUMNS = (
    ("campaign_id", "campaign.id"),
    ("name", "campaign.name"),
    ("purpose", "campaign.purpose"),
    ("source", "campaign.source"),
    ("provider", "campaign.provider"),
    ("provider_campaign_id", "campaign.provider_campaign_id"),
    ("frozen_at", "campaign.frozen_at"),
    ("snapshot_id", "observation.snapshot_id"),
    ("latest_observed_at", "observation.latest_observed_at"),
    ("analytics_status", "observation.analytics_status"),
    ("total_recipients", "audience.total_recipients"),
    ("successful", "normalized.successful"),
    ("failed", "normalized.failed"),
    ("sent", "normalized.sent"),
    ("delivered", "normalized.delivered"),
    ("viewed", "normalized.viewed"),
    ("reach_count", "normalized.reach_count"),
    ("successful_rate", "rates.successful_rate"),
    ("reach_rate", "rates.reach_rate"),
    ("read_rate", "rates.read_rate"),
    ("failure_rate", "rates.failure_rate"),
    ("matched_recipient_count", "coverage.matched_recipient_count"),
    ("unmatched_provider_count", "coverage.unmatched_provider_count"),
    (
        "frozen_without_status",
        "coverage.frozen_recipient_without_provider_status_count",
    ),
    ("status_coverage_rate", "coverage.status_coverage_rate"),
    ("raw_successful", "provider_raw.successful"),
    ("raw_failed", "provider_raw.failed"),
    ("raw_sent", "provider_raw.sent"),
    ("raw_delivered", "provider_raw.delivered"),
    ("raw_viewed", "provider_raw.viewed"),
    ("raw_answered", "provider_raw.answered"),
    (
        "button_interaction_recipient_exposures",
        "interactions.unique_button_recipients",
    ),
    ("responders_aggregate", "interactions.responders_aggregate"),
    ("free_text_aggregate", "interactions.free_text_aggregate"),
    ("cost_status", "cost.status"),
    ("currency", "cost.currency"),
    ("cost_total", "cost.total"),
)
_CAMPAIGN_PERCENT_COLUMNS = {
    "successful_rate",
    "reach_rate",
    "read_rate",
    "failure_rate",
    "status_coverage_rate",
}


def _build_campaigns_sheet(
    workbook: Workbook,
    *,
    report_type: str,
    report: Mapping[str, Any],
) -> None:
    sheet = workbook.create_sheet("Campañas")
    rows = (
        [report]
        if report_type == "INDIVIDUAL"
        else list(report.get("campaigns") or [])
    )
    headers = [column for column, _ in _CAMPAIGN_COLUMNS]
    values = [
        [_nested_value(row, path) for _, path in _CAMPAIGN_COLUMNS]
        for row in rows
    ]
    _write_table(sheet, headers, values)
    for index, (column, _path) in enumerate(_CAMPAIGN_COLUMNS, start=1):
        if column in _CAMPAIGN_PERCENT_COLUMNS:
            _format_column(sheet, index, _PERCENT_FORMAT)


_KPI_DEFINITIONS = (
    (
        "successful_rate",
        "successful",
        "total_recipients",
        "SUITE_DERIVED",
        "successful / total_recipients",
    ),
    (
        "reach_rate",
        "reach_count",
        "total_recipients",
        "SUITE_DERIVED",
        "reach_count / total_recipients",
    ),
    (
        "read_rate",
        "viewed",
        "reach_count",
        "SUITE_DERIVED",
        "viewed / reach_count",
    ),
    (
        "failure_rate",
        "failed",
        "total_recipients",
        "SUITE_DERIVED",
        "failed / total_recipients",
    ),
    (
        "status_coverage_rate",
        "matched_recipient_count",
        "total_recipients",
        "SUITE_DERIVED",
        "matched_recipient_count / total_recipients",
    ),
)


def _build_kpis_sheet(
    workbook: Workbook,
    *,
    report_type: str,
    report: Mapping[str, Any],
) -> None:
    sheet = workbook.create_sheet("KPIs")
    if report_type == "CONSOLIDATED":
        source = report["summary"]
        total_recipients = source["total_recipients"]
        normalized = source["normalized"]
        rates = source["rates"]
        coverage = source["coverage"]
    else:
        total_recipients = report["audience"]["total_recipients"]
        normalized = report["normalized"]
        rates = report["rates"]
        coverage = report["coverage"]

    provider_raw = report["provider_raw"] if report_type == "INDIVIDUAL" else source["provider_raw"]
    interactions = report["interactions"] if report_type == "INDIVIDUAL" else source["interactions"]
    button_interactions = (
        interactions["unique_button_recipients"]
        if report_type == "INDIVIDUAL"
        else interactions["button_interaction_recipient_exposures"]
    )

    rows = [
        [metric, normalized[metric], None, None, "NORMALIZED", f"Normalized {metric}"]
        for metric in (
            "successful",
            "failed",
            "sent",
            "delivered",
            "viewed",
            "reach_count",
        )
    ]
    rows.append(
        [
            "button_interaction_recipient_exposures",
            button_interactions,
            None,
            None,
            "NORMALIZED",
            "Frozen recipient exposures with at least one button interaction",
        ]
    )
    for metric in (
        "successful",
        "failed",
        "sent",
        "delivered",
        "viewed",
        "answered",
    ):
        rows.append(
            [
                f"raw_{metric}",
                provider_raw.get(metric) if provider_raw is not None else None,
                None,
                None,
                "PROVIDER_RAW",
                f"Provider raw {metric}",
            ]
        )
    rows.extend(
        [
            [
                "responders_aggregate",
                interactions["responders_aggregate"],
                None,
                None,
                "PROVIDER_AGGREGATE",
                "Provider aggregate responders",
            ],
            [
                "free_text_aggregate",
                interactions["free_text_aggregate"],
                None,
                None,
                "PROVIDER_AGGREGATE",
                "Provider aggregate freeText",
            ],
        ]
    )

    numerators = {
        **normalized,
        "matched_recipient_count": coverage["matched_recipient_count"],
    }
    denominators = {
        "total_recipients": total_recipients,
        "reach_count": normalized["reach_count"],
    }
    for metric, numerator_key, denominator_key, metric_type, definition in (
        _KPI_DEFINITIONS
    ):
        rows.append(
            [
                metric,
                rates.get(metric)
                if metric != "status_coverage_rate"
                else coverage["status_coverage_rate"],
                numerators[numerator_key],
                denominators[denominator_key],
                metric_type,
                definition,
            ]
        )
    _write_table(
        sheet,
        [
            "metric",
            "value",
            "numerator",
            "denominator",
            "type",
            "definition",
        ],
        rows,
    )
    _format_column(sheet, 2, _PERCENT_FORMAT)


_EVOLUTION_COLUMNS = (
    ("campaign_id", "campaign_id"),
    ("campaign_name", "campaign_name"),
    ("snapshot_id", "snapshot_id"),
    ("observed_at", "observed_at"),
    ("normalized_successful", "normalized.successful"),
    ("normalized_failed", "normalized.failed"),
    ("normalized_sent", "normalized.sent"),
    ("normalized_delivered", "normalized.delivered"),
    ("normalized_viewed", "normalized.viewed"),
    ("normalized_reach", "normalized.reach_count"),
    ("raw_successful", "provider_raw.successful"),
    ("raw_failed", "provider_raw.failed"),
    ("raw_sent", "provider_raw.sent"),
    ("raw_delivered", "provider_raw.delivered"),
    ("raw_viewed", "provider_raw.viewed"),
    (
        "matched_recipient_count",
        "coverage.matched_recipient_count",
    ),
    (
        "unmatched_provider_count",
        "coverage.unmatched_provider_count",
    ),
)


def _build_evolution_sheet(
    workbook: Workbook,
    *,
    evolution: Sequence[Mapping[str, Any]],
) -> None:
    sheet = workbook.create_sheet("Evolución")
    headers = [column for column, _ in _EVOLUTION_COLUMNS]
    headers.append("observed_at_meaning")
    rows = [
        [
            *[_nested_value(row, path) for _, path in _EVOLUTION_COLUMNS],
            "Observed by Suite",
        ]
        for row in evolution
    ]
    _write_table(sheet, headers, rows)


def _build_metadata_sheet(
    workbook: Workbook,
    *,
    report_type: str,
    report: Mapping[str, Any],
    scope: Mapping[str, Any],
    generated_at_iso: str,
) -> None:
    sheet = workbook.create_sheet("Metadata")
    latest = (
        report["observation"]["latest_observed_at"]
        if report_type == "INDIVIDUAL"
        else _latest_observed_at(report.get("campaigns") or [])
    )
    campaign_count = (
        1
        if report_type == "INDIVIDUAL"
        else report["summary"]["campaign_count"]
    )
    rows: list[tuple[str, Any, str | None]] = [
        ("generated_at", generated_at_iso, None),
        ("report_type", report_type, None),
        ("campaign_count", campaign_count, None),
        ("latest_observed_at", latest, None),
        ("scope.is_global", bool(scope.get("is_global")), None),
        (
            "scope.allowed_sucursal_keys",
            _join_text_values(scope.get("allowed_sucursal_keys")),
            None,
        ),
        (
            "definition.total_recipients",
            "Consolidado = recipient exposures; individual = frozen recipients",
            None,
        ),
        (
            "definition.reach_count",
            "Unique frozen recipients in DELIVERED union VIEWED",
            None,
        ),
        (
            "definition.successful_rate",
            "successful / total_recipients",
            None,
        ),
        (
            "definition.reach_rate",
            "reach_count / total_recipients",
            None,
        ),
        (
            "definition.read_rate",
            "viewed / reach_count",
            None,
        ),
        (
            "definition.failure_rate",
            "failed / total_recipients",
            None,
        ),
        (
            "definition.status_coverage_rate",
            "matched / total_recipients",
            None,
        ),
        (
            "definition.provider_raw",
            "Provider raw counts are not normalized Campaign V2 recipients",
            None,
        ),
        (
            "definition.observed_at",
            "Observed by Suite (snapshot.fetched_at)",
            None,
        ),
        (
            "definition.cost",
            "Costos actualmente no disponibles",
            None,
        ),
    ]
    if report_type == "CONSOLIDATED":
        rows.extend(_filter_rows(report.get("filters") or {}))
    else:
        rows.append(
            ("filter.campaign_id", report["campaign"]["id"], None)
        )
    _write_key_value_sheet(sheet, rows)


def _latest_observed_at(
    campaign_rows: Sequence[Mapping[str, Any]],
) -> str | None:
    values = [
        row["observation"]["latest_observed_at"]
        for row in campaign_rows
        if row["observation"]["latest_observed_at"] is not None
    ]
    return max(values) if values else None


def _write_key_value_sheet(
    sheet,
    rows: Sequence[tuple[str, Any, str | None]],
) -> None:
    _write_table(
        sheet,
        ["field", "value"],
        [[field, value] for field, value, _fmt in rows],
    )
    for row_index, (_field, _value, number_format) in enumerate(
        rows,
        start=2,
    ):
        if number_format is not None and sheet.cell(row_index, 2).value is not None:
            sheet.cell(row_index, 2).number_format = number_format


def _write_table(
    sheet,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
) -> None:
    sheet.append([_safe_excel_value(value) for value in headers])
    for row in rows:
        sheet.append([_safe_excel_value(value) for value in row])
    _style_table(sheet)


def _style_table(sheet) -> None:
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _HEADER_ALIGNMENT
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    )
    sheet.sheet_view.showGridLines = False

    for index in range(1, sheet.max_column + 1):
        letter = get_column_letter(index)
        max_length = 0
        for cell in sheet[letter]:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        sheet.column_dimensions[letter].width = min(
            max(max_length + 2, 12),
            48,
        )


def _format_column(sheet, column_index: int, number_format: str) -> None:
    for row_index in range(2, sheet.max_row + 1):
        cell = sheet.cell(row_index, column_index)
        if cell.value is not None:
            cell.number_format = number_format


def _nested_value(
    row: Mapping[str, Any],
    path: str,
) -> Any:
    current: Any = row
    for part in path.split("."):
        if not isinstance(current, Mapping):
            return None
        current = current.get(part)
    return current


def _safe_excel_value(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    probe = value.lstrip(" \t\r\n\x00")
    if probe.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def _generated_at_iso(value: datetime) -> str:
    if not isinstance(value, datetime):
        raise ValueError("generated_at debe ser datetime timezone-aware.")
    if value.tzinfo is None or value.utcoffset() is None:
        raise ValueError("generated_at debe incluir zona horaria.")
    return value.astimezone(timezone.utc).isoformat()


def _filename_timestamp(value: datetime) -> str:
    if value.tzinfo is None or value.utcoffset() is None:
        raise ValueError("generated_at debe incluir zona horaria.")
    return value.astimezone(timezone.utc).strftime("%Y%m%d_%H%M%S")


def _join_text_values(value: Any) -> str | None:
    if value is None:
        return None
    if not isinstance(value, (list, tuple)):
        return str(value)
    return ", ".join(str(item) for item in value)
