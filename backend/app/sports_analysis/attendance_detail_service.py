from __future__ import annotations

from datetime import date
from io import BytesIO
from math import ceil
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import asc, desc

from app.extensions import db
from app.models.attendance import (
    WarehouseAttendanceVisitORM,
)
from app.models.sucursal_model import Sucursal

from .attendance_access import (
    SportsAnalysisScope,
    scoped_branch_catalog,
)
from .attendance_mart_service import (
    ALL_ATTENDANCE_TYPES,
)
from .attendance_query_service import (
    SportsAnalysisValidationError,
    _effective_branch_ids,
    _normalize_attendance_type,
    _validate_date_range,
)


DETAIL_EXPORT_MIMETYPE = (
    "application/vnd.openxmlformats-officedocument."
    "spreadsheetml.sheet"
)
DETAIL_TIMEZONE = ZoneInfo("America/Tijuana")
DEFAULT_PAGE_SIZE = 50
MAX_PAGE_SIZE = 200

_VISITS_COLUMNS = [
    {"key": "business_date", "label": "Fecha"},
    {"key": "branch_name", "label": "Sucursal"},
    {"key": "member_pin", "label": "PIN"},
    {"key": "entered_at", "label": "Entrada"},
    {"key": "exited_at", "label": "Salida"},
    {"key": "visit_status", "label": "Estado"},
    {"key": "duration_minutes", "label": "Estancia (min)"},
    {"key": "age", "label": "Edad"},
    {"key": "city", "label": "Ciudad"},
    {"key": "postal_code", "label": "CP"},
    {"key": "member_since", "label": "Alta"},
    {"key": "attendance_type", "label": "Tipo asistencia"},
    {"key": "has_opening", "label": "Tiene apertura"},
    {"key": "source_branch_name", "label": "Sucursal origen"},
]

_VISITS_SORTS = {
    "business_date": WarehouseAttendanceVisitORM.business_date,
    "branch_name": Sucursal.sucursal,
    "member_pin": WarehouseAttendanceVisitORM.member_pin,
    "entered_at": WarehouseAttendanceVisitORM.entered_at_utc,
    "exited_at": WarehouseAttendanceVisitORM.exited_at_utc,
    "visit_status": WarehouseAttendanceVisitORM.visit_status,
    "duration_minutes": WarehouseAttendanceVisitORM.duration_seconds,
    "age": WarehouseAttendanceVisitORM.age,
    "city": WarehouseAttendanceVisitORM.city,
    "postal_code": WarehouseAttendanceVisitORM.postal_code,
    "member_since": WarehouseAttendanceVisitORM.member_since,
    "attendance_type": WarehouseAttendanceVisitORM.attendance_type,
    "has_opening": WarehouseAttendanceVisitORM.has_opening,
    "source_branch_name": WarehouseAttendanceVisitORM.source_branch_name,
}


def build_attendance_detail(
    scope: SportsAnalysisScope,
    *,
    metric: str,
    date_from: date,
    date_to: date,
    branch_id: int | None,
    region_key: str | None,
    attendance_type: str | None,
    page: object = 1,
    page_size: object = DEFAULT_PAGE_SIZE,
    sort_by: str | None = None,
    sort_dir: str | None = None,
) -> dict:
    normalized_metric = _normalize_metric(metric)
    if normalized_metric != "visits":
        raise SportsAnalysisValidationError(
            "El detalle solicitado todavía no está disponible."
        )

    safe_page = _parse_positive_int(page, "page", default=1)
    safe_page_size = min(
        _parse_positive_int(
            page_size,
            "page_size",
            default=DEFAULT_PAGE_SIZE,
        ),
        MAX_PAGE_SIZE,
    )
    safe_sort_by = _normalize_sort_by(sort_by)
    safe_sort_dir = _normalize_sort_dir(sort_dir)

    query = _visits_query(
        scope,
        date_from=date_from,
        date_to=date_to,
        branch_id=branch_id,
        region_key=region_key,
        attendance_type=attendance_type,
    )

    count = query.order_by(None).count()
    total_pages = max(1, ceil(count / safe_page_size))
    safe_page = min(safe_page, total_pages)

    rows = (
        _apply_visits_sort(
            query,
            sort_by=safe_sort_by,
            sort_dir=safe_sort_dir,
        )
        .offset((safe_page - 1) * safe_page_size)
        .limit(safe_page_size)
        .all()
    )

    return {
        "metric": normalized_metric,
        "title": "Visitas registradas",
        "count": int(count),
        "page": safe_page,
        "page_size": safe_page_size,
        "total_pages": total_pages,
        "sort_by": safe_sort_by,
        "sort_dir": safe_sort_dir,
        "columns": list(_VISITS_COLUMNS),
        "rows": [
            _serialize_visit(visit, branch_name)
            for visit, branch_name in rows
        ],
    }


def build_attendance_detail_export(
    scope: SportsAnalysisScope,
    *,
    metric: str,
    date_from: date,
    date_to: date,
    branch_id: int | None,
    region_key: str | None,
    attendance_type: str | None,
    sort_by: str | None = None,
    sort_dir: str | None = None,
) -> tuple[BytesIO, str]:
    normalized_metric = _normalize_metric(metric)
    if normalized_metric != "visits":
        raise SportsAnalysisValidationError(
            "La exportación solicitada todavía no está disponible."
        )

    safe_sort_by = _normalize_sort_by(sort_by)
    safe_sort_dir = _normalize_sort_dir(sort_dir)

    rows = (
        _apply_visits_sort(
            _visits_query(
                scope,
                date_from=date_from,
                date_to=date_to,
                branch_id=branch_id,
                region_key=region_key,
                attendance_type=attendance_type,
            ),
            sort_by=safe_sort_by,
            sort_dir=safe_sort_dir,
        )
        .all()
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Visitas"

    headers = [column["label"] for column in _VISITS_COLUMNS]
    worksheet.append(headers)

    for visit, branch_name in rows:
        serialized = _serialize_visit(visit, branch_name)
        worksheet.append(
            [serialized[column["key"]] for column in _VISITS_COLUMNS]
        )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F2937",
    )
    header_font = Font(
        color="FFFFFF",
        bold=True,
    )
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    widths = {
        "A": 13,
        "B": 24,
        "C": 15,
        "D": 21,
        "E": 21,
        "F": 16,
        "G": 16,
        "H": 10,
        "I": 20,
        "J": 12,
        "K": 13,
        "L": 22,
        "M": 16,
        "N": 24,
    }
    for column_letter, width in widths.items():
        worksheet.column_dimensions[column_letter].width = width

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    date_fragment = (
        date_from.isoformat()
        if date_from == date_to
        else f"{date_from.isoformat()}_{date_to.isoformat()}"
    )
    filename = (
        f"aforo_asistencia_visitas_{date_fragment}.xlsx"
    )
    return output, filename


def _visits_query(
    scope: SportsAnalysisScope,
    *,
    date_from: date,
    date_to: date,
    branch_id: int | None,
    region_key: str | None,
    attendance_type: str | None,
):
    branches = scoped_branch_catalog(scope)
    branch_ids = _effective_branch_ids(
        branches,
        branch_id=branch_id,
        region_key=region_key,
    )
    normalized_type = _normalize_attendance_type(
        attendance_type
    )
    _validate_date_range(date_from, date_to)

    query = (
        db.session.query(
            WarehouseAttendanceVisitORM,
            Sucursal.sucursal,
        )
        .join(
            Sucursal,
            Sucursal.sucursal_id
            == WarehouseAttendanceVisitORM.sucursal_id,
        )
        .filter(
            WarehouseAttendanceVisitORM.sucursal_id.in_(
                branch_ids or (-1,)
            ),
            WarehouseAttendanceVisitORM.business_date.between(
                date_from,
                date_to,
            ),
        )
    )

    if normalized_type != ALL_ATTENDANCE_TYPES:
        query = query.filter(
            WarehouseAttendanceVisitORM.attendance_type
            == normalized_type
        )

    return query


def _apply_visits_sort(
    query,
    *,
    sort_by: str,
    sort_dir: str,
):
    sort_expression = _VISITS_SORTS[sort_by]
    order = asc if sort_dir == "asc" else desc
    return query.order_by(
        order(sort_expression),
        asc(WarehouseAttendanceVisitORM.id),
    )


def _serialize_visit(
    visit: WarehouseAttendanceVisitORM,
    branch_name: str,
) -> dict:
    duration_minutes = (
        round(float(visit.duration_seconds) / 60, 1)
        if visit.duration_seconds is not None
        else None
    )

    return {
        "business_date": _format_date(visit.business_date),
        "branch_name": str(branch_name or visit.source_branch_name),
        "member_pin": visit.member_pin,
        "entered_at": _format_datetime(visit.entered_at_utc),
        "exited_at": _format_datetime(visit.exited_at_utc),
        "visit_status": visit.visit_status,
        "duration_minutes": duration_minutes,
        "age": int(visit.age) if visit.age is not None else None,
        "city": visit.city,
        "postal_code": visit.postal_code,
        "member_since": _format_date(visit.member_since),
        "attendance_type": visit.attendance_type,
        "has_opening": (
            "Sí"
            if visit.has_opening is True
            else "No"
            if visit.has_opening is False
            else None
        ),
        "source_branch_name": visit.source_branch_name,
    }


def _format_date(value) -> str | None:
    if value is None:
        return None
    return value.strftime("%d/%m/%Y")


def _format_datetime(value) -> str | None:
    if value is None:
        return None
    return (
        value.astimezone(DETAIL_TIMEZONE)
        .strftime("%d/%m/%Y %H:%M:%S")
    )


def _normalize_metric(value: str) -> str:
    return str(value or "").strip().lower()


def _normalize_sort_by(value: str | None) -> str:
    normalized = str(value or "entered_at").strip()
    if normalized not in _VISITS_SORTS:
        raise SportsAnalysisValidationError(
            "sort_by inválido para el detalle de visitas."
        )
    return normalized


def _normalize_sort_dir(value: str | None) -> str:
    normalized = str(value or "asc").strip().lower()
    if normalized not in {"asc", "desc"}:
        raise SportsAnalysisValidationError(
            "sort_dir debe ser asc o desc."
        )
    return normalized


def _parse_positive_int(
    value: object,
    name: str,
    *,
    default: int,
) -> int:
    if value in (None, ""):
        return default
    try:
        parsed = int(value)
    except (TypeError, ValueError) as exc:
        raise SportsAnalysisValidationError(
            f"{name} inválido."
        ) from exc
    if parsed <= 0:
        raise SportsAnalysisValidationError(
            f"{name} inválido."
        )
    return parsed
