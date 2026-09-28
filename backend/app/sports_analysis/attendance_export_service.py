from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


ATTENDANCE_EXPORT_MIMETYPE = (
    "application/vnd.openxmlformats-officedocument."
    "spreadsheetml.sheet"
)


def build_attendance_workbook(
    *,
    columns: list[dict],
    rows: list[dict],
    sheet_name: str,
) -> BytesIO:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name[:31]

    worksheet.append(
        [
            column["label"]
            for column in columns
        ]
    )
    for row in rows:
        worksheet.append(
            [
                row.get(column["key"])
                for column in columns
            ]
        )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F2937",
    )
    header_font = Font(
        color="FFFFFF",
        bold=True,
    )
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    for index, column in enumerate(
        columns,
        start=1,
    ):
        sample_width = max(
            [len(str(column["label"]))]
            + [
                len(
                    str(
                        row.get(
                            column["key"]
                        )
                        or ""
                    )
                )
                for row in rows[:200]
            ]
        )
        worksheet.column_dimensions[
            _excel_column_letter(index)
        ].width = min(
            max(sample_width + 2, 12),
            28,
        )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _excel_column_letter(index: int) -> str:
    result = ""
    value = index
    while value:
        value, remainder = divmod(
            value - 1,
            26,
        )
        result = (
            chr(65 + remainder)
            + result
        )
    return result
