"""Strict Google Ads Spanish ad-group-by-day parser with two action columns.

Workbook contains 43 groups × 15 days plus daily and undated subtotals.
The parser intentionally stores spend once, not per conversion action.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
import re
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from app.warehouse.services.google_ads_daily_xlsx_parser import (
    GoogleAdsXlsxValidationError,
    _campaign_key,
    _date_value,
    _norm,
    _number,
    _integer,
    _report_range,
)

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ROWS = 12000
NEEDED = frozenset({
    "dia", "grupo de anuncios", "campana", "codigo de moneda", "coste",
    "clics", "impr.", "conversiones", "pase gratis ads",
    "compras web ads", "valor de conv. original",
})


@dataclass(frozen=True)
class GoogleAdsAdGroupDailyRow:
    report_date: date
    campaign_name: str
    campaign_key: str
    ad_group_name: str
    ad_group_key: str
    currency_code: str
    cost: Decimal
    clicks: int
    impressions: int
    conversions: Decimal
    free_pass_conversions: Decimal
    web_purchase_conversions: Decimal
    conversion_value: Decimal


@dataclass(frozen=True)
class GoogleAdsAdGroupDailyFile:
    date_from: date
    date_to: date
    currency_code: str
    rows: tuple[GoogleAdsAdGroupDailyRow, ...]
    total_cost: Decimal
    total_free_pass: Decimal
    total_web_purchase: Decimal


def _ad_group_key(campaign_key: str, ad_group_name: str) -> str:
    normalized = _norm(ad_group_name)
    if not normalized:
        raise GoogleAdsXlsxValidationError("Grupo de anuncios vacío.")
    return "name:" + sha256(
        f"{campaign_key}\x1f{normalized}".encode("utf-8")
    ).hexdigest()


def _action_number(value: object, field: str) -> Decimal:
    # Google's custom-column XLSX writes strings like "8,83"; normal cells
    # elsewhere in the same workbook are genuine numeric cells.
    if isinstance(value, str) and re.fullmatch(r"[0-9]+,[0-9]{1,4}", value.strip()):
        value = value.strip().replace(",", ".")
    return _number(value, field, "0.0001")


def _metrics(values: tuple, index: dict[str, int]) -> tuple:
    def cell(column: str):
        position = index[column]
        return values[position] if position < len(values) else None

    currencies = [
        str(values[i] or "").strip().upper()
        for i in index["_currency_positions"]
    ]
    if any(item != "MXN" for item in currencies):
        raise GoogleAdsXlsxValidationError("Solo se admite moneda MXN.")
    return (
        _number(cell("coste"), "coste", "0.01"),
        _integer(cell("clics"), "clics"),
        _integer(cell("impr."), "impresiones"),
        _action_number(cell("conversiones"), "conversiones"),
        _action_number(cell("pase gratis ads"), "pases"),
        _action_number(cell("compras web ads"), "compras"),
        _number(cell("valor de conv. original"), "valor de conv.", "0.01"),
    )


def parse_google_ads_ad_group_daily_xlsx(content: bytes) -> GoogleAdsAdGroupDailyFile:
    if not isinstance(content, bytes) or not 0 < len(content) <= MAX_FILE_BYTES:
        raise GoogleAdsXlsxValidationError("Archivo XLSX vacío o mayor a 10 MB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise GoogleAdsXlsxValidationError("XLSX excede límite descomprimido.")
        workbook = load_workbook(
            BytesIO(content), read_only=True, data_only=True, keep_links=False,
        )
    except (BadZipFile, OSError, KeyError, ValueError) as exc:
        raise GoogleAdsXlsxValidationError("Archivo XLSX inválido.") from exc
    try:
        if len(workbook.sheetnames) != 1:
            raise GoogleAdsXlsxValidationError("Se espera una sola hoja.")
        sheet = workbook.active
        if sheet.max_row is not None and sheet.max_row > MAX_ROWS:
            raise GoogleAdsXlsxValidationError("Demasiadas filas.")
        records = sheet.iter_rows(values_only=True)
        title = next(records, ())
        period = next(records, ())
        columns = next(records, ())
        if not title or _norm(title[0]) != "informe de grupo de anuncios":
            raise GoogleAdsXlsxValidationError("No es un informe de grupos de Google Ads.")
        if not period:
            raise GoogleAdsXlsxValidationError("Periodo ausente.")
        start, end = _report_range(period[0])
        positions: dict[str, list[int]] = {}
        for position, heading in enumerate(columns):
            if heading is not None:
                positions.setdefault(_norm(heading), []).append(position)
        missing = NEEDED - positions.keys()
        if missing:
            raise GoogleAdsXlsxValidationError(
                "Columnas faltantes: " + ", ".join(sorted(missing))
            )
        if any(len(v) != 1 for key, v in positions.items() if key != "codigo de moneda"):
            raise GoogleAdsXlsxValidationError("Columnas duplicadas en el Excel.")
        index = {key: positions[key][0] for key in NEEDED}
        index["_currency_positions"] = positions["codigo de moneda"]

        parsed = []
        unique = set()
        daily = {}
        daily_groups = {}
        daily_account = {}
        observed_dates = set()
        keys_per_day = {}
        for position, values in enumerate(records, start=4):
            if position > MAX_ROWS:
                raise GoogleAdsXlsxValidationError("Demasiadas filas.")
            if not values or all(v is None for v in values):
                continue
            status = str(values[1] or "").strip()
            raw_date = values[index["dia"]]
            if status.startswith("Total:"):
                if status in ("Total: Cuenta", "Total: Grupos de anuncios"):
                    if raw_date is not None and str(raw_date).strip() not in ("", "--"):
                        day = _date_value(raw_date)
                        if not start <= day <= end:
                            raise GoogleAdsXlsxValidationError("Total fuera de rango.")
                        target = daily_account if status == "Total: Cuenta" else daily_groups
                        if day in target:
                            raise GoogleAdsXlsxValidationError("Total diario duplicado.")
                        target[day] = _metrics(values, index)
                continue

            day = _date_value(raw_date)
            if not start <= day <= end:
                raise GoogleAdsXlsxValidationError(f"Fila {position}: día fuera del periodo.")
            campaign_name = str(values[index["campana"]] or "").strip()
            group_name = str(values[index["grupo de anuncios"]] or "").strip()
            if not campaign_name or len(campaign_name) > 255 or campaign_name == "--":
                raise GoogleAdsXlsxValidationError(f"Fila {position}: campaña inválida.")
            if not group_name or len(group_name) > 255 or group_name == "--":
                raise GoogleAdsXlsxValidationError(f"Fila {position}: grupo inválido.")
            campaign_key = _campaign_key(campaign_name)
            group_key = _ad_group_key(campaign_key, group_name)
            key = (day, campaign_key, group_key)
            if key in unique:
                raise GoogleAdsXlsxValidationError(f"Fila {position}: grupo y día duplicados.")
            unique.add(key)
            observed_dates.add(day)
            keys_per_day.setdefault(day, set()).add((campaign_key, group_key))
            cost, clicks, impressions, conv, passes, purchases, value = _metrics(values, index)
            parsed.append(GoogleAdsAdGroupDailyRow(
                report_date=day, campaign_name=campaign_name,
                campaign_key=campaign_key, ad_group_name=group_name,
                ad_group_key=group_key, currency_code="MXN",
                cost=cost, clicks=clicks, impressions=impressions,
                conversions=conv, free_pass_conversions=passes,
                web_purchase_conversions=purchases, conversion_value=value,
            ))
            prev = daily.get(day, (Decimal("0"), 0, 0, Decimal("0"), Decimal("0"), Decimal("0"), Decimal("0")))
            daily[day] = tuple(a + b for a, b in zip(
                prev, (cost, clicks, impressions, conv, passes, purchases, value)
            ))

        expected = {
            date.fromordinal(start.toordinal() + i)
            for i in range((end - start).days + 1)
        }
        if not parsed or observed_dates != expected:
            raise GoogleAdsXlsxValidationError("Faltan días o grupos.")
        if set(daily_groups) != expected or set(daily_account) != expected:
            raise GoogleAdsXlsxValidationError("Faltan subtotales diarios para reconciliar.")
        first_day_keys = keys_per_day[start]
        if any(keys != first_day_keys for keys in keys_per_day.values()):
            raise GoogleAdsXlsxValidationError(
                "No coincide el catálogo de grupos entre días."
            )
        for day in expected:
            if daily[day] != daily_groups[day] or daily[day] != daily_account[day]:
                raise GoogleAdsXlsxValidationError(
                    f"No concilian los totales de Google Ads del {day.isoformat()}."
                )
        return GoogleAdsAdGroupDailyFile(
            date_from=start, date_to=end, currency_code="MXN",
            rows=tuple(parsed),
            total_cost=sum((item.cost for item in parsed), Decimal("0")),
            total_free_pass=sum(
                (item.free_pass_conversions for item in parsed), Decimal("0")
            ),
            total_web_purchase=sum(
                (item.web_purchase_conversions for item in parsed), Decimal("0")
            ),
        )
    finally:
        workbook.close()
