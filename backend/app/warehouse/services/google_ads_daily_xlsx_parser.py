"""Strict parser for Google Ads Spanish campaign-by-day XLSX exports.

No database or Flask dependencies: all validation precedes structured writes.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
import re
import unicodedata
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook


MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 40 * 1024 * 1024
MAX_ROWS = 3000
MAX_DAYS = 31
MONTHS = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5,
    "junio": 6, "julio": 7, "agosto": 8, "septiembre": 9,
    "setiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12,
}
NEEDED = {
    "dia", "campana", "codigo de moneda", "coste", "conversiones",
    "valor de conv.", "clics", "impr.",
}


class GoogleAdsXlsxValidationError(ValueError):
    """Invalid export; nothing structured may be persisted."""


@dataclass(frozen=True)
class GoogleAdsDailyParsedRow:
    report_date: date
    campaign_name: str
    campaign_key: str
    currency_code: str
    cost: Decimal
    impressions: int
    clicks: int
    conversions: Decimal
    conversion_value: Decimal


@dataclass(frozen=True)
class GoogleAdsDailyParsedFile:
    date_from: date
    date_to: date
    currency_code: str
    rows: tuple[GoogleAdsDailyParsedRow, ...]
    total_cost: Decimal


def _norm(value: object) -> str:
    raw = unicodedata.normalize("NFKD", str(value or ""))
    without_accents = "".join(
        char for char in raw if not unicodedata.combining(char)
    )
    return " ".join(without_accents.casefold().split())


def _date_value(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            pass
    raise GoogleAdsXlsxValidationError("La columna Día debe tener fechas ISO YYYY-MM-DD.")


def _report_range(value: object) -> tuple[date, date]:
    pattern = (
        r"^(\d{1,2}) de ([a-z]+) de (\d{4})"
        r" - (\d{1,2}) de ([a-z]+) de (\d{4})$"
    )
    matched = re.fullmatch(pattern, _norm(value))
    if not matched:
        raise GoogleAdsXlsxValidationError(
            "No se reconoció el periodo del encabezado Google Ads."
        )
    d1, m1, y1, d2, m2, y2 = matched.groups()
    if m1 not in MONTHS or m2 not in MONTHS:
        raise GoogleAdsXlsxValidationError("Mes inválido en el encabezado.")
    try:
        start = date(int(y1), MONTHS[m1], int(d1))
        end = date(int(y2), MONTHS[m2], int(d2))
    except ValueError as exc:
        raise GoogleAdsXlsxValidationError("Fechas inválidas.") from exc
    if end < start or (end - start).days >= MAX_DAYS:
        raise GoogleAdsXlsxValidationError("Periodo inválido o mayor a 31 días.")
    return start, end


def _number(value: object, field: str, scale: str) -> Decimal:
    if isinstance(value, bool) or value is None or isinstance(value, str) and not value.strip():
        raise GoogleAdsXlsxValidationError(f"Falta valor numérico: {field}.")
    try:
        n = Decimal(str(value))
        if not n.is_finite() or n < 0 or n != n.quantize(Decimal(scale)):
            raise InvalidOperation
        if n.adjusted() > 15:
            raise InvalidOperation
    except (InvalidOperation, ValueError, TypeError):
        raise GoogleAdsXlsxValidationError(f"Número inválido en {field}.") from None
    return n.quantize(Decimal(scale))


def _integer(value: object, field: str) -> int:
    result = _number(value, field, "1")
    if result > 2**63 - 1:
        raise GoogleAdsXlsxValidationError(f"Entero fuera de rango: {field}.")
    return int(result)


def _campaign_key(name: str) -> str:
    normalized = _norm(name)
    if not normalized:
        raise GoogleAdsXlsxValidationError("Campaña vacía.")
    return "name:" + sha256(normalized.encode("utf-8")).hexdigest()


def parse_google_ads_daily_xlsx(content: bytes) -> GoogleAdsDailyParsedFile:
    if not isinstance(content, bytes) or not 0 < len(content) <= MAX_FILE_BYTES:
        raise GoogleAdsXlsxValidationError("Archivo XLSX vacío o mayor a 8 MB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise GoogleAdsXlsxValidationError("XLSX excede límite descomprimido.")
        workbook = load_workbook(
            BytesIO(content), read_only=True, data_only=True, keep_links=False,
        )
    except (BadZipFile, OSError, ValueError, KeyError) as exc:
        raise GoogleAdsXlsxValidationError("Archivo XLSX inválido.") from exc

    try:
        if len(workbook.sheetnames) != 1:
            raise GoogleAdsXlsxValidationError("Se espera una sola hoja.")
        sheet = workbook.active
        if sheet.max_row is not None and sheet.max_row > MAX_ROWS:
            raise GoogleAdsXlsxValidationError("Demasiadas filas.")
        rows = sheet.iter_rows(values_only=True)
        title = next(rows, ())
        period = next(rows, ())
        headings = next(rows, ())
        if not title or _norm(title[0]) != "informe de campana":
            raise GoogleAdsXlsxValidationError("No es un Informe de campaña Google Ads.")
        if not period:
            raise GoogleAdsXlsxValidationError("Falta periodo.")
        date_from, date_to = _report_range(period[0])
        index = {_norm(col): i for i, col in enumerate(headings) if col is not None}
        if len(index) != len([col for col in headings if col is not None]):
            raise GoogleAdsXlsxValidationError("Columnas duplicadas.")
        missing = sorted(NEEDED - index.keys())
        if missing:
            raise GoogleAdsXlsxValidationError("Columnas faltantes: " + ", ".join(missing))
        parsed = []
        unique_keys = set()
        account_totals = {}
        currencies = set()
        observed_days = set()
        daily_sum = {}
        for position, values in enumerate(rows, start=4):
            if position > MAX_ROWS:
                raise GoogleAdsXlsxValidationError("Demasiadas filas.")
            if not values or all(cell is None for cell in values):
                continue
            def cell(column):
                return values[index[column]] if index[column] < len(values) else None

            status = str(values[1] or "").strip()
            day = values[index["dia"]]
            if status.startswith("Total:"):
                if status == "Total: Cuenta" and day not in (None, "--"):
                    parsed_day = _date_value(day)
                    if parsed_day in account_totals:
                        raise GoogleAdsXlsxValidationError("Total de cuenta duplicado.")
                    account_totals[parsed_day] = (
                        _number(cell("coste"), "total cuenta", "0.01"),
                        _integer(cell("clics"), "clics cuenta"),
                        _integer(cell("impr."), "impresiones cuenta"),
                    )
                continue
            report_date = _date_value(day)
            if not date_from <= report_date <= date_to:
                raise GoogleAdsXlsxValidationError(
                    f"Fila {position}: fecha fuera del periodo."
                )
            name = str(cell("campana") or "").strip()
            if not name or len(name) > 255 or name == "--":
                raise GoogleAdsXlsxValidationError(f"Fila {position}: campaña inválida.")
            campaign_key = _campaign_key(name)
            unique = (report_date, campaign_key)
            if unique in unique_keys:
                raise GoogleAdsXlsxValidationError(
                    f"Fila {position}: campaña y fecha duplicadas."
                )
            unique_keys.add(unique)
            currency = str(cell("codigo de moneda") or "").strip().upper()
            if not re.fullmatch(r"[A-Z]{3}", currency):
                raise GoogleAdsXlsxValidationError(f"Fila {position}: moneda inválida.")
            currencies.add(currency)
            metrics = GoogleAdsDailyParsedRow(
                report_date=report_date,
                campaign_name=name,
                campaign_key=campaign_key,
                currency_code=currency,
                cost=_number(cell("coste"), "coste", "0.01"),
                conversions=_number(cell("conversiones"), "conversiones", "0.0001"),
                conversion_value=_number(cell("valor de conv."), "valor conv.", "0.01"),
                clicks=_integer(cell("clics"), "clics"),
                impressions=_integer(cell("impr."), "impresiones"),
            )
            parsed.append(metrics)
            observed_days.add(report_date)
            prev = daily_sum.get(report_date, (Decimal("0"), 0, 0))
            daily_sum[report_date] = (
                prev[0] + metrics.cost,
                prev[1] + metrics.clicks,
                prev[2] + metrics.impressions,
            )
        if not parsed or len(currencies) != 1 or currencies != {"MXN"}:
            raise GoogleAdsXlsxValidationError(
                "No hay campañas o la moneda no es exclusivamente MXN."
            )
        expected_days = {
            date.fromordinal(date_from.toordinal() + offset)
            for offset in range((date_to - date_from).days + 1)
        }
        if observed_days != expected_days:
            raise GoogleAdsXlsxValidationError("Faltan días del periodo solicitado.")
        if set(account_totals) != expected_days:
            raise GoogleAdsXlsxValidationError(
                "Faltan totales diarios de cuenta para conciliación."
            )
        for day, amount in daily_sum.items():
            if amount != account_totals[day]:
                raise GoogleAdsXlsxValidationError(
                    f"No coincide el total de cuenta del {day.isoformat()}."
                )
        return GoogleAdsDailyParsedFile(
            date_from=date_from, date_to=date_to,
            currency_code="MXN", rows=tuple(parsed),
            total_cost=sum((row.cost for row in parsed), Decimal("0.00")),
        )
    finally:
        workbook.close()
